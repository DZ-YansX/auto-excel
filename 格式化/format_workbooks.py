import argparse
import os
import re
import tempfile
from copy import copy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator


MONTH = re.compile(r"^\s*(1[0-2]|[1-9])\s*月\s*$")
SHEET_NAMES = {
    "profit": ("利润表",),
    "sales": ("销售明细", "销售明细表"),
    "bill": ("账单明细", "账单原表", "收入明细表"),
    "order": ("订单明细", "订单明细表", "订单明细原表"),
    "ads": ("广告", "广告费"),
    "makeup": ("补单明细",),
    "unmatched": ("未匹配",),
    "big_seller": ("大卖家",),
}
EMPTY_SHEETS = ("bill", "order", "ads", "makeup", "big_seller")
CANCEL_TYPES = ("取消未发货", "取消已发货")
CELL_XML = re.compile(rb'<c\b(?P<attrs>[^>]*)>(?P<body>.*?)</c>')
CELL_REF = re.compile(rb'\br="([^"]+)"')
FORMULA_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def get_sheet(workbook, key, required=True):
    aliases = SHEET_NAMES[key]
    for ws in workbook.worksheets:
        if ws.title.strip() in aliases:
            return ws
    if required:
        raise KeyError(f"缺少工作表：{' / '.join(aliases)}")
    return None


def preserve_linked_profit_values(ws, cached_ws, cleared_titles):
    for cell in ws._cells.values():
        if cell.data_type != "f":
            continue
        for token in Tokenizer(cell.value).items:
            if token.type != "OPERAND" or token.subtype != "RANGE" or "!" not in token.value:
                continue
            title = token.value.rsplit("!", 1)[0]
            if title.startswith("'") and title.endswith("'"):
                title = title[1:-1].replace("''", "'")
            if title in cleared_titles:
                value = cached_ws[cell.coordinate].value
                if value is None:
                    raise ValueError(f"{ws.title}!{cell.coordinate}: 关联明细的公式缺少原始计算结果")
                cell.value = value
                break


def add_next_month(ws):
    months = []
    for col in range(2, ws.max_column + 2):
        match = MONTH.match(str(ws.cell(2, col).value or ""))
        if match:
            months.append((col, int(match.group(1))))
    if not months:
        raise ValueError(f"{ws.title}: 第二行未找到月份列")

    source_col, month = months[-1]
    target_col = source_col + 1
    formula_rows = set()
    profit_row = None
    for row in range(3, ws.max_row + 1):
        label = re.sub(r"\s+", "", str(ws.cell(row, 1).value or ""))
        if "本期利润" in label:
            profit_row = row
        if (
            label.startswith(("一、实际主营收入", "二、实际主营成本", "三、营业费用", "四、本期利润"))
            or any(name.casefold() in label.casefold() for name in ("推广费用", "Shopee结算", "TK结算", "团队运营费用"))
        ):
            formula_rows.add(row)
    if profit_row is None:
        raise ValueError(f"{ws.title}: 未找到本期利润行")

    for row in range(2, profit_row + 1):
        source = ws.cell(row, source_col)
        target = ws.cell(row, target_col)
        target._style = copy(source._style)
        target.number_format = source.number_format
        if source.has_style:
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)

        value = source.value
        if row in formula_rows and isinstance(value, str) and value.startswith("="):
            target.value = Translator(value, origin=source.coordinate).translate_formula(target.coordinate)
        else:
            target.value = None

    ws.cell(2, target_col).value = f"{month % 12 + 1}月"
    for row in range(3, profit_row):
        cell = ws.cell(row, target_col)
        font = copy(cell.font)
        font.bold = True
        cell.font = font
        alignment = copy(cell.alignment)
        alignment.horizontal = "right"
        alignment.vertical = "center"
        cell.alignment = alignment
        cell.number_format = "0.00_ "


def clear_sales(ws):
    row_values = {}
    formula_cells = {}
    for cell in ws._cells.values():
        if cell.row >= 2 and cell.value is not None:
            row_values.setdefault(cell.row, []).append(cell.value)
            formula_cells.setdefault(cell.row, []).append(cell)
    total_rows = {
        row for row, values in row_values.items()
        if all(isinstance(value, str) and value.startswith("=SUM(") for value in values)
    }
    removable_rows = {
        row for row in total_rows
        if all(cell.font.color is not None and cell.font.color.type == "rgb" and cell.font.color.rgb == "FF000000"
               for cell in formula_cells[row])
    }

    for cell in list(ws._cells.values()):
        if cell.row >= 2:
            cell.value = None
            cell.comment = None
            cell.hyperlink = None

    for row in removable_rows:
        for key in [key for key in ws._cells if key[0] == row]:
            del ws._cells[key]


def reset_sheet(ws):
    tail_dimension = None
    for dimension in ws.column_dimensions.values():
        if dimension.max == 16384 and (
            dimension.has_style or dimension.width != ws.sheet_format.defaultColWidth
        ):
            tail_dimension = copy(dimension)
            break
    ws._cells.clear()
    ws.row_dimensions.clear()
    ws.column_dimensions.clear()
    if tail_dimension is not None:
        tail_dimension.index = "A"
        tail_dimension.min = 1
        tail_dimension.max = 16384
        ws.column_dimensions["A"] = tail_dimension
    ws.merged_cells.ranges.clear()
    ws.auto_filter.ref = None
    ws._tables.clear()
    ws.conditional_formatting._cf_rules.clear()
    ws.data_validations.dataValidation.clear()
    ws._images.clear()
    ws._charts.clear()


def filter_unmatched(ws):
    keep = [1]
    for row in range(2, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        cancelled = any(
            cancel_type in str(value)
            for value in values if value is not None
            for cancel_type in CANCEL_TYPES
        )
        if not cancelled and any(value not in (None, "") for value in values):
            keep.append(row)

    if len(keep) == 1:
        reset_sheet(ws)
        return

    dimensions = {}
    for new_row, old_row in enumerate(keep, 1):
        if old_row in ws.row_dimensions:
            dimension = copy(ws.row_dimensions[old_row])
            dimension.index = new_row
            dimensions[new_row] = dimension

    keep_set = set(keep)
    for row in range(ws.max_row, 0, -1):
        if row not in keep_set:
            ws.delete_rows(row)

    ws.row_dimensions.clear()
    for row, dimension in dimensions.items():
        ws.row_dimensions[row] = dimension


def patch_saved_file(path, workbook, cached_workbook):
    empty_strings = {}
    formula_values = {}
    plain_styles = {}
    for index, ws in enumerate(workbook.worksheets, 1):
        part = f"xl/worksheets/sheet{index}.xml"
        empty_strings[part] = {cell.coordinate for cell in ws._cells.values() if cell.value == ""}
        plain_styles[part] = {
            cell.coordinate: cell.style_id
            for cell in ws._cells.values() if cell._style is not None and not cell.has_style
        }
        memo = {}

        def calculate(coordinate):
            if coordinate in memo:
                return memo[coordinate]
            cell = ws[coordinate]
            if cell.data_type != "f":
                return cell.value if isinstance(cell.value, (int, float)) else 0
            expression = FORMULA_REF.sub(
                lambda match: repr(calculate(f"{match.group(1)}{match.group(2)}")),
                cell.value[1:],
            )
            if not re.fullmatch(r"[0-9eE+\-*/(). ]+", expression):
                raise ValueError(expression)
            memo[coordinate] = eval(expression, {"__builtins__": {}}, {})
            return memo[coordinate]

        formula_values[part] = {}
        for cell in list(ws._cells.values()):
            if cell.data_type == "f":
                try:
                    formula_values[part][cell.coordinate] = calculate(cell.coordinate)
                except (TypeError, ValueError, SyntaxError, ZeroDivisionError):
                    formula_values[part][cell.coordinate] = cached_workbook[ws.title][cell.coordinate].value or 0

    def patch_xml(data, empty_refs, cached_values, explicit_styles):
        def replace(match):
            ref_match = CELL_REF.search(match.group("attrs"))
            if not ref_match:
                return match.group(0)
            ref = ref_match.group(1).decode("ascii")
            cell = match.group(0)
            if ref in explicit_styles and not re.search(rb'\bs=', match.group("attrs")):
                cell = cell.replace(b"<c", f'<c s="{explicit_styles[ref]}"'.encode("ascii"), 1)
            if ref in empty_refs and b't="inlineStr"' in match.group("attrs") and not match.group("body"):
                return cell.replace(b"></c>", b"><is><t></t></is></c>")
            if ref in cached_values and b"<f" in match.group("body"):
                value = cached_values[ref]
                if isinstance(value, bool):
                    value = int(value)
                encoded = (format(value, ".15g") if isinstance(value, float) else str(value)).encode("utf-8")
                return cell.replace(b"<v></v>", b"<v>" + encoded + b"</v>", 1)
            return cell

        return CELL_XML.sub(replace, data)

    temp = tempfile.NamedTemporaryFile(delete=False, dir=path.parent, suffix=path.suffix)
    temp.close()
    try:
        with ZipFile(path) as source, ZipFile(temp.name, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in empty_strings:
                    data = patch_xml(data, empty_strings[info.filename], formula_values[info.filename], plain_styles[info.filename])
                target.writestr(info, data)
        os.replace(temp.name, path)
    finally:
        if os.path.exists(temp.name):
            os.unlink(temp.name)


def process_file(source, destination):
    keep_vba = source.suffix.lower() == ".xlsm"
    workbook = load_workbook(source, keep_vba=keep_vba)
    cached_workbook = load_workbook(source, data_only=True, read_only=True, keep_vba=keep_vba)
    profit = get_sheet(workbook, "profit")
    cleared_titles = set()
    for key in ("sales", *EMPTY_SHEETS):
        ws = get_sheet(workbook, key, required=False)
        if ws is not None:
            cleared_titles.add(ws.title)
    preserve_linked_profit_values(profit, cached_workbook[profit.title], cleared_titles)
    add_next_month(profit)
    clear_sales(get_sheet(workbook, "sales"))
    for key in EMPTY_SHEETS:
        ws = get_sheet(workbook, key, required=False)
        if ws is not None:
            reset_sheet(ws)
    filter_unmatched(get_sheet(workbook, "unmatched"))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    patch_saved_file(destination, workbook, cached_workbook)
    cached_workbook.close()
    workbook.close()


def main():
    parser = argparse.ArgumentParser(description="批量格式化店铺月度 Excel 工作簿")
    parser.add_argument("--input", type=Path, default=Path("deal"))
    parser.add_argument("--output", type=Path, default=Path("output"))
    args = parser.parse_args()

    files = sorted(
        path for path in args.input.iterdir()
        if path.is_file() and not path.name.startswith("~$") and path.suffix.lower() in {".xlsx", ".xlsm"}
    )
    if not files:
        raise SystemExit(f"未在 {args.input} 中找到 Excel 文件")

    failures = []
    for source in files:
        destination = args.output / source.name
        try:
            process_file(source, destination)
            print(f"{source.name} -> {destination}")
        except Exception as exc:
            failures.append((source.name, str(exc)))
            print(f"失败：{source.name}：{exc}")
    if failures:
        raise SystemExit(f"{len(failures)} 个文件处理失败：{failures}")


if __name__ == "__main__":
    main()
