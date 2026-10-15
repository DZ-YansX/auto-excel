# -*- coding: utf-8 -*-
"""Shopee 销售明细、利润表和未匹配自动填写。"""

import glob
import math
import os
import posixpath
import re
import shutil
import stat
import sys
import tempfile
import warnings
import xml.etree.ElementTree as ET
from collections import Counter
from copy import copy
from datetime import datetime
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter, coordinate_to_tuple
from openpyxl.worksheet.views import Selection

warnings.filterwarnings('ignore')


# ================= 配置区域 =================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TARGET_FOLDER_PATH = r"ShopeeDeal"
OUTPUT_FOLDER_PATH = r"output_wanshan"

CURRENCY_RATES = {
    '菲律宾': 0.1065,
    '印尼': 0.0003754,
    '马来': 1.629,
    '泰国': 0.201,
    'PHP': 0.1065,
    'IDR': 0.0003754,
    'MYR': 1.629,
    'THB': 0.201,
}
UNIT_FEE = 3.5
DISPATCH_FEES = {'菲律宾': 3.5, 'PHP': 3.5, '印尼': 3.0, 'IDR': 3.0,
                 '马来': 3.5, 'MYR': 3.5, '泰国': 3.5, 'THB': 3.5}
USD_RATE = 7
ADS_TAX_FACTORS = {
    '菲律宾': 1.0, 'PHP': 1.0,
    '印尼': 1.11, 'IDR': 1.11,
    '马来': 1.08, 'MYR': 1.08,
    '泰国': 1.07, 'THB': 1.07,
}

COUNTRY_MAP = {
    'Philippines': '菲律宾', 'PH': '菲律宾',
    'Malaysia': '马来', 'MY': '马来',
    'Thailand': '泰国', 'TH': '泰国',
    'Indonesia': '印尼', 'ID': '印尼',
}
COUNTRY_CURRENCY = {
    'Philippines': 'PHP', 'PH': 'PHP', '菲律宾': 'PHP',
    'Malaysia': 'MYR', 'MY': 'MYR', '马来': 'MYR',
    'Thailand': 'THB', 'TH': 'THB', '泰国': 'THB',
    'Indonesia': 'IDR', 'ID': 'IDR', '印尼': 'IDR',
}
SKU_ALIAS = {
    'YY-054': 'YY-KP', 'YY-056': 'YY-KP',
    'YY-057': 'YY-KP', 'YY-059': 'YY-KP',
}
KP_FALLBACK_PREFIXES = ('YY', 'XY')
YELLOW_FILL = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')
OOXML_MAIN_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
OOXML_DOC_REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
OOXML_PKG_REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
OOXML_CONTENT_TYPES_NS = 'http://schemas.openxmlformats.org/package/2006/content-types'

INCOME_HEADER_ROW = 6
ORDER_HEADER_ROW = 1
ADS_HEADER_ROW = 8

SHEET_NAMES = {
    'pnl': ['利润表'],
    'sales': ['销售明细', '销售明细表'],
    'income': ['账单明细', '收入明细表', '收入明细'],
    'order': ['订单明细', '订单明细原表', '订单明细表'],
    'ads': ['广告', '广告费'],
    'brush': ['补单明细 ', '补单明细', '补单'],
    'unmatched': ['未匹配', '未匹配订单'],
    'dmj': ['大卖家'],
}


def safe_float(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return 0.0
        return float(value)
    text = str(value).strip().lower()
    if text in ('', '-', '--', 'none', 'nan'):
        return 0.0
    for token in (',', ' ', '¥', '฿', '₱', 'php', 'rm'):
        text = text.replace(token, '')
    try:
        result = float(text)
        return 0.0 if math.isnan(result) else result
    except ValueError:
        return 0.0


def num_or_int(value):
    value = float(value)
    return int(value) if value.is_integer() else value


def parse_date(value):
    if value in (None, '', '-'):
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(text[:10], fmt)
        except ValueError:
            continue
    return None


def get_sheet(wb, candidates):
    for candidate in candidates:
        for actual in wb.sheetnames:
            if actual.strip() == candidate.strip():
                return wb[actual]
    return None


def find_col(ws, header_row, names, default=None, contains=False):
    targets = [str(name).strip().lower() for name in names]
    for cell in ws[header_row]:
        if cell.value is None:
            continue
        text = str(cell.value).strip().lower()
        if text in targets:
            return cell.column
    if contains:
        for cell in ws[header_row]:
            if cell.value is None:
                continue
            text = str(cell.value).strip().lower()
            if any(target in text for target in targets):
                return cell.column
    return default


def detect_country(file_name):
    for country in CURRENCY_RATES:
        if file_name.startswith(country):
            return country
    return None


def font_is_red(cell):
    color = cell.font.color
    if color is None or color.type != 'rgb':
        return False
    return str(color.rgb).upper().endswith('FF0000')


def is_shipping_status(value):
    text = str(value or '').strip().lower()
    return text in {
        'shipping', 'การจัดส่ง', 'จัดส่งสำเร็จแล้ว',
        '运输中', '運輸中', '运送中', '運送中', '配送中', '出货中', '出貨中',
        'sedang dikirim', 'sedang dihantar',
    }


def clean_label(value):
    if value is None:
        return ''
    text = str(value).strip()
    text = re.sub(r'^[一二三四五六七八九十]+[、.]', '', text)
    text = re.sub(r'^\d+[.、]', '', text)
    text = text.replace('其中:', '').replace('其中：', '')
    return text.rstrip(':：').strip()


def locate_income_cols(ws):
    return {
        'id': find_col(ws, INCOME_HEADER_ROW,
                       ['Order ID', '订单编号', '訂單編號', '订单号', '訂單號'], 2),
        'date': find_col(ws, INCOME_HEADER_ROW,
                         ['Order Creation Date', '订单成立时间', '訂單成立時間',
                          '下单时间', '下單時間'], 5, True),
        'price': find_col(ws, INCOME_HEADER_ROW,
                          ['Original Product Price', '商品原价', '商品原價'], 12, True),
        'promo': find_col(ws, INCOME_HEADER_ROW,
                          ['Seller Product Promotion', 'Your Seller Product Promotion',
                           '商品折扣', '卖家商品折扣', '賣家商品折扣'], 13, True),
        'refund': find_col(ws, INCOME_HEADER_ROW,
                           ['Refund Amount', '退款金额', '退款金額'], 14, True),
        'voucher': find_col(ws, INCOME_HEADER_ROW,
                            ['Voucher Sponsored by Seller', '卖家赞助的优惠券',
                             '賣家贊助的優惠券'], 16),
        'fee_start': 20,  # T
        'fee_end': 31,    # AE
    }


def locate_order_cols(ws):
    return {
        'id': find_col(ws, ORDER_HEADER_ROW,
                       ['Order ID', 'หมายเลขคำสั่งซื้อ', '订单号', '訂單號'], 1),
        'status': find_col(ws, ORDER_HEADER_ROW,
                           ['Order Status', 'สถานะการสั่งซื้อ', '订单状态', '訂單狀態'], 2),
        'ship': find_col(ws, ORDER_HEADER_ROW,
                         ['Ship Time', 'เวลาส่งสินค้า', '发货时间', '發貨時間'], 10),
        'date': find_col(ws, ORDER_HEADER_ROW,
                         ['Order Creation Date', 'วันที่ทำการสั่งซื้อ', 'วันที่สั่งซื้อ',
                          '下单时间', '下單時間'], 11),
        'sku': find_col(ws, ORDER_HEADER_ROW,
                        ['SKU Reference No.', 'SKU参考编号', 'SKU參考編號'], 15, True),
        'qty': find_col(ws, ORDER_HEADER_ROW,
                        ['Quantity', 'จำนวน', '数量', '數量'], 19),
    }


def load_order_cost_map_from_dmj(ws):
    oid_col = find_col(ws, 1, ['订单号'], None)
    total_col = find_col(ws, 1, ['总成本价'], None)
    if oid_col is None or total_col is None:
        raise ValueError("大卖家 Sheet 缺少‘订单号’或‘总成本价’列")

    cost_map = {}
    for row in ws.iter_rows(min_row=2):
        oid_value = row[oid_col - 1].value
        if oid_value is None or str(oid_value).strip() == '':
            continue
        oid = str(oid_value).strip()
        raw_cost = row[total_col - 1].value
        if raw_cost in (None, '', '-', '--'):
            continue
        cost = safe_float(raw_cost)
        if oid not in cost_map or cost > cost_map[oid]:
            cost_map[oid] = cost
    return cost_map


def load_sku_cost_map_from_dmj(ws):
    """仅在订单号缺失时兜底：从大卖家商品SKU/商品成本建立SKU成本映射。"""
    sku_col = find_col(ws, 1, ['商品SKU'], None)
    product_cost_col = find_col(ws, 1, ['商品成本'], None)
    if sku_col is None or product_cost_col is None:
        return {}
    result = {}
    for row in ws.iter_rows(min_row=2):
        skus = split_skus(row[sku_col - 1].value)
        raw_cost = row[product_cost_col - 1].value
        costs = [safe_float(value) for value in re.findall(
            r'[-+]?\d+(?:\.\d+)?', str(raw_cost or '').replace(',', ''))]
        if not skus:
            continue
        if len(costs) == 1 and len(skus) == 1:
            pairs = zip(skus, costs)
        elif len(costs) == len(skus):
            pairs = zip(skus, costs)
        else:
            continue
        for sku, cost in pairs:
            if cost and (sku not in result or cost > result[sku]):
                result[sku] = cost
    return result


def read_income(ws):
    cols = locate_income_cols(ws)
    aggregate = {}
    sequence = []
    for row in ws.iter_rows(min_row=INCOME_HEADER_ROW + 1):
        oid = row[cols['id'] - 1].value
        if oid is None or str(oid).strip() == '':
            continue
        oid = str(oid).strip()
        # 只读取第7行起的订单数据；第5行 subtotal 和第6行表头不参与。
        fee_total = sum(safe_float(row[c - 1].value)
                        for c in range(cols['fee_start'], cols['fee_end'] + 1))
        if oid not in aggregate:
            aggregate[oid] = {
                'date': row[cols['date'] - 1].value,
                'price': 0.0,
                'promo': 0.0,
                'refund': 0.0,
                'voucher': 0.0,
                'fees': 0.0,
            }
            sequence.append(oid)
        item = aggregate[oid]
        item['price'] += safe_float(row[cols['price'] - 1].value)
        item['promo'] += safe_float(row[cols['promo'] - 1].value)
        item['refund'] += safe_float(row[cols['refund'] - 1].value)
        item['voucher'] += safe_float(row[cols['voucher'] - 1].value)
        item['fees'] += fee_total
    return aggregate, sequence


def read_orders(ws, order_cost_map, sku_cost_map=None):
    cols = locate_order_cols(ws)
    aggregate = {}
    sequence = []
    rows = []
    missing_cost_orders = set()
    n_cols = ws.max_column

    for source_row in ws.iter_rows(min_row=ORDER_HEADER_ROW + 1, max_col=n_cols):
        oid_value = source_row[cols['id'] - 1].value
        if oid_value is None or str(oid_value).strip() == '':
            continue
        oid = str(oid_value).strip()
        sku_value = source_row[cols['sku'] - 1].value
        order_cost = order_cost_map.get(oid, 0.0)
        if oid not in order_cost_map:
            missing_cost_orders.add(oid)
        status = str(source_row[cols['status'] - 1].value or '').strip()
        ship = source_row[cols['ship'] - 1].value
        is_red = font_is_red(source_row[cols['id'] - 1])

        if oid not in aggregate:
            aggregate[oid] = {
                'skus': [],
                'qty': 0.0,
                'cost': order_cost,
                'ship': ship,
                'status': status,
                'date': source_row[cols['date'] - 1].value,
                'fallback_cost': False,
            }
            sequence.append(oid)
        item = aggregate[oid]
        sku_text = str(sku_value).strip() if sku_value is not None else ''
        item['skus'].append(sku_text)
        row_qty = safe_float(source_row[cols['qty'] - 1].value)
        item['qty'] += row_qty
        if oid not in order_cost_map and sku_cost_map and sku_text in sku_cost_map:
            item['cost'] += sku_cost_map[sku_text] * row_qty
            item['fallback_cost'] = True
        if item['ship'] in (None, '', '-') and ship not in (None, '', '-'):
            item['ship'] = ship
        if not item['status'] and status:
            item['status'] = status

        rows.append({
            'oid': oid,
            'values': [cell.value for cell in source_row],
            'styles': [copy(cell._style) for cell in source_row],
            'height': ws.row_dimensions[source_row[0].row].height,
            'order_cost': order_cost,
            'status': status,
            'ship': ship,
            'red': is_red,
        })
    missing_cost_orders = {
        oid for oid in missing_cost_orders if not aggregate.get(oid, {}).get('cost')
    }
    for row in rows:
        if row['oid'] in aggregate:
            row['order_cost'] = aggregate[row['oid']]['cost']
    return aggregate, sequence, rows, missing_cost_orders, n_cols


def read_brush(ws):
    if ws is None or ws.max_row < 2:
        return [], 0.0, 0.0, 0.0
    id_col = find_col(ws, 1, ['线上单号'], 6 if ws.max_column >= 6 else None)
    cost_col = find_col(ws, 1, ['垫付成本'], 11 if ws.max_column >= 11 else None)
    sub_col = find_col(ws, 1, ['垫付小计'], 15 if ws.max_column >= 15 else None)
    comm_col = find_col(ws, 1, ['佣金'], 16 if ws.max_column >= 16 else None)
    if id_col is None:
        return [], 0.0, 0.0, 0.0

    ids = []
    cost_sum = sub_sum = comm_sum = 0.0
    for row_number in range(2, ws.max_row + 1):
        oid = ws.cell(row_number, id_col).value
        if oid is None or str(oid).strip() == '':
            break
        ids.append(str(oid).strip())
        if cost_col:
            cost_sum += safe_float(ws.cell(row_number, cost_col).value)
        if sub_col:
            sub_sum += safe_float(ws.cell(row_number, sub_col).value)
        if comm_col:
            comm_sum += safe_float(ws.cell(row_number, comm_col).value)
    return ids, cost_sum, sub_sum, comm_sum


def read_ads(ws):
    """汇总广告表D列明细后取绝对值，表头和合计公式不重复计入。"""
    if ws is None:
        return 0.0

    header_row = None
    amount_col = None
    aliases = ('amount', '总额', '總額', 'expense', '花费', '費用', 'ค่าใช้จ่าย')
    for row in range(1, min(ws.max_row, 20) + 1):
        if str(ws.cell(row, 4).value or '').strip().lower() in aliases:
            header_row, amount_col = row, 4
        if amount_col:
            break
    if amount_col is None:
        if not any(cell.value not in (None, '') for row in ws for cell in row):
            return 0.0
        raise ValueError('广告表D列找不到Amount/总额表头，请检查广告导出格式')

    total = sum(
        safe_float(ws.cell(row, amount_col).value)
        for row in range(header_row + 1, ws.max_row + 1)
        if isinstance(ws.cell(row, amount_col).value, (int, float))
        and not isinstance(ws.cell(row, amount_col).value, bool)
    )
    return abs(total)


def all_thin_border():
    side = Side(style='thin', color='FF000000')
    return Border(left=side, right=side, top=side, bottom=side)


def write_sales_sheet(wb, ws, sales_ids, order_agg, income_agg, rate,
                      missing_cost_orders, cross_month_ids=frozenset()):
    if ws is None:
        ws = wb.create_sheet('销售明细')
        ws.append(['日期', '订单号', '商品SKU', '销售数量', '销售单价\n(本币)',
                   '优惠金额\n(本币)', '优惠劵金额\n(本币)', '退款金额\n(本币)',
                   '销售金额\n(本币)', '销售金额\n(¥)', '商品成本总额\n(¥)', '毛利\n(¥)'])

    template_styles = [copy(ws.cell(2, c)._style) for c in range(1, 13)]
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)

    border = all_thin_border()
    month_counter = Counter()
    result = {'sum_j': 0.0, 'sum_h': 0.0, 'sum_k': 0.0,
              'j_by_id': {}, 'fee_sum': 0.0, 'months': month_counter}

    for index, oid in enumerate(sales_ids, start=2):
        order = order_agg[oid]
        income = income_agg[oid]
        price = income['price']
        promo = income['promo']
        voucher = income['voucher']
        refund = income['refund']
        local_sales = price + promo + voucher + refund
        cny_sales = local_sales * rate
        cross_month = oid in cross_month_ids
        cost = 0.0 if cross_month else order['cost']

        date_value = income['date']
        parsed = parse_date(date_value)
        if parsed:
            month_counter[parsed.month] += 1

        values = [
            date_value,
            oid,
            ','.join(sku for sku in order['skus'] if sku),
            num_or_int(order['qty']),
            num_or_int(price),
            num_or_int(promo),
            num_or_int(voucher),
            num_or_int(refund),
            f'=E{index}+F{index}+G{index}+H{index}',
            f'=I{index}*{rate}',
            num_or_int(cost),
            f'=J{index}-K{index}',
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(index, column)
            cell._style = copy(template_styles[column - 1])
            cell.value = value
            if oid in missing_cost_orders and not cross_month:
                red_font = copy(cell.font)
                red_font.color = 'FFFF0000'
                cell.font = red_font
            elif order.get('fallback_cost'):
                fallback_font = copy(cell.font)
                fallback_font.color = None
                cell.font = fallback_font
        ws.row_dimensions[index].height = 15

        result['sum_j'] += cny_sales
        result['sum_h'] += refund
        result['sum_k'] += cost
        result['j_by_id'][oid] = cny_sales
        result['fee_sum'] += income['fees']

    last_row = len(sales_ids) + 1
    total_row = last_row + 1
    for column in range(4, 13):
        letter = get_column_letter(column)
        cell = ws.cell(total_row, column)
        cell.value = f'=SUM({letter}2:{letter}{last_row})'
        cell.font = Font(name='宋体', size=11, bold=True, charset=134,
                         color='FF000000')
        cell.alignment = Alignment(vertical='center')
    return ws, result


def write_unmatched_sheet(wb, ws, ws_order, order_rows, income_ids, n_order_cols,
                          cross_month_ids=frozenset()):
    if ws is None:
        ws = wb.create_sheet('未匹配')
    if ws.max_row >= 1:
        ws.delete_rows(1, ws.max_row)

    for key, dimension in ws_order.column_dimensions.items():
        if dimension.width is not None:
            ws.column_dimensions[key].width = dimension.width

    for column in range(1, n_order_cols + 1):
        source = ws_order.cell(ORDER_HEADER_ROW, column)
        target = ws.cell(1, column)
        target.value = source.value
        target._style = copy(source._style)
    label_col = n_order_cols + 1
    cost_col = n_order_cols + 2
    if label_col <= ws_order.max_column:
        source = ws_order.cell(ORDER_HEADER_ROW, label_col)
        ws.cell(1, label_col).value = source.value
        ws.cell(1, label_col)._style = copy(source._style)
    else:
        ws.cell(1, label_col).value = None
    ws.cell(1, cost_col).value = None
    ws.cell(1, cost_col).font = Font(name='宋体', size=11, charset=134)
    ws.cell(1, cost_col).alignment = Alignment(horizontal='center', vertical='center',
                                                wrap_text=True)
    ws.column_dimensions['B'].width = 10.875
    ws.column_dimensions['C'].width = 13
    ws.column_dimensions[get_column_letter(cost_col)].width = 6.75
    ws.row_dimensions[1].height = ws_order.row_dimensions[1].height or 28.5

    filtered_rows = [row for row in order_rows
                     if row['oid'] not in income_ids and not row['red']]
    excluded_red = [row for row in order_rows
                    if row['oid'] not in income_ids and row['red']]

    meta = {}
    sequence = []
    for row in filtered_rows:
        oid = row['oid']
        if oid not in meta:
            meta[oid] = {
                'status': row['status'],
                'ship': row['ship'],
                'cost': row['order_cost'],
            }
            sequence.append(oid)
        item = meta[oid]
        if item['ship'] in (None, '', '-') and row['ship'] not in (None, '', '-'):
            item['ship'] = row['ship']
        if is_shipping_status(row['status']):
            item['status'] = 'Shipping'
        elif not item['status'] and row['status']:
            item['status'] = row['status']

    cancel_unshipped = []
    cancel_shipped = []
    unpaid = []
    for oid in sequence:
        item = meta[oid]
        if is_shipping_status(item['status']):
            unpaid.append(oid)
        elif item['ship'] not in (None, '', '-'):
            cancel_shipped.append(oid)
        else:
            cancel_unshipped.append(oid)

    shipped_ids = set(cancel_shipped + unpaid)
    cross_shipped_ids = shipped_ids & set(cross_month_ids)
    for oid in cross_shipped_ids:
        meta[oid]['cost'] = 0.0
    cost_written = set()
    for output_row, row in enumerate(filtered_rows, start=2):
        for column in range(1, n_order_cols + 1):
            cell = ws.cell(output_row, column)
            cell.value = row['values'][column - 1]
            cell._style = copy(row['styles'][column - 1])
        label_cell = ws.cell(output_row, label_col)
        label_style_col = label_col if label_col <= len(row['styles']) else n_order_cols
        label_cell._style = copy(row['styles'][label_style_col - 1])
        label_cell.font = Font(name='宋体', size=10, charset=134)
        if row['oid'] in unpaid:
            label_cell.value = '发货未到账'
        elif row['oid'] in cancel_shipped:
            label_cell.value = '取消已发货'
        else:
            label_cell.value = '取消未发货'
        label_cell.fill = (
            PatternFill(fill_type=None)
            if str(row['status']).strip() == 'ที่ต้องจัดส่ง'
            else PatternFill(fill_type='solid', fgColor='FFF8F8FF',
                             bgColor=Color(indexed=64)))

        cost_cell = ws.cell(output_row, cost_col)
        cost_cell.font = Font(name='宋体', size=11, charset=134)
        cost_cell.alignment = Alignment(horizontal='center', vertical='center')
        if row['oid'] in shipped_ids and row['oid'] not in cost_written:
            cost_cell.value = num_or_int(meta[row['oid']]['cost'])
            cost_written.add(row['oid'])
        else:
            cost_cell.value = None
        height = row['height']
        ws.row_dimensions[output_row].height = height if height is not None else 15

    shipped_cost = sum(meta[oid]['cost'] for oid in shipped_ids)
    return {
        'order_count': len(sequence),
        'row_count': len(filtered_rows),
        'cancel_unshipped': len(cancel_unshipped),
        'cancel_shipped': len(cancel_shipped),
        'unpaid': len(unpaid),
        'shipped_cost': shipped_cost,
        'cross_shipped': len(cross_shipped_ids),
        'excluded_red_rows': len(excluded_red),
        'excluded_red_ids': sorted(set(row['oid'] for row in excluded_red)),
    }


def set_pnl_title(ws):
    ws['A1'] = CellRichText([
        TextBlock(InlineFont(rFont='宋体', sz=12, b=True, charset=134,
                             color=Color(rgb='FF000000')), '利润表'),
        TextBlock(InlineFont(rFont='宋体', sz=8, b=True, charset=134,
                             color=Color(rgb='FFFF0000')), '\n单位：CNY'),
    ])
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 40)


def copy_month_column(ws, source_col, target_col, first_row, last_row):
    source_letter = get_column_letter(source_col)
    target_letter = get_column_letter(target_col)
    for row in range(first_row, last_row + 1):
        source = ws.cell(row, source_col)
        target = ws.cell(row, target_col)
        target._style = copy(source._style)
        if isinstance(source.value, str) and source.value.startswith('='):
            target.value = re.sub(rf'\b{source_letter}(\d+)', rf'{target_letter}\1', source.value)


def write_pnl_sheet(ws, month, values, note_counts):
    set_pnl_title(ws)
    target_header = f'{month}月'
    month_cols = {}
    first_empty = None
    for column in range(2, ws.max_column + 2):
        value = ws.cell(2, column).value
        text = str(value).strip() if value is not None else ''
        match = re.fullmatch(r'(\d{1,2})月', text)
        if match:
            month_cols[int(match.group(1))] = column
        elif first_empty is None:
            first_empty = column

    if month in month_cols:
        target_col = month_cols[month]
    else:
        target_col = first_empty or (ws.max_column + 1)
        if month_cols:
            source_col = max(month_cols.values())
            copy_month_column(ws, source_col, target_col, 2, ws.max_row)
        ws.cell(2, target_col).value = target_header

    row_map = {}
    for row in range(3, ws.max_row + 1):
        label = clean_label(ws.cell(row, 1).value)
        if label:
            row_map[label.lower()] = row

    def row_for(label):
        return row_map.get(label.lower())

    for label, value in values.items():
        row = row_for(label)
        if row:
            ws.cell(row, target_col).value = (
                value if isinstance(value, str) and value.startswith('=')
                else num_or_int(round(float(value), 8))
            )

    letter = get_column_letter(target_col)

    def set_formula(label, plus=(), minus=()):
        target_row = row_for(label)
        if not target_row:
            return
        terms = []
        for item in plus:
            source_row = row_for(item)
            if source_row:
                terms.append(f'+{letter}{source_row}')
        for item in minus:
            source_row = row_for(item)
            if source_row:
                terms.append(f'-{letter}{source_row}')
        if terms:
            formula = ''.join(terms).lstrip('+')
            ws.cell(target_row, target_col).value = f'={formula}'

    set_formula('实际主营收入', plus=['本期销售'], minus=['补单金额', '销售退款'])
    set_formula('实际主营成本', plus=['实际产品成本', '发单费'])
    set_formula('推广费用', plus=['广告费', '补单费'])
    settlement_label = None
    for candidate in ('Shopee结算', 'TK结算', '平台结算'):
        if row_for(candidate):
            settlement_label = candidate
            break
    if settlement_label is None:
        settlement_label = next((label for label in row_map if label.endswith('结算')), None)
    if settlement_label:
        set_formula(settlement_label, plus=['平台费用'])
        set_formula('营业费用', plus=['推广费用', settlement_label])
    else:
        set_formula('营业费用', plus=['推广费用'])
    profit_label = '本期利润' if row_for('本期利润') else '毛利润'
    set_formula(profit_label, plus=['实际主营收入'],
                minus=['实际主营成本', '营业费用'])

    profit_row = row_for(profit_label)
    note_row = None
    for row in range(3, ws.max_row + 1):
        if '未匹配' in str(ws.cell(row, 1).value or ''):
            note_row = row
            break
    if note_row is None:
        note_row = (profit_row or ws.max_row) + 1

    merge_range = f'A{note_row}:D{note_row}'
    containing = [str(item) for item in ws.merged_cells.ranges
                  if item.min_row <= note_row <= item.max_row and item.min_col <= 1 <= item.max_col]
    if not containing:
        ws.merge_cells(merge_range)

    note = (f'{month}月未匹配：取消未发货{note_counts[0]}，'
            f'取消已发货{note_counts[1]}，发货未到账{note_counts[2]}')
    cell = ws.cell(note_row, 1)
    cell.value = note
    cell.font = Font(name='宋体', size=10, bold=True, charset=134,
                     color='FFFF0000')
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[note_row].height = 24.95

    # 目标月份列宽由模板控制，避免改变既有展示比例。


def settlement_find_col(header_cells, keyword, default=None):
    """新版结算表按中英文别名动态定位列号（1-based）。"""
    kws = [keyword.lower()] if isinstance(keyword, str) else [k.lower() for k in keyword]
    for kw in kws:
        for cell in header_cells:
            if cell.value is not None and str(cell.value).strip().lower() == kw:
                return cell.column
    for kw in kws:
        for cell in header_cells:
            if cell.value is not None and kw in str(cell.value).lower():
                return cell.column
    return default


def sum_numeric_cell(x):
    """汇总单元格中的一个或多个数值。
    大卖家导出在同一订单包含多个商品时，成本列可能是 '6.58\n6.58\n'；
    这里按换行/分号拆分后求和，普通数字仍按原数值读取。"""
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return safe_float(x)
    parts = [p.strip() for p in re.split(r'[\r\n;；]+', str(x)) if p.strip()]
    if len(parts) <= 1:
        return safe_float(x)
    return sum(safe_float(p) for p in parts)


def fmt_num(v):
    """数字转字符串: 去掉多余的小数尾零 (用于公式与批注文本)"""
    v = round(float(v), 4)
    if v == int(v):
        return str(int(v))
    return f'{v:f}'.rstrip('0').rstrip('.')


def parse_month(s):
    """从日期(字符串或datetime)提取月份, 失败返回 None"""
    if s is None:
        return None
    if hasattr(s, 'month') and not isinstance(s, str):
        return s.month
    m = re.match(r'\s*(\d{4})[/-](\d{1,2})', str(s))
    if m:
        return int(m.group(2))
    m = re.match(r'\s*(\d{1,2})[/-]\d{1,2}[/-]\d{4}', str(s))  # MM/DD/YYYY
    if m:
        return int(m.group(1))
    return None


def strip_one_newline(v):
    """去掉字符串结尾的一个 \\n (大卖家导出值尾部是 \\r\\n, 正确表保留 \\r)"""
    if isinstance(v, str) and v.endswith('\n'):
        return v[:-1]
    return v


def normalize_multiline_sku(value, trailing_cr=False):
    """清理大卖家导出的尾部空行，并统一多SKU换行为Excel可见的CRLF。"""
    if value in (None, ''):
        return value
    parts = [part.strip() for part in re.split(r'[\r\n]+', str(value)) if part.strip()]
    if not parts:
        return None
    text = ('\r\n' if trailing_cr else '\n').join(parts)
    return text + ('\r' if trailing_cr else '')


def strip_export_tail(value):
    return value.rstrip('\r\n') if isinstance(value, str) else value


def normalize_export_quantity(value):
    """单个数量转数值；多商品数量保留换行显示但去掉末尾空行。"""
    if not isinstance(value, str):
        return value
    parts = [part.strip() for part in re.split(r'[\r\n]+', value) if part.strip()]
    if len(parts) == 1:
        number = safe_float(parts[0])
        return num_or_int(number)
    return '\n\n'.join(parts) if parts else None


def strip_minus_zero(formula):
    """去掉公式中手工失误留下的 '-0' 项 (如 =F4-0-F5 -> =F4-F5), 不影响 -0.5 等正常数字"""
    return re.sub(r'-0(?=[+\-*/)])|-0$', '', formula)


def _normalize_zip_part(base_dir, target):
    """把 OOXML relationship Target 转为 zip 内的规范相对路径。"""
    target = str(target).replace('\\', '/')
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(base_dir, target))


def _worksheet_part(zf, sheet_name):
    """从工作簿关系中取得指定 sheet 对应的 zip part 路径。"""
    workbook = ET.fromstring(zf.read('xl/workbook.xml'))
    rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
    rel_map = {rel.attrib['Id']: rel.attrib['Target'] for rel in rels}
    sheets = workbook.find(f'{{{OOXML_MAIN_NS}}}sheets')
    for sheet in sheets:
        if sheet.attrib.get('name') == sheet_name:
            rel_id = sheet.attrib[f'{{{OOXML_DOC_REL_NS}}}id']
            return _normalize_zip_part('xl', rel_map[rel_id])
    raise KeyError(f'找不到工作表: {sheet_name}')


def _rels_part(part_name):
    return posixpath.join(posixpath.dirname(part_name), '_rels',
                          posixpath.basename(part_name) + '.rels')


def restore_sheet_drawings(source_path, output_path, sheet_name):
    """在 openpyxl 保存后恢复指定工作表的原始嵌入图片。

    WPS 导出的广告截图属于 openpyxl 无法识别的 DrawingML 对象；直接保存会删除
    worksheet 的 drawing 关系、xl/drawings 与 xl/media。此函数只恢复绘图关系和
    原始媒体字节，保留 openpyxl 已写好的单元格、公式、样式及其他工作表内容。
    """
    replacements = {}
    source_infos = {}

    with ZipFile(source_path, 'r') as src, ZipFile(output_path, 'r') as out:
        src_sheet_part = _worksheet_part(src, sheet_name)
        out_sheet_part = _worksheet_part(out, sheet_name)
        src_sheet_rels_part = _rels_part(src_sheet_part)
        out_sheet_rels_part = _rels_part(out_sheet_part)
        if src_sheet_rels_part not in src.namelist():
            return 0

        src_rels = ET.fromstring(src.read(src_sheet_rels_part))
        drawing_rels = [
            rel for rel in src_rels
            if rel.attrib.get('Type', '').endswith('/drawing')
        ]
        if not drawing_rels:
            return 0

        if out_sheet_rels_part in out.namelist():
            out_rels = ET.fromstring(out.read(out_sheet_rels_part))
        else:
            out_rels = ET.Element(f'{{{OOXML_PKG_REL_NS}}}Relationships')
        used_ids = {rel.attrib.get('Id') for rel in out_rels}

        assigned = []
        copied_parts = set()
        for src_rel in drawing_rels:
            rel_id = src_rel.attrib['Id']
            if rel_id in used_ids:
                n = 1
                while f'rId{n}' in used_ids:
                    n += 1
                rel_id = f'rId{n}'
            used_ids.add(rel_id)
            attrs = dict(src_rel.attrib)
            attrs['Id'] = rel_id
            ET.SubElement(out_rels, f'{{{OOXML_PKG_REL_NS}}}Relationship', attrs)
            assigned.append(rel_id)

            drawing_part = _normalize_zip_part(posixpath.dirname(src_sheet_part),
                                               src_rel.attrib['Target'])
            copied_parts.add(drawing_part)
            drawing_rels_part = _rels_part(drawing_part)
            if drawing_rels_part in src.namelist():
                copied_parts.add(drawing_rels_part)
                drawing_part_rels = ET.fromstring(src.read(drawing_rels_part))
                for rel in drawing_part_rels:
                    target_part = _normalize_zip_part(posixpath.dirname(drawing_part),
                                                      rel.attrib['Target'])
                    if target_part in src.namelist():
                        copied_parts.add(target_part)

        # 在 openpyxl 输出的广告费 worksheet 上重新挂回 drawing 节点。
        out_sheet = ET.fromstring(out.read(out_sheet_part))
        for node in list(out_sheet.findall(f'{{{OOXML_MAIN_NS}}}drawing')):
            out_sheet.remove(node)
        for rel_id in assigned:
            ET.SubElement(out_sheet, f'{{{OOXML_MAIN_NS}}}drawing',
                          {f'{{{OOXML_DOC_REL_NS}}}id': rel_id})
        ET.register_namespace('', OOXML_MAIN_NS)
        ET.register_namespace('r', OOXML_DOC_REL_NS)
        replacements[out_sheet_part] = ET.tostring(
            out_sheet, encoding='utf-8', xml_declaration=True)
        ET.register_namespace('', OOXML_PKG_REL_NS)
        replacements[out_sheet_rels_part] = ET.tostring(
            out_rels, encoding='utf-8', xml_declaration=True)

        for part in copied_parts:
            replacements[part] = src.read(part)
            source_infos[part] = src.getinfo(part)

        # 补回图片扩展名与 drawing part 的 Content Types 声明。
        src_types = ET.fromstring(src.read('[Content_Types].xml'))
        out_types = ET.fromstring(out.read('[Content_Types].xml'))
        existing_defaults = {
            node.attrib.get('Extension') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default')
        }
        existing_overrides = {
            node.attrib.get('PartName') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override')
        }
        copied_extensions = {
            posixpath.splitext(part)[1].lstrip('.').lower()
            for part in copied_parts if part.startswith('xl/media/')
        }
        copied_part_names = {'/' + part for part in copied_parts
                             if part.startswith('xl/drawings/') and part.endswith('.xml')}
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default'):
            ext = node.attrib.get('Extension', '').lower()
            if ext in copied_extensions and ext not in existing_defaults:
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Default',
                              dict(node.attrib))
                existing_defaults.add(ext)
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override'):
            part_name = node.attrib.get('PartName')
            if part_name in copied_part_names and part_name not in existing_overrides:
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Override',
                              dict(node.attrib))
                existing_overrides.add(part_name)
        ET.register_namespace('', OOXML_CONTENT_TYPES_NS)
        replacements['[Content_Types].xml'] = ET.tostring(
            out_types, encoding='utf-8', xml_declaration=True)

        out_infos = {info.filename: info for info in out.infolist()}
        fd, temp_path = tempfile.mkstemp(
            prefix='.drawing_restore_', suffix='.xlsx',
            dir=os.path.dirname(os.path.abspath(output_path)))
        os.close(fd)
        with ZipFile(temp_path, 'w', ZIP_DEFLATED, allowZip64=True) as dst:
            for info in out.infolist():
                if info.filename not in replacements:
                    dst.writestr(info, out.read(info.filename))
            for part, data in replacements.items():
                info = out_infos.get(part) or source_infos.get(part)
                if info is not None:
                    dst.writestr(info, data)
                else:
                    dst.writestr(part, data)
    # Windows 不允许替换仍被 ZipFile 打开的文件，因此必须退出上面的 with 后再替换。
    try:
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
    return len(assigned)


def restore_sheet_comments(source_path, output_path, sheet_name):
    """恢复 Excel 原生批注关系和 VML，避免 openpyxl 重建的批注包触发修复。"""
    replacements = {}
    removed_parts = set()
    with ZipFile(source_path, 'r') as src, ZipFile(output_path, 'r') as out:
        src_sheet = _worksheet_part(src, sheet_name)
        out_sheet = _worksheet_part(out, sheet_name)
        src_rels_part = _rels_part(src_sheet)
        out_rels_part = _rels_part(out_sheet)
        if src_rels_part not in src.namelist():
            return 0
        src_rels = ET.fromstring(src.read(src_rels_part))
        src_comment_rels = [
            rel for rel in src_rels
            if rel.attrib.get('Type', '').endswith(('/comments', '/vmlDrawing'))
        ]
        if not src_comment_rels:
            return 0

        if out_rels_part in out.namelist():
            out_rels = ET.fromstring(out.read(out_rels_part))
        else:
            out_rels = ET.Element(f'{{{OOXML_PKG_REL_NS}}}Relationships')
        for rel in list(out_rels):
            if rel.attrib.get('Type', '').endswith(('/comments', '/vmlDrawing')):
                removed_parts.add(_normalize_zip_part(
                    posixpath.dirname(out_sheet), rel.attrib['Target']))
                out_rels.remove(rel)

        used_ids = {rel.attrib.get('Id') for rel in out_rels}
        copied_parts = set()
        vml_rel_id = None
        for src_rel in src_comment_rels:
            rel_id = src_rel.attrib['Id']
            if rel_id in used_ids:
                number = 1
                while f'rId{number}' in used_ids:
                    number += 1
                rel_id = f'rId{number}'
            used_ids.add(rel_id)
            attrs = dict(src_rel.attrib)
            attrs['Id'] = rel_id
            ET.SubElement(out_rels, f'{{{OOXML_PKG_REL_NS}}}Relationship', attrs)
            part = _normalize_zip_part(posixpath.dirname(src_sheet),
                                       src_rel.attrib['Target'])
            copied_parts.add(part)
            replacements[part] = src.read(part)
            if src_rel.attrib.get('Type', '').endswith('/vmlDrawing'):
                vml_rel_id = rel_id

        sheet_xml = ET.fromstring(out.read(out_sheet))
        for node in list(sheet_xml.findall(f'{{{OOXML_MAIN_NS}}}legacyDrawing')):
            sheet_xml.remove(node)
        if vml_rel_id:
            ET.SubElement(sheet_xml, f'{{{OOXML_MAIN_NS}}}legacyDrawing',
                          {f'{{{OOXML_DOC_REL_NS}}}id': vml_rel_id})
        ET.register_namespace('', OOXML_MAIN_NS)
        ET.register_namespace('r', OOXML_DOC_REL_NS)
        replacements[out_sheet] = ET.tostring(
            sheet_xml, encoding='utf-8', xml_declaration=True)
        ET.register_namespace('', OOXML_PKG_REL_NS)
        replacements[out_rels_part] = ET.tostring(
            out_rels, encoding='utf-8', xml_declaration=True)

        src_types = ET.fromstring(src.read('[Content_Types].xml'))
        out_types = ET.fromstring(out.read('[Content_Types].xml'))
        removed_names = {'/' + part for part in removed_parts}
        for node in list(out_types):
            if (node.tag.endswith('Override')
                    and node.attrib.get('PartName') in removed_names):
                out_types.remove(node)
        existing_overrides = {
            node.attrib.get('PartName') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override')
        }
        copied_names = {'/' + part for part in copied_parts}
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override'):
            part_name = node.attrib.get('PartName')
            if part_name in copied_names and part_name not in existing_overrides:
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Override',
                              dict(node.attrib))
                existing_overrides.add(part_name)
        existing_defaults = {
            node.attrib.get('Extension') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default')
        }
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default'):
            if (node.attrib.get('Extension') == 'vml'
                    and 'vml' not in existing_defaults):
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Default',
                              dict(node.attrib))
        ET.register_namespace('', OOXML_CONTENT_TYPES_NS)
        replacements['[Content_Types].xml'] = ET.tostring(
            out_types, encoding='utf-8', xml_declaration=True)

        out_infos = {info.filename: info for info in out.infolist()}
        src_infos = {info.filename: info for info in src.infolist()}
        fd, temp_path = tempfile.mkstemp(prefix='.comment_restore_', suffix='.xlsx',
                                         dir=os.path.dirname(os.path.abspath(output_path)))
        os.close(fd)
        with ZipFile(temp_path, 'w', ZIP_DEFLATED, allowZip64=True) as dst:
            for info in out.infolist():
                if info.filename not in replacements and info.filename not in removed_parts:
                    dst.writestr(info, out.read(info.filename))
            for part, data in replacements.items():
                dst.writestr(out_infos.get(part) or src_infos.get(part) or part, data)
    try:
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
    return len(src_comment_rels)


def preserve_rich_text_spaces(output_path):
    """保护纯空格富文本片段；Excel要求这类文本带xml:space属性。"""
    replacements = {}
    with ZipFile(output_path) as source:
        for name in source.namelist():
            if not (name.startswith('xl/worksheets/') and name.endswith('.xml')):
                continue
            data = source.read(name)
            updated = re.sub(rb'<t>([ \t\r\n]+)</t>',
                             rb'<t xml:space="preserve">\1</t>', data)
            if updated != data:
                replacements[name] = updated
        if not replacements:
            return
        fd, temp_path = tempfile.mkstemp(suffix='.xlsx', dir=os.path.dirname(output_path))
        os.close(fd)
        with ZipFile(temp_path, 'w', ZIP_DEFLATED) as dest:
            for info in source.infolist():
                dest.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    os.replace(temp_path, output_path)


def preserve_sheet_carriage_returns(output_path, sheet_name):
    """把工作表文本中的裸 CR 写成 XML 字符引用，避免再次打开时被规范化成 LF。"""
    with ZipFile(output_path, 'r') as src:
        sheet_part = _worksheet_part(src, sheet_name)
        if not sheet_part:
            return 0
        sheet_xml = src.read(sheet_part)
        count = len(re.findall(b'\r(?!\n)', sheet_xml))
        if not count:
            return 0
        replacement = re.sub(b'\r(?!\n)', b'&#13;', sheet_xml)
        fd, temp_path = tempfile.mkstemp(prefix='.cr_restore_', suffix='.xlsx',
                                         dir=os.path.dirname(output_path))
        os.close(fd)
        try:
            with ZipFile(temp_path, 'w', compression=ZIP_DEFLATED) as dst:
                for info in src.infolist():
                    data = replacement if info.filename == sheet_part else src.read(info.filename)
                    dst.writestr(info, data)
        except Exception:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            raise
    try:
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
    return count


def split_skus(cell):
    """把一个单元格里的多个 SKU 拆成列表 (换行/中英逗号/分号/顿号分隔)。
    注意: 不按空格或 '+'/'*' 拆分 —— SKU 名本身可能含空格(如 'La Roche-B3 30ml')、
    '+'(如 'SPF50+')、'*'(倍数或 'Coffee-*2' 这类含*的正式SKU名)。"""
    s = str(cell or '').strip()
    if not s:
        return []
    return [p.strip() for p in re.split(r'[\r\n,，;；、]+', s) if p.strip()]


def kp_fallback_sku(sku):
    """YY-*/XY-* 等前缀 SKU 的兜底目标 '<前缀>-KP'; 不在配置前缀内返回 None。"""
    m = re.match(r'^([A-Za-z]+)-', str(sku).strip())
    if m and m.group(1).upper() in KP_FALLBACK_PREFIXES:
        return m.group(1).upper() + '-KP'
    return None


def lookup_sku_unit(token, country, cost_map):
    """成本表中解析单个 SKU token 的单位成本(区分国家/仓库):
      1) 整个 token 精确匹配 (兼容 'Coffee-*2' 这类名字里带*的正式SKU)
      2) 'base*N' -> 基础SKU成本 * N
      3) SKU_ALIAS 显式别名
      4) YY-/XY- 前缀兜底到 '<前缀>-KP'
    都查不到返回 None。"""
    token = str(token).strip()
    v = cost_map.get((token, country))
    if v is not None:
        return v
    base, mult = token, 1
    if '*' in token:
        b, _, m = token.rpartition('*')
        if m.strip().isdigit():
            base, mult = b.strip(), int(m.strip())
            v = cost_map.get((base, country))
            if v is not None:
                return v * mult
    alias = SKU_ALIAS.get(base)
    if alias is not None:
        v = cost_map.get((alias, country))
        if v is not None:
            return v * mult
    kp = kp_fallback_sku(base)
    if kp is not None:
        v = cost_map.get((kp, country))
        if v is not None:
            return v * mult
    return None


def resolve_sku_cost(sku_cell, country, cost_map):
    """解析一个 SKU 单元格在指定国家的"单位成本合计"(单元格内多个不同SKU自动求和,
    每个 SKU 支持 'base*N' 倍数)。返回 (成本合计, 未解析token列表)。"""
    total = 0.0
    missing = []
    for token in split_skus(sku_cell):
        unit = lookup_sku_unit(token, country, cost_map)
        if unit is None:
            missing.append(token)
        else:
            total += unit
    return total, missing


def order_sku_qty(oid, det_lines):
    """大卖家表无该订单时, 从订单明细行拼出 SKU 与数量(与大卖家一行一单口径一致):
    多行则 SKU 换行合并、数量求和。返回 (SKU字符串或None, 数量或None)。"""
    lines = det_lines.get(oid, [])
    skus = [str(s).strip() for s, q in lines if str(s or '').strip()]
    qty = sum(safe_float(q) for s, q in lines)
    sku_str = '\n'.join(skus) if skus else None
    return sku_str, (num_or_int(qty) if lines else None)


def indonesia_order_sku_qty(oid, det_lines):
    """新版印尼销售明细按订单明细取SKU/数量，并复现28店对尾部换行的处理。"""
    lines = det_lines.get(oid, [])
    skus = []
    for sku, _ in lines:
        if sku in (None, ''):
            continue
        text = str(sku)
        # 28店未完成表的历史辅助行以两个 LF 结尾，而完成示例中同一批
        # 数据是 CRLF；两者在销售明细统一保留一个 CR。
        cleaned = text.rstrip('\r\n') + '\r' if text.endswith('\n\n') else strip_one_newline(text)
        if cleaned.strip():
            skus.append(cleaned)
    qty = sum(safe_float(q) for _, q in lines)
    return ('\n'.join(skus) if skus else None,
            num_or_int(qty) if lines else None)


def find_pending_pnl_month(ws):
    """查找利润表中已经建好、但“实际产品成本”仍为空的最新月份列。
    本月模板通常会预先建立空月份列；优先使用它可避免少量跨月订单让众数误写历史月份。"""
    product_cost_row = None
    for cell in ws['A']:
        label = str(cell.value or '').replace(' ', '').replace('\t', '')
        if '实际产品成本' in label:
            product_cost_row = cell.row
            break
    if product_cost_row is None:
        return None

    pending = []
    for row in ws.iter_rows(min_row=1, max_row=5):
        for cell in row:
            match = re.fullmatch(r'(\d{1,2})\s*月', str(cell.value or '').strip())
            if match and ws.cell(row=product_cost_row, column=cell.column).value in (None, ''):
                pending.append(int(match.group(1)))
    return max(pending) if pending else None


def parse_money_number(value):
    """解析补单表中的 Rp/¥ 金额文本；仅供新印尼账单逻辑使用。"""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return safe_float(value)
    text = str(value).replace(',', '').strip()
    match = re.search(r'[-+]?\d+(?:\.\d+)?', text)
    return safe_float(match.group(0)) if match else 0.0


def is_excel_number(value):
    """模拟 Excel SUM：只汇总真正的数值单元格，不把数字文本当作金额。"""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def copy_cell_all(source, target):
    """复制值与可见格式，用于新印尼格式的账单/订单明细附表。"""
    target.value = source.value
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


def is_new_indonesia_format(wb, sheet_map):
    """识别 Shopee 印尼新版账单；只在该格式下启用新逻辑，其他国家保持原分支。"""
    ws_bill = wb[sheet_map['bill']]
    ws_det = wb[sheet_map['detail']]
    bill_headers = {str(c.value or '').strip() for c in ws_bill[1]}
    detail_headers = {str(c.value or '').strip() for c in ws_det[1]}
    return ({'Order ID', 'Product Price', 'Payout Completed Date'} <= bill_headers
            and {'No. Pesanan', 'Waktu Pengiriman Diatur', 'Nomor Referensi SKU'} <= detail_headers)


def find_table_header_row(ws, order_id_aliases, max_scan_rows=10):
    """在表格顶部查找含订单号字段的真实表头行。"""
    aliases = {str(v).strip().lower() for v in order_id_aliases}
    max_col = min(ws.max_column, 100)
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = {str(ws.cell(r, c).value or '').strip().lower()
                  for c in range(1, max_col + 1)}
        if values & aliases:
            return r
    return None


def is_new_non_indonesia_settlement_format(wb, sheet_map):
    """识别菲律宾/马来西亚等新版 Shopee 结算账单，不影响印尼专用分支和旧账单。"""
    ws_bill = wb[sheet_map['bill']]
    ws_det = wb[sheet_map['detail']]
    header_row = find_table_header_row(ws_bill, ('Order ID', '订单编号', '订单ID'))
    if not header_row:
        return False
    bill_headers = {str(ws_bill.cell(header_row, c).value or '').strip()
                    for c in range(1, min(ws_bill.max_column, 100) + 1)}
    detail_headers = {str(ws_det.cell(1, c).value or '').strip()
                      for c in range(1, min(ws_det.max_column, 100) + 1)}
    has_price = bool({'Original Product Price', 'Product Price'} & bill_headers)
    has_payout = any('Total Released Amount' in h or '拨款金额' in h for h in bill_headers)
    return (has_price and has_payout
            and {'Order ID', 'SKU Reference No.'} <= detail_headers
            and 'No. Pesanan' not in detail_headers)


def read_indonesia_supplements(ws):
    """读取补单明细顶部业务表，返回 线上单号 -> 补单金额字段。"""
    if ws is None or ws.max_row < 2:
        return {}
    header = ws[1]
    c_oid = settlement_find_col(header, '线上单号')
    if not c_oid:
        return {}
    c_advance = settlement_find_col(header, '垫付成本')
    c_paid = settlement_find_col(header, '实付人民币')
    c_total = settlement_find_col(header, '总金额')
    records = {}
    for r in range(2, ws.max_row + 1):
        oid = ws.cell(r, c_oid).value
        if oid in (None, ''):
            continue
        oid = str(oid).strip()
        paid = parse_money_number(ws.cell(r, c_paid).value) if c_paid else 0.0
        if not paid and c_total:
            paid = parse_money_number(ws.cell(r, c_total).value)
        records[oid] = {
            'advance_idr': parse_money_number(ws.cell(r, c_advance).value) if c_advance else 0.0,
            'paid_cny': paid,
        }
    return records


def write_indonesia_sales_sheet(ws, sales_rows, rate):
    """新版印尼销售明细：A日期、B订单号，按日期升序稳定排列。"""
    template_styles = [copy(ws.cell(2, c)._style) for c in range(1, 13)]

    clear_to = max(ws.max_row, len(sales_rows) + 2)
    for row in ws.iter_rows(min_row=2, max_row=clear_to, min_col=1, max_col=12):
        for cell in row:
            cell.value = None
            cell.comment = None

    thin = Side(border_style='thin', color=Color(auto=True))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font11 = Font(name='宋体', size=11)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(vertical='center')
    supplement_fill = PatternFill(fill_type='solid',
                                  fgColor=Color(theme=5, tint=0.8),
                                  bgColor=Color(indexed=64))
    no_fill = PatternFill(fill_type=None)

    for idx, item in enumerate(sales_rows, start=2):
        values = {
            1: item['date'], 2: item['oid'], 3: item['sku'], 4: item['qty'],
            5: item['sales'], 6: item['disc'], 7: item['coupon'], 8: item['refund'],
            9: f'=E{idx}+F{idx}+G{idx}+H{idx}',
            10: f'=I{idx}*{rate}', 11: item['cost'],
            12: 0 if item['supplement'] else f'=J{idx}-K{idx}',
        }
        for c, value in values.items():
            cell = ws.cell(idx, c)
            cell._style = copy(template_styles[c - 1])
            cell.value = value
            if item['supplement'] and c <= 6:
                cell.fill = supplement_fill
        if item.get('cost_missing'):
            ws.cell(idx, 11).fill = YELLOW_FILL

    total_row = len(sales_rows) + 2
    last_row = len(sales_rows) + 1
    for c in (1, 2):
        ws.cell(total_row, c)._style = copy(ws.cell(total_row + 1, c)._style)
    bold = Font(name='宋体', size=11, bold=True)
    for c in range(4, 13):
        letter = get_column_letter(c)
        cell = ws.cell(total_row, c)
        cell.value = f'=SUM({letter}2:{letter}{last_row})'
    if ws.sheet_view.pane and ws.sheet_view.pane.ySplit == 1:
        ws.sheet_view.pane.topLeftCell = f'A{max(2, len(sales_rows) - 31)}'


def write_indonesia_unmatched_sheet(wb, sheet_map, ws_det, unmatched, last_col,
                                    preserve_label_header=False, country='印尼'):
    """新版印尼未匹配：复制订单明细并增加分类/成本两列，表头保持28店示例为空。"""
    if sheet_map.get('unmatched') in wb.sheetnames:
        ws = wb[sheet_map['unmatched']]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('未匹配')
        sheet_map['unmatched'] = ws.title

    if ws.freeze_panes:
        freeze_col = re.match(r'([A-Z]+)', str(ws.freeze_panes))
        if freeze_col:
            ws.freeze_panes = f'{freeze_col.group(1)}2'
    if country == '印尼':
        ws.conditional_formatting._cf_rules.clear()
        sales_ws = wb[sheet_map['sales']]
        added = False
        for rules in sales_ws.conditional_formatting._cf_rules.values():
            for rule in rules:
                if rule.type == 'duplicateValues':
                    ws.conditional_formatting.add('A1:A1048576', copy(rule))
                    added = True
                    break
            if added:
                break
        if not added:
            ws.conditional_formatting.add(
                'A1:A1048576', Rule(type='duplicateValues', dxfId=0, priority=1))
    if country == '菲律宾':
        ws.row_dimensions[1].height = 28.5

    for c in range(1, last_col + 1):
        copy_cell_all(ws_det.cell(1, c), ws.cell(1, c))
        letter = get_column_letter(c)
        if ws_det.column_dimensions[letter].width is not None:
            ws.column_dimensions[letter].width = ws_det.column_dimensions[letter].width

    label_col = last_col + 1
    cost_col = last_col + 2
    if preserve_label_header and label_col <= ws_det.max_column:
        copy_cell_all(ws_det.cell(1, label_col), ws.cell(1, label_col))
    else:
        ws.cell(1, label_col).value = None
    ws.cell(1, cost_col).value = None
    if country == '菲律宾':
        ws.cell(1, cost_col).font = Font(name='宋体', size=11, charset=134)
        ws.cell(1, cost_col).alignment = Alignment(horizontal='center', vertical='center',
                                                    wrap_text=True)
        header_border = all_thin_border()
        header_border.diagonal = Side()
        ws.cell(1, cost_col).border = header_border
        ws.column_dimensions['C'].width = 13
        ws.column_dimensions[get_column_letter(cost_col)].width = 14
    for column in (label_col, cost_col):
        if column <= ws_det.max_column:
            letter = get_column_letter(column)
            if ws_det.column_dimensions[letter].width is not None:
                ws.column_dimensions[letter].width = ws_det.column_dimensions[letter].width
    if country == '菲律宾':
        ws.column_dimensions[get_column_letter(cost_col)].width = 14
    for i, item in enumerate(unmatched, start=2):
        if country == '菲律宾':
            ws.row_dimensions[i].height = 15
        for c in range(1, last_col + 1):
            copy_cell_all(ws_det.cell(item['r'], c), ws.cell(i, c))
        label = ws.cell(i, label_col)
        if country == '菲律宾':
            label_style_col = label_col if label_col <= ws_det.max_column else last_col
            label._style = copy(ws_det.cell(item['r'], label_style_col)._style)
            label.font = Font(name='宋体', size=10, charset=134)
            label.fill = (PatternFill(fill_type='solid', fgColor='FFF8F8FF',
                                      bgColor=Color(indexed=64))
                          if item['label'] == '取消未发货'
                          else PatternFill(fill_type=None))
        else:
            label.font = Font(name='宋体', size=11, charset=134)
            label.alignment = Alignment(vertical='center')
        label.value = item['label']
        cost_cell = ws.cell(i, cost_col)
        if item['cost_bearing']:
            cost_cell.font = Font(name='宋体', size=11, charset=134)
            cost_cell.alignment = (Alignment(horizontal='center', vertical='center')
                                   if country == '菲律宾'
                                   else Alignment(vertical='center'))
            if item.get('cost') is None:
                cost_cell.value = 0
                cost_cell.fill = YELLOW_FILL
            else:
                cost_cell.value = num_or_int(item['cost'])
        elif country == '菲律宾':
            cost_cell.font = Font(name='宋体', size=11, charset=134)
            cost_cell.alignment = Alignment(horizontal='center', vertical='center')


def write_indonesia_bill_totals(ws, last_data_row, first_col, last_col):
    """按28店示例在账单末行添加数值列合计公式。"""
    total_row = last_data_row + 1
    for c in range(first_col, last_col + 1):
        letter = get_column_letter(c)
        cell = ws.cell(total_row, c)
        cell.value = f'=SUM({letter}2:{letter}{last_data_row})'
        cell.font = Font(name='Calibri', size=11, bold=True)


def append_indonesia_supplement_bills(ws_supplement, ws_bill, supplement_oids,
                                      bill_rows, last_col):
    """把补单订单对应的账单行附在补单明细第19行起。"""
    if ws_supplement is None or not supplement_oids:
        return
    ws_supplement.column_dimensions['E'].width = 20.875
    ws_supplement.column_dimensions['N'].width = 12.625
    ws_supplement.column_dimensions['O'].width = 11.5
    for row in ws_supplement.iter_rows(min_row=19, max_row=max(ws_supplement.max_row, 34),
                                       min_col=1, max_col=last_col):
        for cell in row:
            cell.value = None
    for c in range(1, last_col + 1):
        copy_cell_all(ws_bill.cell(1, c), ws_supplement.cell(19, c))
    out_row = 20
    for bill in bill_rows:
        if bill['oid'] not in supplement_oids:
            continue
        for c in range(1, last_col + 1):
            copy_cell_all(ws_bill.cell(bill['r'], c), ws_supplement.cell(out_row, c))
        out_row += 1
    for row in range(31, 35):
        cell = ws_supplement.cell(row, 12)
        cell.font = Font(name='宋体', size=11, charset=134, scheme='minor',
                         color=Color(theme=1))
        cell.alignment = Alignment(vertical='center')


def find_indonesia_ads_expense(ws):
    """兼容旧调用名；实际按统一广告 Amount/Expense 口径读取。"""
    return read_ads(ws)


def write_indonesia_pnl_sheet(ws, target_month, *, sales_cny, supplement_cny,
                              refund_cny, product_cost, sales_orders,
                              supplement_orders, cross_month_orders,
                              shipped_unmatched_orders,
                              ads_cny, supplement_fee_cny, platform_fee_cny,
                              cancel_unshipped, cancel_shipped, shipped_unpaid):
    """写入新版印尼利润表；示例中的跨月成本和发单费错误按计算说明纠正。"""
    row_map = {}
    for cell in ws['A']:
        if cell.value is not None:
            row_map[str(cell.value).replace(' ', '').replace('\t', '')] = cell.row

    def find_row(key):
        for label, row in row_map.items():
            if key in label:
                return row
        return None

    header_row = None
    month_cols = {}
    for row in ws.iter_rows(min_row=1, max_row=5):
        for cell in row:
            match = re.fullmatch(r'(\d{1,2})\s*月', str(cell.value or '').strip())
            if match:
                header_row = cell.row
                month_cols[int(match.group(1))] = cell.column
    if not month_cols:
        raise ValueError("利润表中找不到月份列")
    if target_month in month_cols:
        col = month_cols[target_month]
    else:
        prev_col = max(month_cols.values())
        col = prev_col + 1
        prev_letter = get_column_letter(prev_col)
        new_letter = get_column_letter(col)
        for r in range(header_row, max(row_map.values()) + 1):
            src, dst = ws.cell(r, prev_col), ws.cell(r, col)
            dst._style = copy(src._style)
            if isinstance(src.value, str) and src.value.startswith('='):
                dst.value = strip_minus_zero(re.sub(rf'\b{prev_letter}(\d+)', rf'{new_letter}\1', src.value))
        ws.cell(header_row, col).value = f'{target_month}月'

    def put(label, value):
        row = find_row(label)
        if row is None:
            raise ValueError(f"利润表找不到行: {label}")
        ws.cell(row, col).value = (
            round(value, 10) if isinstance(value, float) else value)
        ws.cell(row, col).comment = None

    put('本期销售', sales_cny)
    put('补单金额', supplement_cny)
    put('销售退款', refund_cny)
    put('实际产品成本', product_cost)
    dispatch_formula = f'=({sales_orders}-{supplement_orders}'
    if cross_month_orders:
        dispatch_formula += f'-{cross_month_orders}'
    dispatch_formula += f'+{shipped_unmatched_orders})*{fmt_num(DISPATCH_FEES["IDR"])}'
    put('发单费', dispatch_formula)
    put('广告费', ads_cny)
    put('补单费', supplement_fee_cny)
    put('平台费用', platform_fee_cny)

    profit_row = find_row('本期利润')
    if profit_row:
        note = f'{target_month}月未匹配：取消未发货{cancel_unshipped}，取消已发货{cancel_shipped}'
        if shipped_unpaid:
            note += f'，发货未到账{shipped_unpaid}'
        cell = ws.cell(profit_row + 1, 1)
        cell.value = note
        cell.font = Font(name='宋体', size=10, bold=True, charset=134,
                         color='FFFF0000')
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def write_settlement_pnl_sheet(ws, target_month, rate, *, sum_efg, sum_refund,
                    sales_cost_total, unmatched_cost_total,
                    unique_orders, not_shipped_orders,
                    unm_shipped_cnt, unm_unshipped_cnt,
                    platform_fees, adj_total, ads_value, ads_rate=USD_RATE,
                    country=None, cancel_shipped_cnt=None, shipped_unpaid_cnt=0,
                    presale_cnt=0):
    """利润表: 找到/新建目标月份列并填写数据/公式/批注/红字备注"""
    # ---- A1 标题: 富文本, '利润表'宋体12黑色 + 第二行'单位：CNY'宋体8红色 ----
    indent = 9 if country == '菲律宾' else 10
    ws['A1'] = CellRichText(
        TextBlock(InlineFont(rFont='宋体', charset=134, scheme='minor',
                             b=True, sz=12, color=Color(theme=1)), '利润表'),
        TextBlock(InlineFont(rFont='宋体', charset=134, scheme='minor',
                             b=True, sz=8, color=Color(rgb='FFFF0000')),
                  '\n' + ' ' * indent + '单位：CNY'),
    )
    if country == '菲律宾':
        ws.column_dimensions['A'].width = 19.375
    elif country == '马来':
        ws.column_dimensions['A'].width = 23.5

    # ---- A列行标签映射 ----
    row_map = {}
    for row in ws.iter_rows(min_row=1, max_col=1):
        v = row[0].value
        if v is not None:
            row_map[str(v).replace(' ', '').replace('\t', '')] = row[0].row

    def find_row(key):
        for k, r in row_map.items():
            if key in k:
                return r
        return None

    # ---- 定位/新建月份列 ----
    header_row = None
    month_cols = {}   # 月份 -> 列号
    for row in ws.iter_rows(min_row=1, max_row=5):
        for cell in row:
            if cell.value is None:
                continue
            m = re.fullmatch(r'(\d{1,2})\s*月', str(cell.value).strip())
            if m:
                header_row = cell.row
                month_cols[int(m.group(1))] = cell.column
    if not month_cols:
        print("   [错误] 利润表中找不到任何 'N月' 表头")
        return

    if target_month in month_cols:
        col = month_cols[target_month]
        print(f"   利润表已有 {target_month}月 列: 第{col}列, 直接填写")
        # 清理该列公式中手工失误的 '-0' (如 =F4-0-F5)
        for r in range(header_row, max(row_map.values()) + 1):
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = strip_minus_zero(cell.value)
    else:
        prev_col = max(month_cols.values())
        col = prev_col + 1
        prev_letter = get_column_letter(prev_col)
        new_letter = get_column_letter(col)
        print(f"   利润表新建 {target_month}月 列: 第{col}列 (样式/公式复制自 {prev_letter} 列)")
        last_label_row = max(row_map.values())
        for r in range(header_row, last_label_row + 1):
            src = ws.cell(row=r, column=prev_col)
            dst = ws.cell(row=r, column=col)
            dst._style = copy(src._style)
            if isinstance(src.value, str) and src.value.startswith('='):
                # 公式列字母平移 (=E4-E5 -> =F4-F5), 并清理手工失误的 '-0'
                dst.value = strip_minus_zero(
                    re.sub(rf'\b{prev_letter}(\d+)', rf'{new_letter}\1', src.value))
        ws.cell(row=header_row, column=col).value = f'{target_month}月'
        r15 = find_row('团队运营费用')
        if r15:
            ws.cell(row=r15, column=col).value = 0

    letter = get_column_letter(col)

    def put(key, value, formula=None, comment=None):
        r = find_row(key)
        if r is None:
            print(f"   [警告] 利润表找不到行: {key}")
            return
        cell = ws.cell(row=r, column=col)
        cell.value = (formula if formula is not None else
                      (round(value, 10) if isinstance(value, float) else value))
        cell.comment = None
        if comment is not None:
            cell.comment = Comment(f'Administrator:\n{comment}', 'Administrator',
                                   height=79, width=144)

    refund_cny = abs(sum_refund) * rate
    put('本期销售', round(sum_efg * rate + refund_cny, 7))
    put('补单金额', 0)
    put('销售退款', refund_cny)
    put('实际产品成本', round(sales_cost_total, 4) + round(unmatched_cost_total, 4))
    dispatch = str(unique_orders)
    if not_shipped_orders:
        dispatch += f'-{not_shipped_orders}'
    if unm_shipped_cnt:
        dispatch += f'+{unm_shipped_cnt}'
    put('发单费', None, formula=f'=({dispatch})*{fmt_num(DISPATCH_FEES[country])}')
    put('广告费', ads_value * ads_rate)
    put('补单费', 0)
    put('平台费用', abs(platform_fees) * rate)
    if find_row('调整金额') is not None:
        put('调整金额', abs(adj_total) * rate)

    # 统一目标月份公式，不沿用模板中可能存在的“减0”手工写法。
    # 只校正本月和紧邻上月；更早月份保持已确认的历史口径不动。
    all_month_columns = {col}
    if col - 1 in month_cols.values():
        all_month_columns.add(col - 1)
    for formula_col in all_month_columns:
        formula_letter = get_column_letter(formula_col)
        formulas = {
            '实际主营收入': f'={formula_letter}{find_row("本期销售")}-{formula_letter}{find_row("补单金额")}-{formula_letter}{find_row("销售退款")}',
            '实际主营成本': f'={formula_letter}{find_row("实际产品成本")}+{formula_letter}{find_row("发单费")}',
            '推广费用': f'={formula_letter}{find_row("广告费")}+{formula_letter}{find_row("补单费")}',
            '营业费用': f'={formula_letter}{find_row("推广费用")}+{formula_letter}{find_row("Shopee结算")}',
            'Shopee结算': f'={formula_letter}{find_row("平台费用")}',
            '本期利润': f'={formula_letter}{find_row("实际主营收入")}-{formula_letter}{find_row("实际主营成本")}-{formula_letter}{find_row("营业费用")}',
        }
        for label, formula in formulas.items():
            row = find_row(label)
            if row:
                ws.cell(row=row, column=formula_col).value = formula

    # ---- 底部红字备注: '四、本期利润' 的下一行 A 列 ----
    profit_row = find_row('本期利润')
    if profit_row:
        note_cell = ws.cell(row=profit_row + 1, column=1)
        cancel_shipped = unm_shipped_cnt if cancel_shipped_cnt is None else cancel_shipped_cnt
        note = (f'{target_month}月未匹配：取消未发货{unm_unshipped_cnt}，'
                f'取消已发货{cancel_shipped}')
        if shipped_unpaid_cnt:
            note += f'，发货未到账{shipped_unpaid_cnt}'
        if presale_cnt:
            note += f'，预售单{presale_cnt}'
        note_cell.value = note
        note_cell.font = Font(name='宋体', size=10, bold=True, charset=134,
                              color='FFFF0000')
        note_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def process_new_indonesia_workbook(wb, sheet_map, cost_map):
    """处理 Shopee 印尼新版账单；不改动旧账单及其他国家逻辑。"""
    rate = CURRENCY_RATES['IDR']
    ws_bill = wb[sheet_map['bill']]
    ws_det = wb[sheet_map['detail']]
    ws_dmj = wb[sheet_map['dmj']]
    ws_pnl = wb[sheet_map['pnl']]
    ws_sales = wb[sheet_map['sales']]
    ws_ads = wb[sheet_map['ads']] if sheet_map.get('ads') else None
    ws_supplement = wb[sheet_map['supplement']] if sheet_map.get('supplement') else None

    bill_header = ws_bill[1]
    c_b_oid = settlement_find_col(bill_header, 'Order ID', 1)
    c_b_created = settlement_find_col(bill_header, 'Order Creation Date', 5)
    c_b_sales = settlement_find_col(bill_header, 'Product Price', 10)
    c_b_refund = settlement_find_col(bill_header, 'Refund Amount', 11)
    c_b_fee_start = settlement_find_col(bill_header, 'Shipping Fee Paid by Buyer', 12)
    c_b_disc = settlement_find_col(bill_header, 'Voucher Sponsored by Seller', 19)
    c_b_coupon = settlement_find_col(bill_header, 'Rebate Provided by Shopee', 21)
    c_b_ads = settlement_find_col(bill_header, 'Ads Escrow Top Up Fee')
    c_b_fee_end = settlement_find_col(bill_header, 'WHT 22', min(37, ws_bill.max_column))
    c_b_username = settlement_find_col(bill_header, 'Username (Buyer)', c_b_fee_end + 1)

    bill_rows = []
    last_bill_data_row = 1
    for r in range(2, ws_bill.max_row + 1):
        oid = ws_bill.cell(r, c_b_oid).value
        if oid in (None, ''):
            continue
        last_bill_data_row = r
        fee_sum = sum(safe_float(ws_bill.cell(r, c).value)
                      for c in range(c_b_fee_start, c_b_fee_end + 1)
                      if is_excel_number(ws_bill.cell(r, c).value))
        disc = safe_float(ws_bill.cell(r, c_b_disc).value)
        coupon = safe_float(ws_bill.cell(r, c_b_coupon).value)
        ads_escrow = safe_float(ws_bill.cell(r, c_b_ads).value) if c_b_ads else 0.0
        bill_rows.append({
            'r': r, 'oid': str(oid).strip(),
            'date': ws_bill.cell(r, c_b_created).value,
            'sales': safe_float(ws_bill.cell(r, c_b_sales).value),
            'disc': disc, 'coupon': coupon,
            'refund': safe_float(ws_bill.cell(r, c_b_refund).value),
            'fee_component': fee_sum - disc - coupon - ads_escrow,
        })

    detail_header = ws_det[1]
    c_t_oid = settlement_find_col(detail_header, 'No. Pesanan', 1)
    c_t_status = settlement_find_col(detail_header, 'Status Pesanan', 2)
    c_t_ship = settlement_find_col(detail_header, 'Waktu Pengiriman Diatur', 9)
    c_t_created = settlement_find_col(detail_header, 'Waktu Pesanan Dibuat', 10)
    c_t_sku = settlement_find_col(detail_header, 'Nomor Referensi SKU', 16)
    c_t_qty = settlement_find_col(detail_header, 'Jumlah', 20)
    c_t_last = settlement_find_col(detail_header, 'Waktu Pesanan Selesai', ws_det.max_column)
    det_rows = []
    det_lines = {}
    order_status = {}
    order_shipped = {}
    for r in range(2, ws_det.max_row + 1):
        oid = ws_det.cell(r, c_t_oid).value
        if oid in (None, ''):
            continue
        oid = str(oid).strip()
        status = str(ws_det.cell(r, c_t_status).value or '').strip()
        sku_raw = ws_det.cell(r, c_t_sku).value
        sku = str(sku_raw or '').strip()
        qty = safe_float(ws_det.cell(r, c_t_qty).value)
        shipped = str(ws_det.cell(r, c_t_ship).value or '').strip()
        det_rows.append({'r': r, 'oid': oid, 'sku': sku, 'qty': qty,
                         'shipped': shipped, 'status': status,
                         'created': ws_det.cell(r, c_t_created).value})
        det_lines.setdefault(oid, []).append((sku_raw, qty))
        order_status.setdefault(oid, status)
        order_shipped[oid] = order_shipped.get(oid, False) or bool(shipped)
    det_oids = set(det_lines)

    dmj_header = ws_dmj[1]
    c_d_oid = settlement_find_col(dmj_header, '订单号', 1)
    c_d_sku = settlement_find_col(dmj_header, '商品SKU', 16)
    c_d_qty = settlement_find_col(dmj_header, ['实际销售数量', '数量'], 13)
    c_d_cost = settlement_find_col(dmj_header, ['总成本价', '商品成本'], 26)
    dmj = {}
    for r in range(2, ws_dmj.max_row + 1):
        oid = ws_dmj.cell(r, c_d_oid).value
        if oid in (None, ''):
            continue
        oid = str(oid).strip()
        item = {
            'sku': ws_dmj.cell(r, c_d_sku).value,
            'qty': ws_dmj.cell(r, c_d_qty).value,
            'cost': sum_numeric_cell(ws_dmj.cell(r, c_d_cost).value),
        }
        if oid in dmj:
            dmj[oid]['cost'] = max(dmj[oid]['cost'], item['cost'])
        else:
            dmj[oid] = item

    supplement_records = read_indonesia_supplements(ws_supplement)
    supplement_oids = set(supplement_records)
    matched = [b for b in bill_rows if b['oid'] in det_oids]
    if not matched:
        raise ValueError('新版印尼账单与订单明细无任何匹配订单')
    months = Counter(m for m in (parse_month(b['date']) for b in matched) if m)
    pending_month = find_pending_pnl_month(ws_pnl)
    target_month = pending_month if pending_month is not None else months.most_common(1)[0][0]

    sales_rows = []
    sales_cost_credited = set()
    missing_sales = []
    platform_by_oid = {}
    for bill in matched:
        oid = bill['oid']
        platform_by_oid[oid] = platform_by_oid.get(oid, 0.0) + bill['fee_component']
        supplement = oid in supplement_oids
        dm = dmj.get(oid)
        cost_missing = False
        sku, qty = indonesia_order_sku_qty(oid, det_lines)
        order_month = parse_month(bill['date'])
        cross_month = order_month is not None and order_month != target_month
        if supplement or cross_month:
            cost = 0.0
        elif dm is not None:
            if oid not in sales_cost_credited:
                cost = dm['cost']
                sales_cost_credited.add(oid)
            else:
                cost = 0.0
        else:
            cost = 0.0
            cost_missing = True
            missing_sales.extend(split_skus(sku) or [f'订单:{oid}'])
        sales_rows.append({
            'oid': oid, 'date': bill['date'],
            'sku': normalize_multiline_sku(sku, trailing_cr=cross_month), 'qty': qty,
            'sales': num_or_int(bill['sales']), 'disc': num_or_int(bill['disc']),
            'coupon': num_or_int(bill['coupon']), 'refund': num_or_int(bill['refund']),
            'cost': num_or_int(cost), 'cost_missing': cost_missing,
            'supplement': supplement, 'cross_month': cross_month,
        })

    # 行顺序按账单日期升序；同一天保持原账单顺序。
    sales_rows.sort(key=lambda row: str(row['date'] or ''))

    bill_oids = {b['oid'] for b in bill_rows}
    unmatched = []
    category_orders = {'取消未发货': set(), '取消已发货': set(), '发货未到账': set()}
    for detail in det_rows:
        oid = detail['oid']
        if oid in bill_oids:
            continue
        # 新导出会附加仅含订单号/SKU/成本的历史辅助行；如果它没有账单，
        # 28店示例也不把它列入未匹配。
        if not order_status.get(oid, '').strip():
            continue
        canceled = 'batal' in order_status.get(oid, '').lower()
        shipped = order_shipped.get(oid, False)
        if shipped and not canceled:
            label = '发货未到账'
        elif shipped:
            label = '取消已发货'
        else:
            label = '取消未发货'
        category_orders[label].add(oid)
        item = dict(detail)
        item['label'] = label
        item['cost_bearing'] = label in ('取消已发货', '发货未到账')
        order_month = parse_month(item['created'])
        item['cross_month'] = order_month is not None and order_month != target_month
        unmatched.append(item)

    unmatched_cost_total = 0.0
    missing_unmatched = []
    dmj_unmatched_credited = set()
    for item in unmatched:
        if not item['cost_bearing']:
            continue
        if item['cross_month']:
            item['cost'] = 0.0
            continue
        oid = item['oid']
        dm = dmj.get(oid)
        if dm is not None:
            if oid in dmj_unmatched_credited:
                item['cost'] = 0.0
            else:
                item['cost'] = round(safe_float(dm['cost']), 6)
                dmj_unmatched_credited.add(oid)
                unmatched_cost_total += item['cost']
        else:
            unit, missing = resolve_sku_cost(item['sku'], '印尼', cost_map)
            if missing:
                item['cost'] = None
                missing_unmatched.extend(missing)
            else:
                item['cost'] = round(unit * item['qty'], 6)
                unmatched_cost_total += item['cost']

    raw_sales_idr = sum(row['sales'] + row['disc'] + row['coupon'] + row['refund']
                        for row in sales_rows)
    refund_cny = abs(sum(row['refund'] for row in sales_rows)) * rate
    sales_cny = raw_sales_idr * rate + refund_cny
    supplement_cny = sum((row['sales'] + row['disc'] + row['coupon'] + row['refund']) * rate
                         for row in sales_rows if row['supplement'])
    sales_cost_total = sum(safe_float(row['cost']) for row in sales_rows)
    product_cost = round(sales_cost_total, 4) + round(unmatched_cost_total, 4)
    # 补单订单的平台扣费已经计入补单费，不再重复计入利润表“平台费用”。
    platform_fee_cny = abs(sum(
        bill['fee_component'] for bill in matched
        if bill['oid'] not in supplement_oids
    )) * rate
    ads_idr = find_indonesia_ads_expense(ws_ads)
    ads_cny = ads_idr * ADS_TAX_FACTORS['IDR'] * rate
    supplement_fee_cny = 0.0
    for oid, record in supplement_records.items():
        supplement_fee_cny += (record['paid_cny'] - record['advance_idr'] * rate
                               + abs(platform_by_oid.get(oid, 0.0)) * rate)

    unique_sales_orders = len({row['oid'] for row in sales_rows})
    matched_supplement_orders = len({row['oid'] for row in sales_rows if row['supplement']})
    shipped_unmatched_orders = len(category_orders['取消已发货'] | category_orders['发货未到账'])
    # 发单费只计算本期实际发出的订单：补单不重复计算，跨月订单也不在本期重复计费。
    # 补单若同时属于跨月订单，只扣一次；跨月未匹配已发货订单抵消上面的“+未匹配已发货”。
    cross_month_dispatch_orders = {
        row['oid'] for row in sales_rows
        if row['cross_month'] and not row['supplement']
    } | {
        item['oid'] for item in unmatched
        if item['cost_bearing'] and item['cross_month']
    }

    write_indonesia_sales_sheet(ws_sales, sales_rows, rate)
    write_indonesia_unmatched_sheet(
        wb, sheet_map, ws_det, unmatched, c_t_last, country='印尼')
    append_indonesia_supplement_bills(ws_supplement, ws_bill, supplement_oids,
                                      bill_rows, c_b_username)
    write_indonesia_pnl_sheet(
        ws_pnl, target_month,
        sales_cny=sales_cny, supplement_cny=supplement_cny,
        refund_cny=refund_cny, product_cost=product_cost,
        sales_orders=unique_sales_orders, supplement_orders=matched_supplement_orders,
        cross_month_orders=len(cross_month_dispatch_orders),
        shipped_unmatched_orders=shipped_unmatched_orders,
        ads_cny=ads_cny, supplement_fee_cny=supplement_fee_cny,
        platform_fee_cny=platform_fee_cny,
        cancel_unshipped=len(category_orders['取消未发货']),
        cancel_shipped=len(category_orders['取消已发货']),
        shipped_unpaid=len(category_orders['发货未到账']),
    )

    print(f"   新版印尼账单: 销售 {len(sales_rows)} 行/{unique_sales_orders} 单，"
          f"补单 {matched_supplement_orders} 单，汇率 {rate}")
    print(f"   未匹配: 取消未发货 {len(category_orders['取消未发货'])} 单，"
          f"取消已发货 {len(category_orders['取消已发货'])} 单，"
          f"发货未到账 {len(category_orders['发货未到账'])} 单")
    print(f"   跨月订单: 成本按0，发单费扣除 {len(cross_month_dispatch_orders)} 单")
    print(f"   利润表: 销售 {fmt_num(sales_cny)}，产品成本 {fmt_num(product_cost)}，"
          f"广告费 {fmt_num(ads_cny)}，补单费 {fmt_num(supplement_fee_cny)}，"
          f"平台费用 {fmt_num(platform_fee_cny)}")
    if missing_sales or missing_unmatched:
        print(f"   [注意] 无成本SKU已标黄: {sorted(set(missing_sales + missing_unmatched))}")


def write_settlement_sales_sheet(ws, sales_rows, rate, country):
    """写菲律宾/马来西亚新版销售明细：A日期、B订单号、E:H为四项本币金额。"""
    template_styles = [copy(ws.cell(2, c)._style) for c in range(1, 13)]
    ws.freeze_panes = 'A2'
    clear_to = max(ws.max_row, len(sales_rows) + 2)
    for row in ws.iter_rows(min_row=2, max_row=clear_to, min_col=1, max_col=12):
        for cell in row:
            cell.value = None
            cell.comment = None

    thin = Side(border_style='thin', color=Color(auto=True))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font11 = Font(name='宋体', size=11)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(vertical='center')
    no_fill = PatternFill(fill_type=None)

    for r, item in enumerate(sales_rows, start=2):
        values = {
            1: item['date'], 2: item['oid'], 3: item['sku'], 4: item['qty'],
            5: item['sales'], 6: item['disc'], 7: item['coupon'], 8: item['refund'],
            9: f'=E{r}+F{r}+G{r}+H{r}',
            10: f'=I{r}*{rate}', 11: item['cost'], 12: f'=J{r}-K{r}',
        }
        for c, value in values.items():
            cell = ws.cell(r, c)
            cell._style = copy(template_styles[c - 1])
            cell.value = value
        if item.get('cost_missing'):
            ws.cell(r, 11).fill = YELLOW_FILL

    last_row = len(sales_rows) + 1
    total_row = last_row + 1
    bold = Font(name='宋体', size=11, bold=True)
    for c in range(4, 13):
        letter = get_column_letter(c)
        cell = ws.cell(total_row, c)
        cell.value = f'=SUM({letter}2:{letter}{last_row})'
        if country == '菲律宾':
            cell.font = Font(name='宋体', size=11, bold=True, charset=134,
                             color='FF000000')
            cell.alignment = Alignment(vertical='center')


def process_new_non_indonesia_settlement_workbook(wb, sheet_map, cost_map):
    """处理菲律宾/马来西亚新版结算账单，沿用非印尼成本及发单费口径。"""
    ws_bill = wb[sheet_map['bill']]
    ws_det = wb[sheet_map['detail']]
    ws_dmj = wb[sheet_map['dmj']]
    ws_pnl = wb[sheet_map['pnl']]
    ws_sales = wb[sheet_map['sales']]
    ws_ads = wb[sheet_map['ads']] if sheet_map.get('ads') else None

    # 订单明细：新版英文导出从第2行开始，列名按当前导出动态定位。
    detail_header = list(ws_det.iter_rows(min_row=1, max_row=1,
                                          max_col=min(ws_det.max_column, 100)))[0]
    c_t_oid = settlement_find_col(detail_header, 'Order ID', 1)
    c_t_status = settlement_find_col(detail_header, 'Order Status', 2)
    c_t_sku = settlement_find_col(detail_header, 'SKU Reference No.', 15)
    c_t_qty = settlement_find_col(detail_header, 'Quantity', 19)
    c_t_ship = settlement_find_col(detail_header, ['Ship Time', 'Shipped Time'], 10)
    c_t_created = settlement_find_col(detail_header, 'Order Creation Date', 11)
    c_t_estimated = settlement_find_col(detail_header, 'Estimated Ship Out Date', 9)
    c_t_country = settlement_find_col(detail_header, 'Country')
    c_t_last = settlement_find_col(detail_header, ['Zip Code', 'Postal Code'], ws_det.max_column)

    det_rows = []
    det_lines = {}
    order_shipped = {}
    order_status = {}
    order_estimated = {}
    order_red = {}
    country_raw = None
    for r in range(2, ws_det.max_row + 1):
        oid_value = ws_det.cell(r, c_t_oid).value
        if oid_value in (None, ''):
            continue
        oid = str(oid_value).strip()
        sku_raw = ws_det.cell(r, c_t_sku).value
        sku = str(sku_raw or '').strip()
        qty = safe_float(ws_det.cell(r, c_t_qty).value)
        shipped = str(ws_det.cell(r, c_t_ship).value or '').strip()
        status = str(ws_det.cell(r, c_t_status).value or '').strip()
        if country_raw is None and c_t_country:
            value = ws_det.cell(r, c_t_country).value
            if value not in (None, ''):
                country_raw = str(value).strip()
        det_rows.append({
            'r': r, 'oid': oid, 'sku': sku, 'qty': qty,
            'shipped': shipped, 'status': status,
            'created': ws_det.cell(r, c_t_created).value,
            'estimated': ws_det.cell(r, c_t_estimated).value,
            'red': font_is_red(ws_det.cell(r, c_t_oid)),
        })
        det_lines.setdefault(oid, []).append((sku_raw, qty))
        order_shipped[oid] = order_shipped.get(oid, False) or bool(shipped)
        if status:
            order_status[oid] = status
        else:
            order_status.setdefault(oid, '')
        order_estimated.setdefault(oid, ws_det.cell(r, c_t_estimated).value)
        order_red[oid] = order_red.get(oid, False) or font_is_red(ws_det.cell(r, c_t_oid))
    det_oids = set(det_lines)
    country = COUNTRY_MAP.get(country_raw, country_raw)
    currency = COUNTRY_CURRENCY.get(country_raw) or COUNTRY_CURRENCY.get(country)
    if currency not in CURRENCY_RATES:
        raise ValueError(f'新版结算账单无法根据国家 {country_raw!r} 确定币种')
    rate = CURRENCY_RATES[currency]

    # 大卖家：销售与未匹配已发货订单都优先使用整单“总成本价”。
    dmj_header = list(ws_dmj.iter_rows(min_row=1, max_row=1,
                                       max_col=min(ws_dmj.max_column, 100)))[0]
    c_d_oid = settlement_find_col(dmj_header, '订单号', 1)
    c_d_sku = settlement_find_col(dmj_header, '商品SKU', 4)
    c_d_qty = settlement_find_col(dmj_header, ['实际销售数量', '数量'], 13)
    c_d_cost = settlement_find_col(dmj_header, ['总成本价', '商品成本'], 26)
    dmj = {}
    for r in range(2, ws_dmj.max_row + 1):
        oid_value = ws_dmj.cell(r, c_d_oid).value
        if oid_value in (None, ''):
            continue
        oid = str(oid_value).strip()
        item = {
            'sku': ws_dmj.cell(r, c_d_sku).value,
            'qty': ws_dmj.cell(r, c_d_qty).value,
            'cost': sum_numeric_cell(ws_dmj.cell(r, c_d_cost).value),
        }
        if oid in dmj:
            dmj[oid]['cost'] = max(dmj[oid]['cost'], item['cost'])
        else:
            dmj[oid] = item

    # 账单：菲律宾表头在第6行，马来西亚在第1行；字段动态定位。
    bill_header_row = find_table_header_row(ws_bill, ('Order ID', '订单编号', '订单ID'))
    if not bill_header_row:
        raise ValueError('新版结算账单找不到订单号表头')
    bill_header = list(ws_bill.iter_rows(min_row=bill_header_row, max_row=bill_header_row,
                                         max_col=min(ws_bill.max_column, 100)))[0]
    c_b_oid = settlement_find_col(bill_header, ['Order ID', '订单编号', '订单ID'])
    c_b_created = settlement_find_col(bill_header, ['Order Creation Date', '订单成立时间', '订单创建时间'])
    c_b_sales = settlement_find_col(bill_header, ['Original Product Price', 'Product Price'])
    c_b_refund = settlement_find_col(bill_header, ['Refund Amount', '退款金額', '退款金额'])
    c_b_shopee_rebate = settlement_find_col(
        bill_header, ['Rebate Provided by Shopee', 'Shopee回扣金额'])
    c_b_payout = settlement_find_col(bill_header, ['Total Released Amount', '拨款金额'])
    c_b_ads = settlement_find_col(bill_header, ['Ads Sales Top Up Fee', 'Ads Escrow Top Up Fee',
                                     '广告托管充值费用'])
    c_b_disc = settlement_find_col(
        bill_header, ['Seller Product Promotion', 'Your Seller Product Promotion',
                      '商品折扣', '卖家商品折扣'])
    coupon_aliases = (
        'Voucher Sponsored by Seller', '卖家赞助的优惠券',
        'Cofund Voucher Sponsored by Seller', '卖家赞助的合资优惠券',
        'Coin Cashback Sponsored by Seller', '卖家赞助的 Shopee 币回扣',
        'Cofund Coin Cashback Sponsored by Seller', '卖家赞助的合资 Shopee 币回扣',
    )
    coupon_cols = []
    for alias in coupon_aliases:
        col = settlement_find_col(bill_header, alias)
        if col and col not in coupon_cols:
            coupon_cols.append(col)
    required = {
        '订单号': c_b_oid, '订单成立时间': c_b_created, '商品价': c_b_sales,
        '退款': c_b_refund, '平台回扣': c_b_shopee_rebate, '拨款金额': c_b_payout,
    }
    missing_cols = [name for name, col in required.items() if not col]
    if missing_cols:
        raise ValueError(f'新版结算账单缺少必要字段: {missing_cols}')

    bill_rows = []
    for r in range(bill_header_row + 1, ws_bill.max_row + 1):
        oid_value = ws_bill.cell(r, c_b_oid).value
        if oid_value in (None, ''):
            continue
        sales = safe_float(ws_bill.cell(r, c_b_sales).value)
        seller_coupon = sum(safe_float(ws_bill.cell(r, c).value) for c in coupon_cols)
        disc = safe_float(ws_bill.cell(r, c_b_disc).value) if c_b_disc else 0.0
        coupon = safe_float(ws_bill.cell(r, c_b_shopee_rebate).value)
        if country == '菲律宾':
            coupon += seller_coupon
        else:
            disc += seller_coupon
        refund = safe_float(ws_bill.cell(r, c_b_refund).value)
        payout = safe_float(ws_bill.cell(r, c_b_payout).value)
        ads_topup = safe_float(ws_bill.cell(r, c_b_ads).value) if c_b_ads else 0.0
        # 拨款=销售四项+平台结算项+广告充值；广告充值单独计入广告费，故从平台费用剔除。
        fee_component = payout - sales - disc - coupon - refund - ads_topup
        bill_rows.append({
            'r': r, 'oid': str(oid_value).strip(),
            'date': ws_bill.cell(r, c_b_created).value,
            'sales': sales, 'disc': disc, 'coupon': coupon, 'refund': refund,
            'payout': payout, 'ads_topup': ads_topup, 'fee_component': fee_component,
        })

    matched = [row for row in bill_rows if row['oid'] in det_oids]
    if not matched:
        raise ValueError('新版结算账单与订单明细无任何匹配订单')
    months = Counter(m for m in (parse_month(row['date']) for row in matched) if m)
    pending_month = find_pending_pnl_month(ws_pnl)
    if pending_month is not None:
        target_month = pending_month
    elif months:
        target_month = months.most_common(1)[0][0]
    else:
        raise ValueError('无法识别利润表待填月份或订单月份')

    sales_rows = []
    sales_cost_credited = set()
    missing_sales = []
    for bill in matched:
        oid = bill['oid']
        dm = dmj.get(oid)
        cost_missing = False
        cross_month = (parse_month(bill['date']) is not None
                       and parse_month(bill['date']) != target_month)
        if dm is not None:
            if country == '菲律宾':
                sku = ','.join(split_skus(dm['sku']))
                sku = re.sub(r'\s+\+', '\n+', sku)
                qty = num_or_int(sum_numeric_cell(dm['qty']))
            else:
                sku = strip_export_tail(dm['sku'])
                qty = normalize_export_quantity(dm['qty'])
            if cross_month:
                cost = 0.0
            elif oid not in sales_cost_credited:
                cost = dm['cost']
                sales_cost_credited.add(oid)
            else:
                cost = 0.0
        else:
            sku, qty = order_sku_qty(oid, det_lines)
            sku = normalize_multiline_sku(sku)
            qty = normalize_export_quantity(qty)
            cost = 0.0
            # 跨月订单成本已在下单月计入，大卖家本月导出缺失属正常，不标黄。
            if not cross_month:
                cost_missing = True
                missing_sales.extend(split_skus(sku) or [f'订单:{oid}'])
        sales_rows.append({
            'date': bill['date'], 'oid': oid, 'sku': sku, 'qty': qty,
            'sales': num_or_int(bill['sales']), 'disc': num_or_int(bill['disc']),
            'coupon': num_or_int(bill['coupon']), 'refund': num_or_int(bill['refund']),
            'cost': num_or_int(cost), 'cost_missing': cost_missing,
            'cross_month': cross_month,
        })

    # 顶部即真实表头的新账单按日期升序；带元数据前导行的旧账单保持原顺序。
    if bill_header_row == 1:
        sales_rows.sort(key=lambda row: str(row['date'] or ''))

    unique_orders = len({row['oid'] for row in sales_rows})
    sales_cost_total = sum(safe_float(row['cost']) for row in sales_rows)
    sum_refund = sum(safe_float(row['refund']) for row in sales_rows)
    sum_sales_local = sum(row['sales'] + row['disc'] + row['coupon'] + row['refund']
                          for row in sales_rows)
    platform_fees = sum(row['fee_component'] for row in matched)

    bill_oids = {row['oid'] for row in bill_rows}
    unmatched = []
    category_orders = {
        '取消未发货': set(), '取消已发货': set(),
        '发货未到账': set(), '预售单': set(),
    }
    for detail in det_rows:
        oid = detail['oid']
        if oid in bill_oids:
            continue
        status = order_status.get(oid, '').strip()
        shipped = order_shipped.get(oid, False)
        count_in_summary = True
        if not status:
            # 菲律宾标红辅助行全部忽略；马来仅保留确实已有发货时间的跨月行。
            if country != '马来' or not shipped:
                continue
            label = '取消已发货'
            count_in_summary = False
        elif status.lower() in ('shipping', 'delivered'):
            label = '发货未到账'
        elif (status.lower() == 'to ship'
              and parse_month(order_estimated.get(oid)) not in (None, target_month)):
            label = '预售单'
        elif shipped:
            label = '取消已发货'
        else:
            label = '取消未发货'
        item = dict(detail)
        item['label'] = label
        item['cost_bearing'] = label in ('取消已发货', '发货未到账')
        order_month = parse_month(item['created'])
        item['cross_month'] = (count_in_summary and order_month is not None
                               and order_month != target_month)
        if count_in_summary:
            category_orders[label].add(oid)
        unmatched.append(item)
    unm_shipped = [item for item in unmatched if item['cost_bearing']]
    unm_unshipped_cnt = len(category_orders['取消未发货'])

    unmatched_cost_total = 0.0
    missing_unmatched = []
    dmj_unmatched_credited = set()
    for item in unm_shipped:
        oid = item['oid']
        if item['cross_month']:
            item['cost'] = 0.0
            continue
        dm = dmj.get(oid)
        if dm is not None:
            if oid in dmj_unmatched_credited:
                item['cost'] = 0.0
            else:
                item['cost'] = round(safe_float(dm['cost']), 6)
                dmj_unmatched_credited.add(oid)
                unmatched_cost_total += item['cost']
        else:
            unit, missing = resolve_sku_cost(item['sku'], country, cost_map)
            if missing:
                item['cost'] = None
                missing_unmatched.extend(missing)
            else:
                item['cost'] = round(unit * item['qty'], 6)
                unmatched_cost_total += item['cost']

    ads_local = find_indonesia_ads_expense(ws_ads)
    shipped_unmatched_orders = len(
        category_orders['取消已发货'] | category_orders['发货未到账'])
    cross_month_dispatch_orders = {
        row['oid'] for row in sales_rows if row['cross_month']
    } | {item['oid'] for item in unm_shipped if item['cross_month']}
    write_settlement_sales_sheet(ws_sales, sales_rows, rate, country)
    write_indonesia_unmatched_sheet(
        wb, sheet_map, ws_det, unmatched, c_t_last,
        preserve_label_header=(country == '菲律宾'), country=country)
    write_settlement_pnl_sheet(
        ws_pnl, target_month, rate,
        sum_efg=sum_sales_local, sum_refund=sum_refund,
        sales_cost_total=sales_cost_total, unmatched_cost_total=unmatched_cost_total,
        unique_orders=unique_orders,
        not_shipped_orders=len(cross_month_dispatch_orders),
        unm_shipped_cnt=shipped_unmatched_orders,
        unm_unshipped_cnt=unm_unshipped_cnt,
        platform_fees=platform_fees, adj_total=0.0,
        ads_value=ads_local * ADS_TAX_FACTORS[currency], ads_rate=rate,
        country=country,
        cancel_shipped_cnt=len(category_orders['取消已发货']),
        shipped_unpaid_cnt=len(category_orders['发货未到账']),
        presale_cnt=len(category_orders['预售单']),
    )

    print(f"   新版{country}账单: 销售 {len(sales_rows)} 行/{unique_orders} 单，"
          f"币种 {currency}，汇率 {rate}")
    print(f"   大卖家成本匹配: {len({row['oid'] for row in sales_rows} & set(dmj))}/"
          f"{unique_orders} 单 | 商品成本 {fmt_num(sales_cost_total)}")
    print(f"   未匹配: 取消未发货 {len(category_orders['取消未发货'])} 单，"
          f"取消已发货 {len(category_orders['取消已发货'])} 单，"
          f"发货未到账 {len(category_orders['发货未到账'])} 单，"
          f"预售单 {len(category_orders['预售单'])} 单")
    print(f"   利润表 {target_month}月: 广告费 "
          f"{fmt_num(ads_local * ADS_TAX_FACTORS[currency] * rate)}，"
          f"平台费用 {fmt_num(abs(platform_fees) * rate)}")
    if missing_sales or missing_unmatched:
        print(f"   [注意] 无成本SKU已标黄: {sorted(set(missing_sales + missing_unmatched))}")


def normalize_workbook_view_state(sheets, country):
    """同步示例工作簿的冻结窗格、打开位置和展示列宽。"""
    if sheets.get('pnl') is not None:
        set_pnl_title(sheets['pnl'])
    ads = sheets.get('ads')
    if country in ('泰国', '菲律宾') and ads is not None:
        ads.freeze_panes = None
        active = (f'F{ads.max_row + 4}' if country == '泰国' else 'F33')
        ads.sheet_view.selection = [Selection(activeCell=active, sqref=active)]
    if country == '泰国' and ads is not None:
        ads.column_dimensions['D'].width = 13.75
    if country == '印尼' and ads is not None and ads.sheet_view.pane:
        ads.sheet_view.pane.topLeftCell = f'A{max(2, ads.max_row - 22)}'
    if country == '印尼' and sheets.get('income') is not None:
        income = sheets['income']
        if income.sheet_view.pane:
            income.sheet_view.pane.topLeftCell = f'B{max(2, income.max_row - 38)}'
    if country == '菲律宾' and sheets.get('dmj') is not None:
        dmj = sheets['dmj']
        if dmj.sheet_view.pane:
            dmj.sheet_view.pane.topLeftCell = f'B{max(3, dmj.max_row - 50)}'
    if country == '马来':
        income = sheets.get('income')
        order = sheets.get('order')
        if income is not None and income.sheet_view.pane:
            income.sheet_view.pane.topLeftCell = 'B2'
        if order is not None:
            order.freeze_panes = None
            active = f'H{max(2, order.max_row - 23)}'
            order.sheet_view.selection = [Selection(activeCell=active, sqref=active)]
        if ads is not None:
            ads.freeze_panes = 'A8'
            ads.sheet_view.selection = [
                Selection(activeCell='A1', sqref='A1'),
                Selection(pane='bottomLeft', activeCell='H22', sqref='H22'),
            ]

    # 少量订单/广告行的表格也必须把滚动位置放在冻结区域之外。
    for ws in sheets.values():
        if ws is None:
            continue
        pane = ws.sheet_view.pane
        if pane is not None and pane.state in ('frozen', 'frozenSplit'):
            row, col = coordinate_to_tuple(pane.topLeftCell or 'A1')
            row = max(row, int(pane.ySplit or 0) + 1)
            col = max(col, int(pane.xSplit or 0) + 1)
            pane.topLeftCell = f'{get_column_letter(col)}{row}'


def process_file(file_path, out_dir):
    file_name = os.path.basename(file_path)
    country = detect_country(file_name)
    if country is None:
        print(f"[跳过] 文件名无法识别国家: {file_name}")
        return None
    rate = CURRENCY_RATES[country]
    print(f"\n{'=' * 60}\n>>> 处理 {file_name} | {country}汇率 {rate}")

    output_path = os.path.join(out_dir, f'新_{file_name}')
    if os.path.exists(output_path):
        try:
            os.chmod(output_path, stat.S_IWRITE)
            os.remove(output_path)
        except PermissionError:
            stem, extension = os.path.splitext(file_name)
            suffix = 1
            while True:
                label = '修正版' if suffix == 1 else f'修正版{suffix}'
                candidate = os.path.join(out_dir, f'新_{stem}_{label}{extension}')
                if not os.path.exists(candidate):
                    output_path = candidate
                    print(f"   原输出文件正在打开，改为输出: {os.path.basename(output_path)}")
                    break
                suffix += 1
    shutil.copy2(file_path, output_path)
    os.chmod(output_path, stat.S_IWRITE)

    wb = openpyxl.load_workbook(output_path, rich_text=True)
    sheets = {key: get_sheet(wb, names) for key, names in SHEET_NAMES.items()}
    missing = [key for key in ('pnl', 'income', 'order', 'dmj') if sheets[key] is None]
    if missing:
        wb.close()
        os.remove(output_path)
        raise ValueError(f"{file_name} 缺少必要 Sheet: {missing}")

    sheet_map = {
        'pnl': sheets['pnl'].title,
        'bill': sheets['income'].title,
        'detail': sheets['order'].title,
        'dmj': sheets['dmj'].title,
    }
    for source_key, target_key in (
            ('sales', 'sales'), ('ads', 'ads'), ('brush', 'supplement'),
            ('unmatched', 'unmatched')):
        if sheets[source_key] is not None:
            sheet_map[target_key] = sheets[source_key].title

    new_format = None
    if is_new_indonesia_format(wb, sheet_map):
        if 'sales' not in sheet_map:
            raise ValueError(f"{file_name} 缺少销售明细 Sheet")
        process_new_indonesia_workbook(wb, sheet_map, {})
        new_format = '新版印尼'
    elif is_new_non_indonesia_settlement_format(wb, sheet_map):
        if 'sales' not in sheet_map:
            raise ValueError(f"{file_name} 缺少销售明细 Sheet")
        process_new_non_indonesia_settlement_workbook(wb, sheet_map, {})
        new_format = '新版非印尼'

    if new_format:
        normalize_workbook_view_state(sheets, country)
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = 'auto'
        except AttributeError:
            pass
        wb.save(output_path)
        wb.close()
        preserve_sheet_carriage_returns(output_path, sheet_map['sales'])
        restore_sheet_comments(file_path, output_path, sheet_map['pnl'])
        if sheet_map.get('ads'):
            restored = restore_sheet_drawings(file_path, output_path, sheet_map['ads'])
            if restored:
                print(f"   广告表图片已保留: {restored} 张")
        preserve_rich_text_spaces(output_path)
        print(f">>> 完成: {output_path}")
        return {'output': output_path, 'format': new_format}

    cost_map = load_order_cost_map_from_dmj(sheets['dmj'])
    sku_cost_map = load_sku_cost_map_from_dmj(sheets['dmj'])
    income_agg, income_seq = read_income(sheets['income'])
    order_agg, order_seq, order_rows, missing_cost_orders, n_order_cols = \
        read_orders(sheets['order'], cost_map, sku_cost_map)

    income_ids = set(income_agg)
    order_ids = set(order_agg)
    sales_ids = [oid for oid in income_seq if oid in order_ids]
    sales_months = Counter(d.month for d in (parse_date(income_agg[oid]['date'])
                                             for oid in sales_ids) if d)
    if not sales_months:
        wb.close()
        os.remove(output_path)
        raise ValueError(f"{file_name} 销售明细没有可识别月份")
    month = sales_months.most_common(1)[0][0]
    brush_ids, brush_cost, brush_sub, brush_comm = read_brush(sheets['brush'])
    # 跨月订单的成本和发单费已在下单月计入（当时为未匹配已发货），本月不重复计算。
    cross_month_ids = {
        oid for oid, order in order_agg.items()
        if oid not in brush_ids
        and (d := parse_date(order['date'])) is not None and d.month != month
    }
    sheets['sales'], sales = write_sales_sheet(
        wb, sheets['sales'], sales_ids, order_agg, income_agg, rate,
        missing_cost_orders, cross_month_ids)

    unmatched_last_col = find_col(
        sheets['order'], ORDER_HEADER_ROW,
        ['ประเภทคำสั่งซื้อ', 'Order Type', 'Zip Code', 'Postal Code'],
        n_order_cols)
    unmatched = write_unmatched_sheet(
        wb, sheets['unmatched'], sheets['order'], order_rows, income_ids,
        unmatched_last_col, cross_month_ids)
    ads_sum = read_ads(sheets['ads'])

    refund = -sales['sum_h'] * rate
    current_sales = sales['sum_j'] + refund
    brush_amount = sum(sales['j_by_id'].get(oid, 0.0) for oid in brush_ids)
    # 取消未发货只保留在未匹配明细和利润表注释中，不计产品成本或发单费。
    product_cost = sales['sum_k'] + unmatched['shipped_cost']
    shipped_unmatched = unmatched['cancel_shipped'] + unmatched['unpaid']
    dispatch = str(len(sales_ids))
    if brush_ids:
        dispatch += f'-{len(brush_ids)}'
    cross_dispatch = (len([oid for oid in sales_ids if oid in cross_month_ids])
                      + unmatched['cross_shipped'])
    if cross_dispatch:
        dispatch += f'-{cross_dispatch}'
    if shipped_unmatched:
        dispatch += f'+{shipped_unmatched}'
    order_fee = f'=({dispatch})*{fmt_num(DISPATCH_FEES[country])}'
    ad_fee = ads_sum * ADS_TAX_FACTORS[country] * rate
    brush_platform = abs(sum(income_agg.get(oid, {}).get('fees', 0.0)
                             for oid in brush_ids)) * rate
    brush_exchange = brush_sub - brush_cost * rate
    brush_fee = brush_comm + brush_platform + brush_exchange
    platform_fee = abs(sales['fee_sum']) * rate

    values = {
        '本期销售': current_sales,
        '补单金额': brush_amount,
        '销售退款': refund,
        '实际产品成本': product_cost,
        '发单费': order_fee,
        '广告费': ad_fee,
        '补单费': brush_fee,
        '平台费用': platform_fee,
    }
    write_pnl_sheet(
        sheets['pnl'], month, values,
        (unmatched['cancel_unshipped'], unmatched['cancel_shipped'], unmatched['unpaid']))

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = 'auto'
    except AttributeError:
        pass
    normalize_workbook_view_state(sheets, country)
    wb.save(output_path)
    wb.close()
    restore_sheet_comments(file_path, output_path, sheet_map['pnl'])

    preserve_rich_text_spaces(output_path)

    if cross_dispatch:
        print(f"   跨月订单: 成本按0，发单费扣除 {cross_dispatch} 单")
    missing_cost_orders -= cross_month_ids
    print(f"   大卖家成本订单 {len(cost_map)} | 销售订单 {len(sales_ids)} | "
          f"未匹配 {unmatched['order_count']}单/{unmatched['row_count']}行")
    print(f"   未匹配：取消未发货{unmatched['cancel_unshipped']}，"
          f"取消已发货{unmatched['cancel_shipped']}，发货未到账{unmatched['unpaid']}")
    if unmatched['excluded_red_ids']:
        print(f"   标红且未匹配订单已忽略: {unmatched['excluded_red_ids']}")
    if missing_cost_orders:
        sample = sorted(missing_cost_orders)[:20]
        print(f"   [警告] 大卖家无成本订单 {len(missing_cost_orders)} 个（按0），"
              f"前20个: {sample}")
    print(f"   本期销售 {current_sales:.2f} | 产品成本 {product_cost:.2f} | "
          f"广告费 {ad_fee:.2f} | 补单费 {brush_fee:.2f} | 平台费用 {platform_fee:.2f}")
    print(f">>> 完成: {output_path}")
    return {
        'output': output_path,
        'month': month,
        'sales_orders': len(sales_ids),
        'unmatched': unmatched,
        'values': values,
        'missing_cost_orders': sorted(missing_cost_orders),
    }


def resolve_path(path):
    if not path:
        return None
    return path if os.path.isabs(path) else os.path.join(SCRIPT_DIR, path)


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else resolve_path(TARGET_FOLDER_PATH)
    if not folder or not os.path.isdir(folder):
        print(f"找不到目标文件夹: {folder}")
        return 1
    out_dir = resolve_path(OUTPUT_FOLDER_PATH) or folder
    os.makedirs(out_dir, exist_ok=True)

    source_files = []
    for path in sorted(glob.glob(os.path.join(folder, '*.xlsx'))):
        name = os.path.basename(path)
        if name.startswith(('新_', '~$')):
            continue
        source_files.append(path)
    print(f"找到 {len(source_files)} 个店铺文件: {[os.path.basename(p) for p in source_files]}")

    failures = 0
    for path in source_files:
        try:
            process_file(path, out_dir)
        except Exception as exc:
            failures += 1
            print(f"[失败] {os.path.basename(path)}: {exc}")
            import traceback
            traceback.print_exc()
    print(f"\n全部执行完毕：成功 {len(source_files) - failures}，失败 {failures}")
    return 1 if failures else 0


if __name__ == '__main__':
    raise SystemExit(main())
