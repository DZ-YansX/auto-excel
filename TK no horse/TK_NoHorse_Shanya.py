import pandas as pd
import numpy as np
import os
import shutil
import warnings
import openpyxl
import time
import glob
import stat
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# 忽略警告
warnings.filterwarnings('ignore')

# 脚本所在目录（自动定位，无需手动修改）
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# ================= 配置区域 =================

# 1. 山亚数据文件夹（自动定位到脚本同级的 TKdeal_Shanya 子文件夹）
TARGET_FOLDER = os.path.join(SCRIPT_DIR, 'TKdeal_Shanya')

# 2. 唯一成本表文件名（位于 TKdeal_Shanya 文件夹内）
COST_FILE_NAME = '山亚成本价.xlsx'

# 3. 参数配置
UNIT_FEE = 3.5              # 发单费/空包费单价

# 3z. 美金汇率（广告费等以美元结算的费用统一按此汇率换算为 CNY）
USD_RATE = 7.0             # 1 USD = 7 CNY

# 3a. 多国汇率配置（脚本根据文件名关键词自动匹配，单位：1本国货币 = ? CNY）
#     若文件名不含以下关键词，则使用 DEFAULT_EXCHANGE_RATE
EXCHANGE_RATE_MAP = {
    '菲律宾': 0.10871,     # 1 PHP = 0.10871 CNY
    '印尼':   0.000375636,  # 1 IDR = 0.000375636 CNY
    '马来':   1.6373,      # 1 MYR = 1.6373 CNY
    '泰国':   0.2011,   # 1 THB = 0.2011 CNY
    '越南':   0.000028, # 1 VND ≈ 0.000028 CNY
    '新加坡': 5.28,     # 1 SGD ≈ 5.28 CNY
}
DEFAULT_EXCHANGE_RATE = 0.2011   # 未匹配时的默认汇率

# 3b. 货币代码 → 汇率（优先级高于文件名匹配；脚本会自动读取账单“货币”列识别）
CURRENCY_RATE_MAP = {
    'THB': 0.2011,
    'PHP': 0.10871,
    'IDR': 0.000375636,
    'MYR': 1.6373,
    'VND': 0.000028,
    'SGD': 5.28,
}

# 4. 输出文件夹（自动创建在脚本同级目录，不污染 TKdeal_Shanya 原始数据）
OUTPUT_FOLDER = os.path.join(SCRIPT_DIR, 'output_Shanya')

# 5. 时间筛选配置 (可随时修改)
#    TARGET_MONTH = 0 → 自动按销售明细日期识别月份；非 0 → 强制指定（如 5 表示“5月”）
#    利润表会写入对应“{月}月”列：已存在该列则更新，不存在则新增到首个空列，
#    从而保留之前几个月的数据（多月累计）。
TARGET_YEAR = 2026      # 目标年份
TARGET_MONTH = 0        # 目标月份 (0=自动识别, 1-12=强制指定)

# Sheet 名称映射
SHEET_NAMES = {
    'order':     ['订单明细表', '订单明细原表', '订单明细'],          # TK: 订单明细表
    'income':    ['账单原表', '收入明细表', '收入明细'],              # TK: 账单原表
    'brush':     ['补单明细', '补单'],
    'ads':       ['广告费', '广告'],
    'sales':     ['销售明细表', '销售明细'],                          # TK: 销售明细表
    'pnl':       ['利润表', 'Profit'],
    'unmatched': ['未匹配', '未匹配订单']
}

# ================= 1. 基础工具函数 =================

def find_header_row(df, keywords):
    for idx, row in df.head(20).iterrows():
        row_str = " ".join([str(x) for x in row.values])
        if all(k in row_str for k in keywords):
            return idx
    return 0

def load_and_clean_sheet(excel_file, sheet_candidates, keywords):
    sheet_name = None
    for name in sheet_candidates:
        if name in excel_file.sheet_names:
            sheet_name = name
            break
    if not sheet_name: return pd.DataFrame()
    try:
        df_preview = pd.read_excel(excel_file, sheet_name=sheet_name, header=None, nrows=20)
        header_idx = find_header_row(df_preview, keywords)
        df = pd.read_excel(excel_file, sheet_name=sheet_name, header=header_idx, dtype=str)
        return df
    except: return pd.DataFrame()

def clean_currency(x):
    if pd.isna(x): return 0.0
    s = str(x).replace('PHP', '').replace('RM', '').replace(',', '').replace(' ', '').strip()
    try: return float(s)
    except: return 0.0

def build_income_dict(df_income, amount_columns):
    """按订单汇总销售、退款及费用流水，保留其他字段的最后一条记录。"""
    result = df_income.drop_duplicates(subset='Order ID', keep='last').set_index('Order ID').to_dict('index')
    columns = list(dict.fromkeys(c for c in amount_columns if c in df_income.columns))
    if columns:
        amounts = df_income[columns].apply(lambda col: col.map(clean_currency))
        totals = amounts.groupby(df_income['Order ID']).sum()
        for oid, row in totals.iterrows():
            result[oid].update(row.to_dict())
    return result

def safe_border(ws, row_idx, col_idx):
    """安全地给单元格加边框"""
    thin = Side(border_style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    cell = ws.cell(row=row_idx, column=col_idx)
    cell.border = border
    cell.alignment = Alignment(horizontal='center', vertical='center')

# ================= 2. 核心处理逻辑 =================

def get_exchange_rate(file_name):
    """根据文件名关键词自动匹配汇率"""
    for keyword, rate in EXCHANGE_RATE_MAP.items():
        if keyword in file_name:
            return rate
    return DEFAULT_EXCHANGE_RATE

def copy_as_empty(source_path, file_name):
    """当无有效数据时，仅复制原表并增加前缀作为结果"""
    new_file_name = f"新_{file_name}"
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    file_output = os.path.join(OUTPUT_FOLDER, new_file_name)
    try:
        shutil.copy(source_path, file_output)
        os.chmod(file_output, stat.S_IWRITE)
        wb = openpyxl.load_workbook(file_output)
        for sheet_name in ('销售明细1', '销售明细表1'):
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]
        wb.save(file_output)
    except Exception as e:
        print(f"   [错误] 空文件生成副本失败: {e}")

def process_single_file(source_path, cost_map):
    file_name = os.path.basename(source_path)
    EXCHANGE_RATE = get_exchange_rate(file_name)  # 按文件名自动选汇率
    print(f"\n{'='*60}")
    print(f">>> 正在处理: {file_name}  (汇率: {EXCHANGE_RATE})")

    # 1. 读取数据
    xls = pd.ExcelFile(source_path)
    # TK 订单关键词: 'Order ID' 和 'Seller SKU' 均存在
    df_order = load_and_clean_sheet(xls, SHEET_NAMES['order'], ['Order ID', 'Seller SKU'])
    # TK 账单关键词: 'Order/adjustment ID' 和 'Total settlement amount'
    df_income = load_and_clean_sheet(xls, SHEET_NAMES['income'],
                                     ['Order/adjustment ID', 'Total settlement amount'])
    # 兼容 Shopee 收入表
    if df_income.empty:
        df_income = load_and_clean_sheet(xls, SHEET_NAMES['income'], ['Order ID', 'Total Released Amount'])
    df_brush = load_and_clean_sheet(xls, SHEET_NAMES['brush'], ['线上单号', '佣金'])
    df_ads = load_and_clean_sheet(xls, SHEET_NAMES['ads'], ['花费'])

    if df_order.empty or df_income.empty:
        print("   [跳过] 缺少订单或收入明细表，正在直接生成空原样副本")
        copy_as_empty(source_path, file_name)
        return

    # 将账单的 'Order/adjustment ID  '(含空格) 标准化为 'Order ID'
    inc_id_col = next((c for c in df_income.columns if 'Order' in c and ('ID' in c or 'adjustment' in c)), None)
    # TK 中文账单：订单ID列名为“订单ID/调整单ID”
    if inc_id_col is None:
        inc_id_col = next((c for c in df_income.columns if '订单ID' in c.replace(' ', '')), None)
    if inc_id_col and inc_id_col != 'Order ID':
        df_income = df_income.rename(columns={inc_id_col: 'Order ID'})
    # 去掉非订单行（如 adjustment、refund 等）——保留 Type==Order / 交易类型==订单 的行
    if 'Type ' in df_income.columns or 'Type' in df_income.columns:
        type_col = 'Type ' if 'Type ' in df_income.columns else 'Type'
        df_income = df_income[df_income[type_col].astype(str).str.strip() == 'Order'].copy()
    else:
        zh_type_col = next((c for c in df_income.columns if c.strip() == '交易类型'), None)
        if zh_type_col:
            df_income = df_income[df_income[zh_type_col].astype(str).str.strip() == '订单'].copy()

    # 根据账单“货币”列自动识别汇率（优先级高于文件名匹配）
    cur_col = next((c for c in df_income.columns if c.strip() in ('货币', 'Currency')), None)
    if cur_col:
        cur_vals = df_income[cur_col].dropna()
        if len(cur_vals):
            cur_code = str(cur_vals.iloc[0]).strip().upper()
            if cur_code in CURRENCY_RATE_MAP:
                EXCHANGE_RATE = CURRENCY_RATE_MAP[cur_code]
                print(f"   按账单货币 {cur_code} 自动选用汇率: {EXCHANGE_RATE}")

    # 最终校验：若 Order ID 仍不存在（表头识别失败），视为无账单数据，跳过
    if 'Order ID' not in df_income.columns:
        print(f"   [警告] 账单原表表头识别失败（实际列: {list(df_income.columns[:4])}），跳过并生成原样副本")
        copy_as_empty(source_path, file_name)
        return

    # 2. 数据计算 — 日期列定位（TK: 'Created Time', Shopee: 'Order Creation Date'/'Order created time'）
    date_col = next((c for c in df_order.columns if c in ('Created Time', 'Order Creation Date', 'Order created time')), df_order.columns[0])

    # TK 订单明细的时间为 DD/MM/YYYY 格式，必须 dayfirst=True，否则月份会错位
    df_order['Created Time'] = pd.to_datetime(df_order[date_col], dayfirst=True, errors='coerce')

    # 目标月份：TARGET_MONTH==0 时按订单日期自动识别（取目标年份内出现最多的月份）
    if TARGET_MONTH == 0:
        yr_months = df_order[df_order['Created Time'].dt.year == TARGET_YEAR]['Created Time'].dt.month.dropna()
        if len(yr_months) == 0:
            print(f"   [跳过] {TARGET_YEAR}年内无可识别的订单日期，无法自动确定月份，生成空原样副本")
            copy_as_empty(source_path, file_name)
            return
        current_month = int(yr_months.mode().iloc[0])
        print(f"   自动识别月份: {current_month}月（按 {TARGET_YEAR} 年订单日期众数）")
    else:
        current_month = TARGET_MONTH

    mask_order = (df_order['Created Time'].dt.year == TARGET_YEAR) & (df_order['Created Time'].dt.month == current_month)
    df_order_dec = df_order[mask_order].copy()
    print(f"   {current_month}月订单数: {len(df_order_dec)}")

    if len(df_order_dec) == 0:
        print(f"   [跳过] {current_month}月无订单，将直接生成空原样副本")
        copy_as_empty(source_path, file_name)
        return

    # 收入明细时间筛选（TK 账单: 'Order created time'，Shopee: 'Order Creation Date'）
    inc_date_col = next((c for c in df_income.columns
                         if c.strip() in ('Order created time', 'Order Creation Date', '订单创建时间')), None)
    if inc_date_col:
        df_income['Created Time'] = pd.to_datetime(df_income[inc_date_col], errors='coerce')
        mask_income = (df_income['Created Time'].dt.year == TARGET_YEAR) & (df_income['Created Time'].dt.month == current_month)
        df_income_dec = df_income[mask_income].copy()
    else:
        df_income_dec = df_income.copy()  # 无时间列则全量使用


    sku_col = next((c for c in df_order_dec.columns if 'Seller SKU' in c), None)
    if not sku_col: sku_col = next((c for c in df_order_dec.columns if 'Reference' in c and 'Parent' not in c), None)
    if not sku_col: sku_col = next((c for c in df_order_dec.columns if 'SKU' in c), 'Seller SKU')
    qty_col = 'Quantity'
    # TK 物流单号: 'Tracking ID'， Shopee: 'Tracking Number*'
    track_col = next((c for c in df_order_dec.columns if 'Tracking' in c), 'Tracking ID')
    
    def resolve_sku_cost(sku, cmap):
        """解析单个 SKU 的单位成本：
        成本表只收录基础 SKU；组合装 SKU 形如 “基础SKU*N”，成本 = 基础成本 * N。"""
        sku = str(sku).strip()
        if sku in cmap:
            return cmap[sku]
        if '*' in sku:
            base, _, mult = sku.rpartition('*')
            base = base.strip()
            try:
                n = float(mult.strip())
            except ValueError:
                n = None
            if base in cmap and n is not None:
                return cmap[base] * n
        return 0

    def get_row_cost(row, cmap):
        qty = clean_currency(row[qty_col])
        return resolve_sku_cost(row[sku_col], cmap) * qty

    df_order_dec['Line_Total_Cost'] = df_order_dec.apply(get_row_cost, axis=1, cmap=cost_map)
    track_col = next((c for c in df_order_dec.columns if 'Tracking' in c), 'Tracking Number*')
    status_col = 'Order Status'

    order_agg = df_order_dec.groupby('Order ID').agg({
        'Created Time': 'first',
        sku_col: lambda x: '、'.join([str(i) for i in x if str(i)!='nan']),
        qty_col: lambda x: sum([clean_currency(i) for i in x]),
        'Line_Total_Cost': 'sum',
        status_col: 'first',
        track_col: 'first'
    }).reset_index()
    order_agg.rename(columns={sku_col: 'Merged_SKU', 'Line_Total_Cost': 'Total_Cost_CNY'}, inplace=True)

    cols = df_income.columns
    # TK 账单列名， Shopee 列名兼容
    c_price   = next((c for c in cols if 'Subtotal before discounts' in c
                      or 'Original Product Price' in c), None)
    c_promo   = next((c for c in cols if 'Seller discounts' in c
                      and 'co-funded' not in c and 'Refund' not in c
                      or 'Seller Product Promotion' in c), None)
    c_voucher = next((c for c in cols if 'Seller co-funded voucher discount' in c
                      and 'Refund' not in c
                      or 'Voucher Sponsored by Seller' in c), None)
    c_refund  = next((c for c in cols if c.strip() == 'Customer refund'
                      or 'Refund Amount' in c), None)
    # TK 中文账单列名兼容：销售单价/优惠金额/退款金额
    #   销售单价 = 享受折扣前小计；优惠金额 = 商家折扣（不含优惠劵）；
    #   退款金额 = 享受商家折扣后的退款小计（注意与“客户退款”不同）
    if c_price is None:
        c_price = next((c for c in cols if c.strip() == '享受折扣前小计'), None)
    if c_promo is None:
        c_promo = next((c for c in cols if c.strip() == '商家折扣'), None)
    if c_refund is None:
        c_refund = next((c for c in cols if c.strip() == '享受商家折扣后的退款小计'), None)
    # 同一订单可能分多行结算、退款、补扣费用，金额必须累加，不能仅取最后一行。
    if 'Order ID' in df_income.columns:
        amount_columns = [c_price, c_promo, c_voucher, c_refund]
        amount_columns += [c for c in cols if c.strip() == '总费用'
                           or any(word in c.lower() for word in ('fee', 'tax', 'commission', 'rebate'))]
        income_dict = build_income_dict(df_income, amount_columns)
    else:
        print(f"   [警告] 账单原表未找到 Order ID 列（实际列名: {list(df_income.columns[:5])}），收入金额将为 0")
        income_dict = {}
    
    sale_details = []
    for _, row in order_agg.iterrows():
        oid = row['Order ID']
        inc_row = income_dict.get(oid, {})
        sale_details.append({
            '日期': row['Created Time'],
            '订单号': oid,
            '商品SKU': row['Merged_SKU'],
            '销售数量': row[qty_col],
            '销售单价': clean_currency(inc_row.get(c_price, 0)),
            '优惠金额': clean_currency(inc_row.get(c_promo, 0)),
            '优惠劵金额': clean_currency(inc_row.get(c_voucher, 0)),
            '退款金额': clean_currency(inc_row.get(c_refund, 0)),
            '商品成本总额(CNY)': row['Total_Cost_CNY'],
            'Status': row[status_col],
            'Tracking': row[track_col]
        })
    df_sales_detail = pd.DataFrame(sale_details)

    df_sales_detail = df_sales_detail[df_sales_detail['销售单价'] > 0.001].copy()
    df_sales_detail.reset_index(drop=True, inplace=True)

    ids_order = set(df_order_dec['Order ID'])
    ids_income = set(df_income['Order ID'])
    unmatched_ids = ids_order - ids_income
    df_unmatched = df_order_dec[df_order_dec['Order ID'].isin(unmatched_ids)].copy()
    
    count_cancel_shipped = 0
    count_cancel_unshipped = 0
    cancel_shipped_cost = 0.0
    ship_time_col = next((c for c in df_order_dec.columns if c.strip() in ('Shipped Time', 'Ship Time')), None)
    if not df_unmatched.empty:
        def check_status(row):
            status = str(row.get(status_col, ''))
            # 取消订单：以是否“已发货(Shipped Time 非空)”判断是否已发货
            if 'Cancel' in status or '取消' in status:
                ship = str(row.get(ship_time_col, '')) if ship_time_col else ''
                shipped = ship.strip() not in ('', 'nan', 'NaT', 'None')
                return 'Shipped' if shipped else 'Unshipped'
            return 'Other'
        df_unmatched['Check_Status'] = df_unmatched.apply(check_status, axis=1)
        count_cancel_shipped = df_unmatched[df_unmatched['Check_Status'] == 'Shipped']['Order ID'].nunique()
        count_cancel_unshipped = df_unmatched[df_unmatched['Check_Status'] == 'Unshipped']['Order ID'].nunique()
        # 已发货却被取消的订单货物已寄出，其商品成本计入实际产品成本
        cancel_shipped_cost = df_unmatched[df_unmatched['Check_Status'] == 'Shipped']['Line_Total_Cost'].sum()

    brushing_ids = set()
    if not df_brush.empty:
        brush_id_col = next((c for c in df_brush.columns if '单号' in c), df_brush.columns[0])
        brushing_ids = set(df_brush[brush_id_col].astype(str).str.strip())
    
    c_rebate = next((c for c in cols if 'Rebate' in c and 'Shopee' in c), None)
    cols_brush_amt = [c for c in [c_price, c_promo, c_voucher, c_refund, c_rebate] if c]
    
    all_cols = list(df_income.columns)
    try:
        # TK: 平台费用用 'Total Fees' / 中文“总费用” 一列表示； Shopee: 范围列汇总
        zh_total_fee = next((c for c in all_cols if c.strip() == '总费用'), None)
        if 'Total Fees' in all_cols:
            cols_platform_fee_range = ['Total Fees']
        elif zh_total_fee:
            cols_platform_fee_range = [zh_total_fee]
        else:
            idx_start = next(i for i, c in enumerate(all_cols) if 'Seller shipping fee' in c or 'Shipping Fee Paid by Seller' in c)
            idx_end   = next(i for i, c in enumerate(all_cols) if 'Withholding Tax' in c)
            cols_platform_fee_range = all_cols[idx_start : idx_end+1]
    except:
        cols_platform_fee_range = [c for c in all_cols if 'Fee' in c or 'Tax' in c or 'Commission' in c]
    


    brush_refund_sum_php = 0
    brush_platform_fee_php = 0
    for bid in brushing_ids:
        if bid in income_dict:
            row = income_dict[bid]
            for c in cols_brush_amt: brush_refund_sum_php += clean_currency(row.get(c, 0))
            for c in cols_platform_fee_range: brush_platform_fee_php += clean_currency(row.get(c, 0))

    cols_pure_platform = [c for c in cols_platform_fee_range if 'Tax' not in c]
    total_platform_fee_php = sum(df_income_dec[c].apply(clean_currency).sum() for c in cols_pure_platform if c in df_income_dec.columns)
    c_tax = next((c for c in all_cols if 'Withholding Tax' in c), None)
    total_tax_php = df_income_dec[c_tax].apply(clean_currency).sum() if c_tax else 0

    # 广告费以美元结算，按 USD_RATE 换算为 CNY（列名可能是“花费”“Expense”或直接是“广告费”）
    ads_col = next((c for c in df_ads.columns if '花费' in c or 'Expense' in c or '广告' in str(c)), None)
    ads_cny = df_ads[ads_col].apply(clean_currency).sum() * USD_RATE if ads_col else 0
    
    brush_comm = 0; brush_diff = 0
    if not df_brush.empty:
        c_comm = next((c for c in df_brush.columns if '佣金' in c), None)
        c_diff = next((c for c in df_brush.columns if '汇差' in c), None)
        if c_comm: brush_comm = df_brush[c_comm].apply(clean_currency).sum()
        if c_diff: brush_diff = df_brush[c_diff].apply(clean_currency).sum()

    metrics = {
        'total_platform_fee_php': total_platform_fee_php,
        'total_tax_php': total_tax_php,
        'brush_amt_php': brush_refund_sum_php,
        'brush_platform_fee_php': brush_platform_fee_php,
        'brush_commission': brush_comm,
        'brush_diff': brush_diff,
        'brush_empty_pkg': len(brushing_ids) * UNIT_FEE,
        'ads_cny': ads_cny,
        'cancel_shipped': count_cancel_shipped,
        'cancel_unshipped': count_cancel_unshipped,
        'cancel_shipped_cost': cancel_shipped_cost,
        'order_count': len(df_sales_detail)
    }

    # 销售明细按日期升序排列（与人工整理顺序一致）
    df_sales_detail = df_sales_detail.sort_values('日期').reset_index(drop=True)

    # ================= 3. 写入文件 =================
    
    new_file_name = f"新_{file_name}"
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)  # 确保输出目录存在
    file_output = os.path.join(OUTPUT_FOLDER, new_file_name)

    print(f"   正在生成副本并进行写表操作...")
    try:
        shutil.copy(source_path, file_output)
        os.chmod(file_output, stat.S_IWRITE)
        time.sleep(1)
    except Exception as e:
        print(f"   [错误] 复制失败: {e}")
        return

    try:
        wb = openpyxl.load_workbook(file_output)
        
        # 删除模板中不再需要的第二张销售明细
        for sheet_name in ('销售明细1', '销售明细表1'):
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]

        # --- A. 处理“销售明细” ---
        curr_sym = next((sym for key, sym in {'菲律宾':'₱', '印尼':'Rp', '马来':'RM', '泰国':'฿', '越南':'₫', '新加坡':'S$'}.items() if key in file_name), '外币')

        def write_sales_sheet(sheet_candidates, default_name, cost_field):
            """使用唯一成本表写入销售明细 sheet。"""
            sheet_name = next((n for n in sheet_candidates if n in wb.sheetnames), None)
            if not sheet_name:
                ws_sales = wb.create_sheet(default_name)
                headers = ['日期', '订单号', '商品SKU', '销售数量', f'销售单价({curr_sym})', f'优惠金额({curr_sym})', f'退款金额({curr_sym})',
                           f'销售金额({curr_sym})', '销售金额(¥)', '商品成本总额(¥)', '毛利(¥)']
                ws_sales.append(headers)
                sheet_name = default_name
            else:
                ws_sales = wb[sheet_name]
                max_row = ws_sales.max_row
                # 【修复点 1】：放弃极其耗时的单元格迭代清空，直接使用 delete_rows 瞬间删除旧的几十万行幽灵数据
                if max_row >= 2:
                    ws_sales.delete_rows(2, max_row)

            for i, row in df_sales_detail.iterrows():
                r = i + 2
                try:
                    ws_sales.cell(r, 1).value = pd.Timestamp(row['日期']).strftime('%Y/%m/%d')
                except Exception:
                    ws_sales.cell(r, 1).value = str(row['日期'])[:10]
                ws_sales.cell(r, 2).value = row['订单号']
                ws_sales.cell(r, 3).value = row['商品SKU']
                ws_sales.cell(r, 4).value = row['销售数量']
                ws_sales.cell(r, 5).value = row['销售单价']
                ws_sales.cell(r, 6).value = row['优惠金额'] + row['优惠劵金额']
                ws_sales.cell(r, 7).value = row['退款金额']

                ws_sales.cell(r, 8).value = f'=E{r}+F{r}+G{r}'
                ws_sales.cell(r, 9).value = f'=H{r}*{EXCHANGE_RATE}'
                ws_sales.cell(r, 10).value = row[cost_field]
                ws_sales.cell(r, 11).value = f'=I{r}-J{r}'

                for c in range(1, 12): safe_border(ws_sales, r, c)

            last_row = len(df_sales_detail) + 1

            # 末尾追加合计行（E~K 列求和）
            if len(df_sales_detail) > 0:
                sum_r = last_row + 1
                for c in range(5, 12):
                    col_let = get_column_letter(c)
                    ws_sales.cell(sum_r, c).value = f'=SUM({col_let}2:{col_let}{last_row})'
                    safe_border(ws_sales, sum_r, c)
            return sheet_name, last_row

        sales_sheet_name, last_sales_row = write_sales_sheet(SHEET_NAMES['sales'], '销售明细表', '商品成本总额(CNY)')

        # --- B. 处理“补单明细” ---
        brush_sheet_name = None
        for n in SHEET_NAMES['brush']:
            if n in wb.sheetnames: brush_sheet_name = n; break

        v_brush_amt = metrics['brush_amt_php'] * EXCHANGE_RATE
        v_shopee_fee = abs(metrics['total_platform_fee_php']) * EXCHANGE_RATE 
        v_tax = abs(metrics['total_tax_php']) * EXCHANGE_RATE                  
        v_brush_plat = abs(metrics['brush_platform_fee_php']) * EXCHANGE_RATE
        
        brush_commission = metrics['brush_commission']
        brush_empty_pkg = metrics['brush_empty_pkg']
        brush_diff_expr = metrics['brush_diff']

        if brush_sheet_name and not df_brush.empty:
            ws_brush = wb[brush_sheet_name]
            
            header_row = None
            col_idx_xiaoji = None
            col_idx_chengben = None
            col_idx_huicha = None

            for row in ws_brush.iter_rows(min_row=1, max_row=20):
                for cell in row:
                    if cell.value is not None:
                        val_str = str(cell.value).strip()
                        if '垫付小计' in val_str: col_idx_xiaoji = cell.column
                        elif '垫付成本' in val_str: col_idx_chengben = cell.column
                        elif '汇率差' in val_str or '汇差' in val_str: col_idx_huicha = cell.column
                if col_idx_xiaoji and col_idx_chengben:
                    header_row = row[0].row
                    break

            if header_row and col_idx_xiaoji and col_idx_chengben:
                if not col_idx_huicha:
                    col_idx_huicha = ws_brush.max_column + 1
                    ws_brush.cell(row=header_row, column=col_idx_huicha, value="汇率差")

                col_let_xiaoji = get_column_letter(col_idx_xiaoji)
                col_let_chengben = get_column_letter(col_idx_chengben)
                col_let_huicha = get_column_letter(col_idx_huicha)

                max_brush_row = ws_brush.max_row
                for r in range(header_row + 1, max_brush_row + 1):
                    if ws_brush.cell(row=r, column=col_idx_xiaoji).value is not None or ws_brush.cell(row=r, column=col_idx_chengben).value is not None:
                        ws_brush.cell(row=r, column=col_idx_huicha).value = f'={col_let_xiaoji}{r}-{EXCHANGE_RATE}*{col_let_chengben}{r}'
                
                brush_diff_expr = f'=SUM({brush_sheet_name}!{col_let_huicha}{header_row+1}:{col_let_huicha}{max_brush_row})'
        else:
            v_brush_amt = 0
            brush_commission = 0
            v_brush_plat = 0
            brush_diff_expr = 0
            brush_empty_pkg = 0

        # --- C. 处理“利润表” ---
        pnl_sheet_name = None
        for n in SHEET_NAMES['pnl']:
            if n in wb.sheetnames: pnl_sheet_name = n; break
        
        if pnl_sheet_name:
            ws_pnl = wb[pnl_sheet_name]
            
            # 定位填报的列：在“项目”行寻找“{当前月}月”列。
            #   已存在该月列 → 复用（更新本月数据）；
            #   不存在 → 取该行首个空列新增“{当前月}月”表头（保留之前几个月的数据，多月累计）。
            month_label = f"{current_month}月"
            target_col_idx = 2  # 默认在第二列(B列)
            proj_row = None
            for r_idx, row in enumerate(ws_pnl.iter_rows(min_row=1, max_row=5), start=1):
                row_vals = [str(c.value).strip() if c.value else "" for c in row]
                if any("项目" in val for val in row_vals):
                    proj_row = r_idx
                    proj_col = next((c.column for c in row if c.value and '项目' in str(c.value)), 1)
                    # 1) 优先复用已存在的“{当前月}月”列
                    found = None
                    for cell in row:
                        if cell.value and str(cell.value).replace(' ', '').strip() == month_label:
                            found = cell.column
                            break
                    if found is not None:
                        target_col_idx = found
                    else:
                        # 2) 否则取“项目”列右侧首个空列，写入“{当前月}月”表头
                        scan = proj_col + 1
                        while ws_pnl.cell(proj_row, scan).value not in (None, ''):
                            scan += 1
                        target_col_idx = scan
                        ws_pnl.cell(proj_row, target_col_idx).value = month_label
                    break
            target_col_letter = get_column_letter(target_col_idx)

            sname = sales_sheet_name or "销售明细表"
            cancel_cost = round(metrics['cancel_shipped_cost'], 4)
            # 叶子项（直接填数值/公式）；汇总项（实际主营收入/成本/营业费用/结算/利润等）保留模板原有公式
            val_map = {
                # 本期销售 = 销售金额(¥)合计 + 销售退款（即未扣退款前的销售额）
                '本期销售': f'=SUM({sname}!I2:I{last_sales_row})-SUM({sname}!G2:G{last_sales_row})*{EXCHANGE_RATE}',
                '其中:补单金额': v_brush_amt,
                '销售退款': f'=ABS(SUM({sname}!G2:G{last_sales_row}))*{EXCHANGE_RATE}',
                # 实际产品成本 = 销售明细成本合计 + 已发货被取消订单的商品成本
                '实际产品成本': f'=SUM({sname}!J2:J{last_sales_row})+{cancel_cost}',
                '发单费': f'=({metrics["order_count"]}+{metrics["cancel_shipped"]})*{UNIT_FEE}',
                '广告费': metrics['ads_cny'],
                '补单佣金': brush_commission,
                '补单平台费用': v_brush_plat,
                '补单汇率差额': brush_diff_expr,
                '补单空包费': brush_empty_pkg,
                '平台费用': v_shopee_fee,
                '调整金额': 0, # 新增或可选预留
                '预扣税': v_tax,
                '团队运营费用': 0,
            }

            row_idx_map = {}
            note_written = False
            for row in ws_pnl.iter_rows(min_row=1, max_row=100): # 限制读取范围防卡死
                cell_a = row[0]
                cell_val_str = str(cell_a.value).strip() if cell_a.value else ""

                for key in ['实际主营收入', '本期销售', '其中:补单金额', '销售退款',
                            '实际主营成本', '实际产品成本', '发单费',
                            '营业费用', '推广费用', '广告费', '补单佣金', '补单平台费用', '补单汇率差额', '补单空包费',
                            'TK结算', 'shopee结算', '平台费用', '调整金额', '预扣税', '团队运营费用', '毛利润', '本期利润']:
                    if key in cell_val_str:
                        row_idx_map[key] = cell_a.row
                        if key in val_map:
                            row[target_col_idx - 1].value = val_map[key]
                        # 每个项目行只允许匹配一个项目；避免“补单平台费用”再次误匹配“平台费用”
                        break

                if "未匹配" in cell_val_str or "补充" in cell_val_str:
                    cell_a.value = f"{current_month}月未匹配：取消未发货{metrics['cancel_unshipped']}，取消已发货{metrics['cancel_shipped']}"
                    note_written = True

            def get_cell(key):
                r = row_idx_map.get(key)
                return f"{target_col_letter}{r}" if r else "0"

            def _is_empty(addr):
                v = ws_pnl[addr].value
                return v is None or (isinstance(v, str) and v.strip() == "")

            # 汇总项仅在模板未预置公式时才补写，避免覆盖人工表已有的合计公式
            if '实际主营收入' in row_idx_map and _is_empty(get_cell('实际主营收入')):
                ws_pnl[get_cell('实际主营收入')] = f"={get_cell('本期销售')}-{get_cell('其中:补单金额')}-{get_cell('销售退款')}"

            if '实际主营成本' in row_idx_map and _is_empty(get_cell('实际主营成本')):
                ws_pnl[get_cell('实际主营成本')] = f"={get_cell('实际产品成本')}+{get_cell('发单费')}"

            if '推广费用' in row_idx_map and _is_empty(get_cell('推广费用')):
                sub_items = [get_cell(k) for k in ['广告费', '补单佣金', '补单平台费用', '补单汇率差额', '补单空包费'] if k in row_idx_map]
                ws_pnl[get_cell('推广费用')] = f"={'+'.join(sub_items)}"

            # 兼容：如果有shopee结算则用shopee结算，如果是TK结算则用TK结算
            settle_key = 'TK结算' if 'TK结算' in row_idx_map else ('shopee结算' if 'shopee结算' in row_idx_map else None)
            if settle_key and _is_empty(get_cell(settle_key)):
                settle_items = [get_cell(k) for k in ['平台费用', '调整金额', '预扣税'] if k in row_idx_map]
                if settle_items:
                    ws_pnl[get_cell(settle_key)] = f"={'+'.join(settle_items)}"

            if '营业费用' in row_idx_map and _is_empty(get_cell('营业费用')):
                biz_fee_items = [get_cell('推广费用'), get_cell('团队运营费用')]
                if settle_key: biz_fee_items.append(get_cell(settle_key))
                ws_pnl[get_cell('营业费用')] = f"={'+'.join(biz_fee_items)}"

            profit_key = '本期利润' if '本期利润' in row_idx_map else '毛利润'
            if profit_key in row_idx_map and _is_empty(get_cell(profit_key)):
                ws_pnl[get_cell(profit_key)] = f"={get_cell('实际主营收入')}-{get_cell('实际主营成本')}-{get_cell('营业费用')}"

            # 若模板没有专门的“未匹配”说明行，则在利润表末尾补一行说明
            if not note_written:
                anchor = row_idx_map.get(profit_key) or (max(row_idx_map.values()) if row_idx_map else ws_pnl.max_row)
                ws_pnl.cell(anchor + 1, 1).value = f"{current_month}月未匹配：取消未发货{metrics['cancel_unshipped']}，取消已发货{metrics['cancel_shipped']}"

        # --- D. 处理“未匹配” ---
        un_sheet_name = None
        for n in SHEET_NAMES['unmatched']:
            if n in wb.sheetnames: un_sheet_name = n; break
            
        if un_sheet_name:
            ws_un = wb[un_sheet_name]
            # 无论本期是否存在未匹配订单，都先清除模板中的历史内容，避免旧数据残留
            if ws_un.max_row >= 1:
                ws_un.delete_rows(1, ws_un.max_row)
        elif not df_unmatched.empty:
            ws_un = wb.create_sheet('未匹配订单')

        if not df_unmatched.empty:
                
            # 只保留原始订单列，去掉脚本计算用的辅助列
            df_un_out = df_unmatched.drop(columns=[c for c in ['Line_Total_Cost', 'Check_Status'] if c in df_unmatched.columns])
            headers = df_un_out.columns.tolist()
            ws_un.append(headers)
            for r in dataframe_to_rows(df_un_out, index=False, header=False):
                ws_un.append([str(x) if isinstance(x, pd.Timestamp) else x for x in r])

        wb.save(file_output)
        wb.close()
        print("   [成功] 处理完成")

    except Exception as e:
        print(f"   [失败] 处理出错: {e}")

# ================= 3. 主入口 =================

def main():
    print(f"启动批量处理...")
    print(f"  数据来源: {TARGET_FOLDER}")
    print(f"  输出目录: {OUTPUT_FOLDER}")
    month_desc = "自动识别(按订单日期)" if TARGET_MONTH == 0 else f"{TARGET_MONTH}月"
    print(f"  当前设置分析时间: {TARGET_YEAR}年 {month_desc}")
    
    if not os.path.exists(TARGET_FOLDER):
        print(f"错误: 找不到 TKdeal_Shanya 文件夹 '{TARGET_FOLDER}'")
        print(f"  请确认脚本与 TKdeal_Shanya 文件夹在同一目录下。")
        input("按回车键退出...")
        return

    def load_cost_map(cost_path):
        df_cost = pd.read_excel(cost_path, dtype=str)
        sku_col = next((c for c in df_cost.columns if 'sku' in c.lower()), 'SKU')
        cost_col = next((c for c in df_cost.columns if '成本' in c), '人民币成本')
        return dict(zip(df_cost[sku_col].astype(str).str.strip(), df_cost[cost_col].apply(clean_currency)))

    # 唯一成本表从 TKdeal_Shanya 文件夹内读取
    cost_path = os.path.join(TARGET_FOLDER, COST_FILE_NAME)
    if not os.path.exists(cost_path):
        print(f"错误: 找不到成本表 '{cost_path}'")
        print(f"  请将 '{COST_FILE_NAME}' 放在 TKdeal_Shanya 文件夹内: {TARGET_FOLDER}")
        input("按回车键退出...")
        return

    print(f">>> 加载商品成本表: {cost_path}")
    try:
        cost_map = load_cost_map(cost_path)
        print(f"   已加载 {len(cost_map)} 条 SKU 成本记录")
    except Exception as e:
        print(f"   读取成本表失败: {e}")
        input("按回车键退出...")
        return

    # 搜索 TKdeal_Shanya 内所有 xlsx，跳过成本表和临时文件
    xlsx_files = sorted(glob.glob(os.path.join(TARGET_FOLDER, "*.xlsx")))
    count = 0
    for file_path in xlsx_files:
        file_name = os.path.basename(file_path)
        if file_name == COST_FILE_NAME or file_name.startswith("~$"):  # 跳过成本表和临时锁定文件
            continue
        process_single_file(file_path, cost_map)
        count += 1

    print(f"\n{'='*60}")
    print(f"批量处理完成！共处理了 {count} 个文件。")
    print(f"新生成的文件存放在: {OUTPUT_FOLDER}")
    input("按回车键退出...")

if __name__ == '__main__':
    main()
