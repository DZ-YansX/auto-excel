# -*- coding: utf-8 -*-
"""
批量补齐当前目录中所有大卖家表的“商品成本”和“总成本价”。

处理口径：
1. 以“店铺SKU”为匹配键；商品成本按店铺 SKU 逐行对应，格式为两位小数并以 CRLF 换行。
2. 除基础 SKU 为 EB-S 外，0 成本视为缺失并重新查价；EB-S 成本固定允许为 0。
3. 空白、--、非数值或与 SKU 子行数不一致的商品成本需要重建：
   先从本表同国家、同基础 SKU 的已有成本反推单价，并取正成本众数；
   本表没有时再查“万山成本价.xlsx”；YY/XY/THS/WSS 精确查不到时按对应 *-KP 兜底。
4. 总成本价等于商品成本各子行之和，不再重复乘 SKU 倍数或销售数量。
5. 只有真实商品 SKU 在两处都找不到成本时才标黄；空 SKU 和 -- 占位符不标黄。

用法：
    python 补成本.py
    python 补成本.py 大卖家表1.xlsx 大卖家表2.xlsx 万山成本价.xlsx

输出：
    <原文件名>_已填成本.xlsx
"""

import os
import re
import sys
import math
from collections import Counter, defaultdict
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import COLOR_INDEX


try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


COUNTRY_KEYWORDS = [
    ("印度尼西亚", "印尼"),
    ("马来西亚", "马来"),
    ("印尼", "印尼"),
    ("马来", "马来"),
    ("菲律宾", "菲律宾"),
    ("泰国", "泰国"),
    ("越南", "越南"),
    ("新加坡", "新加坡"),
]

# 这两个店铺名没有国家字样，但表内 YY 系列单位成本均为印尼口径。
SHOP_COUNTRY_OVERRIDES = {"保健1店": "印尼", "保健7店": "印尼"}
WAREHOUSE_ALIASES = {"印度尼西亚": "印尼", "马来西亚": "马来"}

MULT_RE = re.compile(r"^(.*?)\s*[*＊×]\s*(\d+)\s*$")
MULT_ONLY_RE = re.compile(r"^[*＊×]\s*\d+\s*$")
MISSING_MARKERS = {"", "-", "--"}
KP_FALLBACK_PREFIXES = {"YY", "XY", "THS", "WSS"}

COST_COL_CANDIDATES = {
    "sku": ["库存SKU编号", "商品SKU", "SKU编号", "SKU"],
    "wh": ["仓库", "国家", "站点"],
    "price": ["成本单价(￥)", "成本单价", "成本价", "单价", "成本"],
}
COST_COL_FALLBACK = {"sku": 1, "wh": 2, "price": 3}

FILL_COL_CANDIDATES = {
    "order": ["订单号", "订单编号"],
    "shop": ["BigSeller店铺昵称", "BigSeller店铺", "店铺昵称", "店铺"],
    "store_sku": ["店铺SKU"],
    "qty": ["实际销售数量", "销售数量", "数量"],
    "cost": ["商品成本"],
    "total": ["总成本价"],
}
FILL_COL_FALLBACK = {
    "order": 1,
    "shop": 3,
    "store_sku": 11,
    "qty": 13,
    "cost": 25,
    "total": 26,
}

YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
NO_FILL = PatternFill()


def norm(value):
    return "" if value is None else str(value).strip()


def sku_key(value):
    # 仅清理外围空白并忽略大小写；SKU 内部空格保留。
    return norm(value).upper()


def warehouse_key(value):
    text = norm(value)
    return WAREHOUSE_ALIASES.get(text, text)


def first_sheet(workbook):
    if "Sheet1" in workbook.sheetnames:
        return workbook["Sheet1"]
    if "sheet1" in workbook.sheetnames:
        return workbook["sheet1"]
    return workbook.worksheets[0]


def header_row(worksheet):
    return [
        (column, norm(worksheet.cell(row=1, column=column).value))
        for column in range(1, worksheet.max_column + 1)
    ]


def find_col(headers, candidates, fallback):
    for candidate in candidates:
        for index, text in headers:
            if text == candidate:
                return index
    for candidate in candidates:
        for index, text in headers:
            if text and candidate in text:
                return index
    return fallback


def split_lines(value):
    if value is None:
        return []
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return [part.strip() for part in text.split("\n") if part.strip()]


def split_store_skus(value):
    """拆分店铺SKU；将断行的 *N 和以 + 开头的套装续行并回前一项。"""
    result = []
    for part in split_lines(value):
        if MULT_ONLY_RE.fullmatch(part) and result:
            result[-1] += re.sub(r"\s+", "", part)
        elif part.startswith("+") and part[1:].strip() and result:
            result[-1] += part
        else:
            result.append(part)
    return result


def quantity_lines(value, sku_count):
    values = number_lines(value)
    # 套装店铺SKU可能拆成多个组件，但订单数量只导出一行；同一数量作用于各组件。
    if sku_count > 1 and len(values) == 1:
        return values * sku_count
    return values


def parse_sku_line(line):
    match = MULT_RE.match(norm(line))
    if match:
        return match.group(1).strip(), int(match.group(2))
    return norm(line), 1


def country_of(shop):
    text = norm(shop)
    if text in SHOP_COUNTRY_OVERRIDES:
        return SHOP_COUNTRY_OVERRIDES[text]
    for keyword, warehouse in COUNTRY_KEYWORDS:
        if keyword in text:
            return warehouse
    return None


def to_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    try:
        text = str(value).replace(",", "").replace("￥", "").replace("¥", "").strip()
        number = float(text)
        return number if math.isfinite(number) else None
    except Exception:
        return None


def number_lines(value):
    return [to_number(part) for part in split_lines(value)]


def format_cost_lines(values):
    parts = []
    for value in values:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            parts.append(f"{round(value, 2):.2f}\r\n")
        else:
            parts.append(f"{norm(value) or '--'}\r\n")
    return "".join(parts)


def is_yellow(cell):
    if cell.fill.fill_type != "solid":
        return False
    color = cell.fill.fgColor
    if color.type == "rgb" and isinstance(color.rgb, str):
        return color.rgb.upper().endswith("FFFF00")
    if color.type == "indexed" and isinstance(color.indexed, int):
        try:
            return COLOR_INDEX[color.indexed].upper().endswith("FFFF00")
        except IndexError:
            return False
    return False


def clear_yellow(cell):
    if is_yellow(cell):
        cell.fill = NO_FILL


def is_real_sku(token):
    base_sku, _ = parse_sku_line(token)
    return sku_key(base_sku) not in MISSING_MARKERS


def is_valid_cost(base_sku, value):
    normalized_base, _ = parse_sku_line(base_sku)
    return value is not None and (
        value != 0 or sku_key(normalized_base) == "EB-S"
    )


def effective_max_row(worksheet, key_columns):
    # 普通模式下 openpyxl 只保留实际单元格；倒序检查可避开导出文件虚高的维度。
    for row in range(worksheet.max_row, 1, -1):
        if any(norm(worksheet.cell(row=row, column=column).value) for column in key_columns):
            return row
    return 1


def load_cost(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    worksheet = first_sheet(workbook)
    headers = header_row(worksheet)
    sku_col = find_col(headers, COST_COL_CANDIDATES["sku"], COST_COL_FALLBACK["sku"])
    wh_col = find_col(headers, COST_COL_CANDIDATES["wh"], COST_COL_FALLBACK["wh"])
    price_col = find_col(headers, COST_COL_CANDIDATES["price"], COST_COL_FALLBACK["price"])

    raw = defaultdict(list)
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        sku = norm(row[sku_col - 1]) if sku_col <= len(row) else ""
        warehouse = warehouse_key(row[wh_col - 1]) if wh_col <= len(row) else ""
        price = to_number(row[price_col - 1]) if price_col <= len(row) else None
        if not sku or not warehouse or not is_valid_cost(sku, price):
            continue
        raw[(warehouse, sku_key(sku))].append((row_number, price))
    workbook.close()

    # 沿用原脚本对万山重复价格取最大值的口径，并在处理报告中列明。
    costs = {}
    for key, values in raw.items():
        prices = [price for _, price in values]
        costs[key] = (
            0.0
            if key[1] == "EB-S" and 0.0 in prices
            else max(prices)
        )
    duplicates = {key: values for key, values in raw.items() if len(values) > 1}
    return costs, duplicates


def choose_mode(counter, first_rows, normalized_sku):
    if normalized_sku == "EB-S" and 0.0 in counter:
        return 0.0
    candidates = [price for price in counter if price != 0]
    return min(candidates, key=lambda price: (-counter[price], first_rows[price], price))


def build_existing_costs(worksheet, columns, max_row):
    observations = defaultdict(Counter)
    first_rows = defaultdict(dict)

    for row in range(2, max_row + 1):
        skus = split_store_skus(
            worksheet.cell(row=row, column=columns["store_sku"]).value
        )
        quantities = quantity_lines(
            worksheet.cell(row=row, column=columns["qty"]).value, len(skus)
        )
        costs = number_lines(worksheet.cell(row=row, column=columns["cost"]).value)
        warehouse = country_of(worksheet.cell(row=row, column=columns["shop"]).value)
        if (
            not warehouse
            or not skus
            or len(skus) != len(quantities)
            or len(skus) != len(costs)
        ):
            continue

        for index, token in enumerate(skus):
            if not is_real_sku(token):
                continue
            base_sku, multiple = parse_sku_line(token)
            if (
                not is_valid_cost(base_sku, costs[index])
                or quantities[index] is None
            ):
                continue
            denominator = multiple * quantities[index]
            if denominator == 0:
                continue
            unit_price = round(costs[index] / denominator, 6)
            key = (warehouse, sku_key(base_sku))
            observations[key][unit_price] += 1
            first_rows[key].setdefault(unit_price, row)

    selected = {}
    conflicts = []
    for key, counter in observations.items():
        price = choose_mode(counter, first_rows[key], key[1])
        selected[key] = price
        if len(counter) > 1:
            detail = ", ".join(
                f"{value:g}×{count}"
                for value, count in sorted(counter.items(), key=lambda item: item[0])
            )
            conflicts.append((key[0], key[1], detail, price))
    return selected, conflicts


def resolve_unit_price(table_costs, external_costs, warehouse, base_sku):
    key = (warehouse, sku_key(base_sku))
    if warehouse and key in table_costs:
        return table_costs[key], "本表"
    if warehouse and key in external_costs:
        return external_costs[key], "万山精确"

    normalized = sku_key(base_sku)
    prefix = normalized.split("-", 1)[0].strip() if "-" in normalized else ""
    if prefix not in KP_FALLBACK_PREFIXES:
        return None, None
    kp_sku = f"{prefix}-KP"
    kp_key = (warehouse, kp_sku)
    if warehouse and kp_key in table_costs:
        return table_costs[kp_key], f"本表{kp_sku}"
    if warehouse and kp_key in external_costs:
        return external_costs[kp_key], f"万山{kp_sku}"
    return None, None


def classify_file(path):
    try:
        workbook = openpyxl.load_workbook(path, read_only=True)
        worksheet = first_sheet(workbook)
        row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        workbook.close()
    except Exception:
        return None

    headers = {norm(value) for value in row if norm(value)}
    text = " ".join(headers)
    has_cost_sku = any(name in headers for name in COST_COL_CANDIDATES["sku"])
    has_cost_wh = any(name in headers for name in COST_COL_CANDIDATES["wh"])
    has_cost_price = any(name in headers for name in COST_COL_CANDIDATES["price"])
    if has_cost_sku and has_cost_wh and has_cost_price:
        return "cost"
    if "店铺SKU" in text and "商品成本" in text:
        return "fill"
    return None


def detect_files(args):
    candidates = args or [
        name
        for name in os.listdir(".")
        if name.lower().endswith(".xlsx")
        and not name.startswith("~$")
        and "_已填成本" not in name
    ]
    fills = []
    cost = None
    for path in candidates:
        role = classify_file(path)
        if role == "cost" and cost is None:
            cost = path
        elif role == "fill":
            fills.append(path)
    return sorted(dict.fromkeys(fills)), cost


def locate_fill_columns(worksheet):
    headers = header_row(worksheet)
    header_names = {text for _, text in headers}
    required = {"订单号", "BigSeller店铺", "店铺SKU", "数量", "商品成本", "总成本价"}
    missing = sorted(required - header_names)
    if missing:
        raise ValueError(f"大卖家表缺少必要列：{missing}")
    return {
        key: find_col(headers, FILL_COL_CANDIDATES[key], FILL_COL_FALLBACK[key])
        for key in FILL_COL_CANDIDATES
    }


def process(fill_path, cost_path):
    external_costs, external_duplicates = load_cost(cost_path)
    workbook = openpyxl.load_workbook(fill_path)
    worksheet = first_sheet(workbook)
    columns = locate_fill_columns(worksheet)
    max_row = effective_max_row(
        worksheet, [columns["order"], columns["store_sku"]]
    )
    table_costs, table_conflicts = build_existing_costs(worksheet, columns, max_row)

    summary = Counter(
        {
            "本表成本键": len(table_costs),
            "本表多价键": len(table_conflicts),
            "万山成本键": len(external_costs),
            "万山重复键": len(external_duplicates),
        }
    )
    total_sum = 0.0
    unmatched = []
    review = []
    external_used = []
    rebuilt = []

    for row in range(2, max_row + 1):
        sku_value = worksheet.cell(row=row, column=columns["store_sku"]).value
        skus = split_store_skus(sku_value)
        cost_cell = worksheet.cell(row=row, column=columns["cost"])
        total_cell = worksheet.cell(row=row, column=columns["total"])
        if not skus:
            clear_yellow(cost_cell)
            continue
        summary["SKU数据行"] += 1

        real_indices = [index for index, token in enumerate(skus) if is_real_sku(token)]
        if not real_indices:
            clear_yellow(cost_cell)
            summary["无真实SKU跳过行"] += 1
            continue

        quantities = quantity_lines(
            worksheet.cell(row=row, column=columns["qty"]).value,
            len(skus),
        )
        cost_parts = split_lines(cost_cell.value)
        costs = [to_number(part) for part in cost_parts]
        costs_aligned = len(cost_parts) == len(skus)

        if costs_aligned and all(
            is_valid_cost(parse_sku_line(skus[index])[0], costs[index])
            for index in real_indices
        ):
            clear_yellow(cost_cell)
            row_total = round(sum(costs[index] for index in real_indices), 2)
            summary["原成本保留行"] += 1
            if len(real_indices) != len(skus):
                summary["含占位符但真实成本完整行"] += 1
        else:
            original_cost = cost_cell.value
            shop = norm(worksheet.cell(row=row, column=columns["shop"]).value)
            warehouse = country_of(shop)
            order_id = norm(worksheet.cell(row=row, column=columns["order"]).value)
            values = []
            missing_costs = []
            row_review = []
            row_sources = []
            row_external_used = []
            row_zero_repriced = 0

            quantities_aligned = len(quantities) == len(skus)
            for index, token in enumerate(skus):
                if not is_real_sku(token):
                    marker = (
                        cost_parts[index]
                        if costs_aligned and cost_parts[index]
                        else norm(token) or "--"
                    )
                    values.append(marker)
                    continue

                base_sku, multiple = parse_sku_line(token)
                if costs_aligned and is_valid_cost(base_sku, costs[index]):
                    values.append(round(costs[index], 2))
                    continue
                if (
                    costs_aligned
                    and costs[index] == 0
                    and sku_key(base_sku) != "EB-S"
                ):
                    row_zero_repriced += 1

                unit_price, source = resolve_unit_price(
                    table_costs, external_costs, warehouse, base_sku
                )
                if unit_price is None:
                    values.append(None)
                    reason = (
                        "店铺国家无法识别"
                        if not warehouse
                        else "本表和万山均无成本"
                    )
                    missing_costs.append((token, reason))
                    continue

                quantity = quantities[index] if quantities_aligned else None
                if quantity is None:
                    values.append(None)
                    row_review.append((token, "数量子行数不一致或数量非数字"))
                    continue

                line_cost = round(unit_price * multiple * quantity, 2)
                values.append(line_cost)
                row_sources.append(source)
                if source != "本表":
                    row_external_used.append(
                        (
                            row,
                            order_id,
                            warehouse,
                            token,
                            source,
                            unit_price,
                            line_cost,
                        )
                    )

            if missing_costs:
                cost_cell.fill = YELLOW
                total_cell.value = None
                summary["未完成并标黄行"] += 1
                for token, reason in missing_costs:
                    unmatched.append(
                        (row, order_id, shop, warehouse or "", token, reason)
                    )
                for token, reason in row_review:
                    review.append((row, order_id, shop, token, reason))
                continue

            if row_review:
                clear_yellow(cost_cell)
                total_cell.value = None
                summary["结构异常未计算行"] += 1
                for token, reason in row_review:
                    review.append((row, order_id, shop, token, reason))
                continue

            for source in row_sources:
                summary[f"{source}取价子行"] += 1
            external_used.extend(row_external_used)
            summary["0成本重取子行"] += row_zero_repriced
            new_cost = format_cost_lines(values)
            cost_cell.value = new_cost
            cost_cell.number_format = "@"
            clear_yellow(cost_cell)
            row_total = round(
                sum(
                    value
                    for index, value in enumerate(values)
                    if index in real_indices and isinstance(value, (int, float))
                ),
                2,
            )
            summary["重建商品成本行"] += 1
            if len(costs) != len(skus) and any(value is not None for value in costs):
                rebuilt.append(
                    (row, order_id, norm(sku_value), norm(original_cost), norm(new_cost))
                )

        total_cell.value = row_total
        total_cell.number_format = "0.00"
        total_sum += row_total
        summary["已填总成本价行"] += 1

    summary["总成本价合计"] = round(total_sum, 2)
    write_report(
        workbook,
        fill_path,
        cost_path,
        summary,
        unmatched,
        review,
        external_used,
        rebuilt,
        table_conflicts,
        external_duplicates,
    )

    base, extension = os.path.splitext(fill_path)
    output_path = f"{base}_已填成本{extension or '.xlsx'}"
    workbook.save(output_path)
    workbook.close()

    table_kp_count = sum(
        value
        for key, value in summary.items()
        if key.startswith("本表")
        and key.endswith("取价子行")
        and key != "本表取价子行"
    )
    wan_kp_count = sum(
        value
        for key, value in summary.items()
        if key.startswith("万山")
        and key.endswith("取价子行")
        and key != "万山精确取价子行"
    )
    print("-" * 56)
    print(f"补成本表       : {fill_path}")
    print(f"SKU 数据行     : {summary['SKU数据行']}")
    print(f"原成本保留     : {summary['原成本保留行']}")
    print(f"重建商品成本   : {summary['重建商品成本行']}")
    print(f"0成本重取子行  : {summary['0成本重取子行']}")
    print(f"本表精确取价   : {summary['本表取价子行']}")
    print(f"本表KP兜底     : {table_kp_count}")
    print(f"万山精确取价   : {summary['万山精确取价子行']}")
    print(f"万山KP兜底     : {wan_kp_count}")
    print(f"未完成并标黄   : {summary['未完成并标黄行']}")
    print(f"无真实SKU跳过  : {summary['无真实SKU跳过行']}")
    print(f"结构异常未计算 : {summary['结构异常未计算行']}")
    print(f"已填总成本价   : {summary['已填总成本价行']}")
    print(f"总成本价合计   : {summary['总成本价合计']:.2f}")
    print(f"输出文件       : {output_path}")
    return output_path


def write_report(
    workbook,
    fill_path,
    cost_path,
    summary,
    unmatched,
    review,
    external_used,
    rebuilt,
    table_conflicts,
    external_duplicates,
):
    if "处理报告" in workbook.sheetnames:
        del workbook["处理报告"]
    sheet = workbook.create_sheet("处理报告")
    bold = Font(name="Arial", bold=True)
    normal = Font(name="Arial")
    row = 1

    def put(target_row, column, value, is_bold=False):
        cell = sheet.cell(row=target_row, column=column, value=value)
        cell.font = bold if is_bold else normal
        return cell

    put(row, 1, "处理报告", True)
    row += 1
    put(row, 1, f"补成本表：{fill_path}")
    row += 1
    put(row, 1, f"万山成本表：{cost_path}")
    row += 1
    put(row, 1, "本表多价选择：非EB-S忽略0后取单位成本众数；EB-S固定允许0；并列取表中靠上的值")
    row += 1
    put(row, 1, "万山重复价格选择：通常取最大值；EB-S存在0时固定取0")
    row += 2

    put(row, 1, "汇总", True)
    row += 1
    for key in [
        "SKU数据行",
        "原成本保留行",
        "重建商品成本行",
        "0成本重取子行",
        "本表取价子行",
        "万山精确取价子行",
        "未完成并标黄行",
        "无真实SKU跳过行",
        "含占位符但真实成本完整行",
        "结构异常未计算行",
        "已填总成本价行",
        "总成本价合计",
        "本表成本键",
        "本表多价键",
        "万山成本键",
        "万山重复键",
    ]:
        put(row, 1, key)
        put(row, 2, summary[key])
        row += 1
    table_kp_count = sum(
        value
        for key, value in summary.items()
        if key.startswith("本表")
        and key.endswith("取价子行")
        and key != "本表取价子行"
    )
    wan_kp_count = sum(
        value
        for key, value in summary.items()
        if key.startswith("万山")
        and key.endswith("取价子行")
        and key != "万山精确取价子行"
    )
    put(row, 1, "本表KP兜底取价子行")
    put(row, 2, table_kp_count)
    row += 1
    put(row, 1, "万山KP兜底取价子行")
    put(row, 2, wan_kp_count)
    row += 2

    def table(title, headers, rows):
        nonlocal row
        put(row, 1, title, True)
        row += 1
        for column, header in enumerate(headers, start=1):
            put(row, column, header, True)
        row += 1
        if not rows:
            put(row, 1, "（无）")
            row += 2
            return
        for values in rows:
            for column, value in enumerate(values, start=1):
                put(row, column, value)
            row += 1
        row += 1

    table(
        "① 真实SKU在本表和万山均无成本（商品成本已标黄）",
        ["行号", "订单号", "BigSeller店铺", "国家", "SKU", "原因"],
        unmatched,
    )
    table(
        "② 数量或结构异常（不标黄，总成本价留空）",
        ["行号", "订单号", "BigSeller店铺", "SKU", "原因"],
        review,
    )
    table(
        "③ 非本表精确匹配的取价明细",
        ["行号", "订单号", "国家", "SKU", "来源", "单位成本", "子行成本"],
        external_used,
    )
    table(
        "④ 原商品成本子行数不完整、已重建的行",
        ["行号", "订单号", "店铺SKU", "原商品成本", "新商品成本"],
        rebuilt,
    )
    table(
        "⑤ 本表单位成本多价清单（已按非零众数选择，EB-S例外）",
        ["国家", "基础SKU", "全部单位成本×次数", "取用值"],
        table_conflicts,
    )
    table(
        "⑥ 万山重复价格清单（通常取最大值，EB-S例外）",
        ["国家", "SKU", "原表行号与价格", "取用值"],
        [
            (
                key[0],
                key[1],
                ", ".join(f"{row_number}:{price:g}" for row_number, price in values),
                (
                    0.0
                    if key[1] == "EB-S"
                    and any(price == 0 for _, price in values)
                    else max(price for _, price in values)
                ),
            )
            for key, values in external_duplicates.items()
        ],
    )

    for column, width in {
        "A": 14,
        "B": 24,
        "C": 28,
        "D": 18,
        "E": 34,
        "F": 28,
        "G": 14,
    }.items():
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows():
        for cell in cells:
            alignment = copy(cell.alignment)
            alignment.vertical = "center"
            alignment.wrap_text = True
            cell.alignment = alignment


def main():
    fills, cost_path = detect_files(sys.argv[1:])
    if not fills or not cost_path:
        print("未能同时找到大卖家表和“万山成本价.xlsx”。")
        print(f"识别到大卖家表：{fills}")
        print(f"识别到成本价表：{cost_path}")
        sys.exit(1)

    print(f"成本价表：{cost_path}")
    print(f"待处理大卖家表（{len(fills)} 个）：{fills}")
    outputs = [process(fill_path, cost_path) for fill_path in fills]
    print("\n全部处理完成：")
    for output in outputs:
        print(f"  {output}")


if __name__ == "__main__":
    main()
