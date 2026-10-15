# -*- coding: utf-8 -*-
"""按店铺名称拆分虾皮大卖家和补单明细数据，并写入各店铺模板。"""

from __future__ import annotations

import argparse
import json
import math
import os
import posixpath
import sys
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.cell import Cell


BASE = Path(__file__).resolve().parent
DATA_DIR = BASE / "数据源"
TEMPLATE_DIR = DATA_DIR
OUTPUT_DIR = BASE / "输出"

SOURCE_SPECS = {
    "大卖家": {
        "filename_markers": (("虾皮", "大卖家"), ("虾皮", "成本")),
        "shop_header": "BigSeller店铺",
        "order_header": "订单号",
        "required_headers": {"订单号", "BigSeller店铺", "商品成本", "总成本价"},
    },
    "补单明细": {
        "filename_markers": (("测评",),),
        "shop_header": "店铺名称",
        "order_header": "线上单号",
        "required_headers": {"店铺名称", "系统单号", "业务单号", "线上单号", "实付人民币"},
    },
}

# 两组新旧名称的平台店铺分别同为 Beauty & Health、Avitaglow.。
SHOP_ALIASES = {"印尼保健1店": "保健1店", "印尼保健7店": "保健7店"}

STYLE_ATTRS = ("font", "fill", "border", "alignment", "number_format", "protection")
ROW_DIM_ATTRS = (
    "height",
    "hidden",
    "outlineLevel",
    "collapsed",
    "thickTop",
    "thickBot",
)
COL_DIM_ATTRS = ("width", "hidden", "bestFit", "outlineLevel", "collapsed")
VIEW_ATTRS = (
    "showGridLines",
    "rightToLeft",
    "zoomScale",
    "zoomScaleNormal",
    "showRowColHeaders",
    "showZeros",
    "view",
    "topLeftCell",
)
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


@dataclass
class RowData:
    source_row: int
    cells: tuple
    dimension: tuple | None


@dataclass
class SheetSettings:
    sheet_format: object
    sheet_properties: object
    page_margins: object
    page_setup: object
    print_options: object
    sheet_state: str
    freeze_panes: object
    auto_filter: object
    view: tuple
    column_dimensions: dict[str, tuple]
    merged_ranges: tuple[str, ...]


@dataclass
class SourceData:
    header: RowData
    max_column: int
    shop_col: int
    order_col: int
    index: dict[str, list[RowData]]
    styles: dict[int, tuple | None]
    settings: SheetSettings
    strict_lookup: dict[str, list[str]]
    compact_lookup: dict[str, list[str]]
    duplicate_orders: dict[str, int]
    blank_order_rows: list[int]


def configure_console():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", line_buffering=True)


def normalized(value):
    return unicodedata.normalize("NFKC", str(value)).strip().casefold()


def compact_normalized(value):
    return "".join(char for char in normalized(value) if not char.isspace())


def header_columns(ws):
    result = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value is not None and str(value).strip():
            result[str(value).strip()] = col
    return result


def find_source_sheet(wb, required_headers):
    matches = [ws for ws in wb.worksheets if required_headers <= set(header_columns(ws))]
    if len(matches) > 1:
        raise ValueError(f"源文件含多个符合表头的 sheet，无法唯一确定：{[ws.title for ws in matches]}")
    if matches:
        return matches[0]
    raise ValueError(f"源文件中找不到包含这些表头的 sheet：{sorted(required_headers)}")


def discover_source_files(skip_review=False):
    files = [path for path in DATA_DIR.glob("*.xlsx") if not path.name.startswith("~$")]
    specs = {}
    for sheet_name, spec in SOURCE_SPECS.items():
        if skip_review and sheet_name == "补单明细":
            continue
        matches = [path for path in files if any(
            all(marker in path.stem for marker in markers)
            for markers in spec["filename_markers"]
        )]
        if len(matches) != 1:
            raise ValueError(f"{sheet_name} 源表必须唯一，实际找到 {len(matches)} 个：{[p.name for p in matches]}")
        specs[sheet_name] = {**spec, "file": matches[0]}
    return specs


def build_lookup(shops, normalizer):
    lookup = defaultdict(list)
    for shop in shops:
        lookup[normalizer(shop)].append(shop)
    return dict(lookup)


def capture_sheet_settings(ws):
    if ws.merged_cells.ranges:
        raise ValueError(f"源 sheet {ws.title} 含合并单元格，无法安全拆分数据行")
    if ws.tables:
        raise ValueError(f"源 sheet {ws.title} 含 Excel 表对象，无法安全拆分")
    return SheetSettings(
        sheet_format=copy(ws.sheet_format),
        sheet_properties=copy(ws.sheet_properties),
        page_margins=copy(ws.page_margins),
        page_setup=copy(ws.page_setup),
        print_options=copy(ws.print_options),
        sheet_state=ws.sheet_state,
        freeze_panes=ws.freeze_panes,
        auto_filter=ws.auto_filter.ref,
        view=tuple(getattr(ws.sheet_view, attr) for attr in VIEW_ATTRS),
        column_dimensions={
            key: dimension_signature(dim, COL_DIM_ATTRS)
            for key, dim in ws.column_dimensions.items()
        },
        merged_ranges=tuple(sorted(str(rng) for rng in ws.merged_cells.ranges)),
    )


def snapshot_cell(cell, style_keys, styles):
    if cell.hyperlink is not None or cell.comment is not None:
        raise ValueError(f"源数据单元格 {cell.coordinate} 含超链接或批注，当前不支持拆分")
    if not cell.has_style:
        style_code = 0
    else:
        key = tuple(cell._style)
        style_code = style_keys.get(key)
        if style_code is None:
            style_code = len(style_keys) + 1
            style_keys[key] = style_code
            styles[style_code] = tuple(
                copy(getattr(cell, attr)) if attr != "number_format" else cell.number_format
                for attr in STYLE_ATTRS
            ) + (cell.quotePrefix, cell.pivotButton)
    return (cell._value, cell.data_type, style_code)


def load_source(spec):
    path = spec["file"]
    if not path.exists():
        raise FileNotFoundError(f"找不到源文件：{path}")
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = find_source_sheet(wb, spec["required_headers"])
    columns = header_columns(ws)
    shop_col = columns[spec["shop_header"]]
    order_col = columns[spec["order_header"]]

    max_column = ws.max_column
    settings = capture_sheet_settings(ws)
    style_keys = {}
    styles = {0: None}
    index = defaultdict(list)
    order_counts = Counter()
    blank_order_rows = []
    header = None
    for row_number, row in enumerate(
        ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=max_column),
        start=1,
    ):
        cells = tuple(snapshot_cell(cell, style_keys, styles) for cell in row)
        src_dim = ws.row_dimensions.get(row_number)
        row_data = RowData(
            source_row=row_number,
            cells=cells,
            dimension=dimension_signature(src_dim, ROW_DIM_ATTRS) if src_dim else None,
        )
        if row_number == 1:
            header = row_data
            continue
        shop_value = cells[shop_col - 1][0]
        if shop_value is None or not str(shop_value).strip():
            if any(cell[0] not in (None, "") for cell in cells):
                raise ValueError(f"{path.name}/{ws.title} 第 {row_number} 行有数据但缺少店铺名称")
            continue
        shop = str(shop_value).strip()
        index[shop].append(row_data)
        order_value = cells[order_col - 1][0]
        if order_value is None or not str(order_value).strip():
            blank_order_rows.append(row_number)
        else:
            order_counts[str(order_value).strip()] += 1

    duplicate_orders = {
        order: count for order, count in order_counts.items() if count > 1
    }
    index = dict(index)
    wb.close()
    return SourceData(
        header=header,
        max_column=max_column,
        shop_col=shop_col,
        order_col=order_col,
        index=index,
        styles=styles,
        settings=settings,
        strict_lookup=build_lookup(index, normalized),
        compact_lookup=build_lookup(index, compact_normalized),
        duplicate_orders=duplicate_orders,
        blank_order_rows=blank_order_rows,
    )


def candidate_names(stem):
    stem = stem.strip()
    candidates = [stem]
    if stem.endswith("店"):
        candidates.append(stem[:-1])
    else:
        candidates.append(stem + "店")
    return candidates


def lookup_unique(lookup, key, stem):
    matches = lookup.get(key, [])
    if len(matches) > 1:
        raise ValueError(f"店铺名匹配不唯一：{stem} -> {matches}")
    return matches[0] if matches else None


def resolve_shop_name(stem, source):
    candidates = candidate_names(stem)
    for candidate in candidates:
        match = lookup_unique(source.strict_lookup, normalized(candidate), stem)
        if match is not None:
            return match
    for candidate in candidates:
        match = lookup_unique(source.compact_lookup, compact_normalized(candidate), stem)
        if match is not None:
            return match
    alias = SHOP_ALIASES.get(stem.strip())
    if alias is not None:
        return lookup_unique(source.strict_lookup, normalized(alias), stem)
    return None


def copy_cell(cell_data, dst, source_styles, style_cache):
    value, data_type, style_code = cell_data
    dst._value = value
    dst.data_type = data_type
    if style_code:
        style = style_cache.get(style_code)
        if style is None:
            proxy = Cell(dst.parent, row=1, column=1)
            style_values = source_styles[style_code]
            for attr, style_value in zip(STYLE_ATTRS, style_values):
                setattr(proxy, attr, copy(style_value) if attr != "number_format" else style_value)
            proxy.quotePrefix = style_values[-2]
            proxy.pivotButton = style_values[-1]
            style = copy(proxy._style)
            style_cache[style_code] = style
        dst._style = style


def copy_row_dimension(dimension, dst_ws, dst_row):
    if dimension is None:
        return
    dst_dim = dst_ws.row_dimensions[dst_row]
    for attr, value in zip(ROW_DIM_ATTRS, dimension):
        setattr(dst_dim, attr, value)


def copy_sheet_settings(settings, dst):
    dst.sheet_format = copy(settings.sheet_format)
    dst.sheet_properties = copy(settings.sheet_properties)
    dst.page_margins = copy(settings.page_margins)
    dst.page_setup = copy(settings.page_setup)
    dst.print_options = copy(settings.print_options)
    dst.sheet_state = settings.sheet_state
    dst.freeze_panes = settings.freeze_panes
    dst.auto_filter.ref = settings.auto_filter
    for attr, value in zip(VIEW_ATTRS, settings.view):
        setattr(dst.sheet_view, attr, value)

    for key, signature in settings.column_dimensions.items():
        dst_dim = dst.column_dimensions[key]
        for attr, value in zip(COL_DIM_ATTRS, signature):
            setattr(dst_dim, attr, value)


def rebuild_sheet(wb, sheet_name, source, rows):
    active_title = wb.active.title
    created = sheet_name not in wb.sheetnames
    if created:
        sheet_index = len(wb.sheetnames)
    else:
        sheet_index = wb.sheetnames.index(sheet_name)
        del wb[sheet_name]

    dst_ws = wb.create_sheet(sheet_name, sheet_index)
    copy_sheet_settings(source.settings, dst_ws)
    style_cache = {}

    for col, cell_data in enumerate(source.header.cells, start=1):
        copy_cell(cell_data, dst_ws.cell(1, col), source.styles, style_cache)
    copy_row_dimension(source.header.dimension, dst_ws, 1)

    for dst_row, row_data in enumerate(rows, start=2):
        for col, cell_data in enumerate(row_data.cells, start=1):
            copy_cell(cell_data, dst_ws.cell(dst_row, col), source.styles, style_cache)
        copy_row_dimension(row_data.dimension, dst_ws, dst_row)

    if active_title in wb.sheetnames:
        wb.active = wb.sheetnames.index(active_title)
    return created


def package_sheet_parts(archive):
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{PKG_REL_NS}}}Relationship")
    }
    result = {}
    for sheet in workbook_root.findall(f".//{{{MAIN_NS}}}sheet"):
        rel_id = sheet.attrib[f"{{{DOC_REL_NS}}}id"]
        target = targets[rel_id].lstrip("/")
        if not target.startswith("xl/"):
            target = posixpath.normpath(posixpath.join("xl", target))
        result[sheet.attrib["name"]] = target
    return result


def sheet_rel_part(sheet_part):
    folder, name = posixpath.split(sheet_part)
    return posixpath.join(folder, "_rels", f"{name}.rels")


def merge_content_types(original_data, generated_data):
    original = ET.fromstring(original_data)
    generated = ET.fromstring(generated_data)
    default_tag = f"{{{CONTENT_NS}}}Default"
    override_tag = f"{{{CONTENT_NS}}}Override"
    defaults = {item.attrib["Extension"] for item in generated.findall(default_tag)}
    overrides = {item.attrib["PartName"] for item in generated.findall(override_tag)}
    for item in original.findall(default_tag):
        if item.attrib["Extension"] not in defaults:
            generated.append(copy(item))
    for item in original.findall(override_tag):
        if item.attrib["PartName"] not in overrides:
            generated.append(copy(item))
    ET.register_namespace("", CONTENT_NS)
    return ET.tostring(generated, encoding="utf-8", xml_declaration=True)


def merge_workbook_relationships(original_data, generated_data):
    original = ET.fromstring(original_data)
    generated = ET.fromstring(generated_data)
    rel_tag = f"{{{PKG_REL_NS}}}Relationship"
    existing = {
        (item.attrib.get("Type"), item.attrib.get("Target"), item.attrib.get("TargetMode"))
        for item in generated.findall(rel_tag)
    }
    used_ids = {item.attrib.get("Id") for item in generated.findall(rel_tag)}
    number = 1
    for item in original.findall(rel_tag):
        key = (
            item.attrib.get("Type"),
            item.attrib.get("Target"),
            item.attrib.get("TargetMode"),
        )
        if key in existing:
            continue
        while f"rId{number}" in used_ids:
            number += 1
        cloned = copy(item)
        cloned.attrib["Id"] = f"rId{number}"
        generated.append(cloned)
        used_ids.add(f"rId{number}")
        existing.add(key)
    ET.register_namespace("", PKG_REL_NS)
    return ET.tostring(generated, encoding="utf-8", xml_declaration=True)


def preserve_original_sheets(generated_path, template_path, preserved_sheet_names):
    patched_path = generated_path.with_suffix(".patched.xlsx")
    try:
        with zipfile.ZipFile(template_path, "r") as original, zipfile.ZipFile(
            generated_path, "r"
        ) as generated:
            original_parts = package_sheet_parts(original)
            generated_parts = package_sheet_parts(generated)
            preserved = {}
            skipped_generated = set()

            for sheet_name in preserved_sheet_names:
                src_part = original_parts[sheet_name]
                dst_part = generated_parts[sheet_name]
                preserved[dst_part] = original.read(src_part)
                skipped_generated.add(dst_part)
                src_rel = sheet_rel_part(src_part)
                dst_rel = sheet_rel_part(dst_part)
                skipped_generated.add(dst_rel)
                if src_rel in original.namelist():
                    preserved[dst_rel] = original.read(src_rel)

            generated_data = {
                info.filename: generated.read(info.filename)
                for info in generated.infolist()
                if info.filename not in skipped_generated
            }
            generated_data["[Content_Types].xml"] = merge_content_types(
                original.read("[Content_Types].xml"),
                generated.read("[Content_Types].xml"),
            )
            generated_data["xl/_rels/workbook.xml.rels"] = merge_workbook_relationships(
                original.read("xl/_rels/workbook.xml.rels"),
                generated.read("xl/_rels/workbook.xml.rels"),
            )

            with zipfile.ZipFile(patched_path, "w") as patched:
                written = set()
                for info in original.infolist():
                    name = info.filename
                    data = preserved.get(name, generated_data.get(name, original.read(name)))
                    patched.writestr(info, data)
                    written.add(name)
                for name, data in {**generated_data, **preserved}.items():
                    if name in written:
                        continue
                    if name in generated.namelist():
                        info = generated.getinfo(name)
                    else:
                        info = zipfile.ZipInfo(name)
                        info.compress_type = zipfile.ZIP_DEFLATED
                    patched.writestr(info, data)
        os.replace(patched_path, generated_path)
    finally:
        if patched_path.exists():
            patched_path.unlink()


def save_atomic(wb, output_path, template_path, preserved_sheet_names):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    )
    temp_path = Path(handle.name)
    handle.close()
    try:
        wb.save(temp_path)
        preserve_original_sheets(temp_path, template_path, preserved_sheet_names)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def values_equal(left, right):
    if left in (None, "") and right in (None, ""):
        return True
    if isinstance(left, float) and isinstance(right, float):
        if math.isnan(left) and math.isnan(right):
            return True
    return left == right


def style_signature(cell, cache):
    if not cell.has_style:
        return None
    key = tuple(cell._style)
    if key not in cache:
        cache[key] = tuple(
            copy(getattr(cell, attr)) if attr != "number_format" else cell.number_format
            for attr in STYLE_ATTRS
        ) + (cell.quotePrefix, cell.pivotButton)
    return cache[key]


def hyperlink_signature(cell):
    link = cell.hyperlink
    if link is None:
        return None
    return (link.target, link.location, link.tooltip, link.display)


def comment_signature(cell):
    comment = cell.comment
    if comment is None:
        return None
    return (comment.text, comment.author)


def dimension_signature(dim, attrs):
    return tuple(getattr(dim, attr) for attr in attrs)


def worksheet_settings_signature(ws):
    return (
        ws.freeze_panes,
        ws.auto_filter.ref,
        copy(ws.sheet_format),
        copy(ws.sheet_properties),
        copy(ws.page_margins),
        copy(ws.page_setup),
        copy(ws.print_options),
        tuple(getattr(ws.sheet_view, attr) for attr in VIEW_ATTRS),
        tuple(sorted(str(rng) for rng in ws.merged_cells.ranges)),
    )


def captured_settings_signature(settings):
    return (
        settings.freeze_panes,
        settings.auto_filter,
        settings.sheet_format,
        settings.sheet_properties,
        settings.page_margins,
        settings.page_setup,
        settings.print_options,
        settings.view,
        settings.merged_ranges,
    )


def validate_target_sheet(ws, rows, source, shop):
    expected_rows = len(rows) + 1
    if ws.max_row != expected_rows or ws.max_column != source.max_column:
        raise ValueError(
            f"{ws.title} 尺寸不正确：实际 {ws.max_row}x{ws.max_column}，"
            f"应为 {expected_rows}x{source.max_column}"
        )

    dst_style_cache = {}
    for dst_row, row_data in enumerate([source.header, *rows], start=1):
        for col, cell_data in enumerate(row_data.cells, start=1):
            value, data_type, style_code = cell_data
            dst_cell = ws.cell(dst_row, col)
            if not values_equal(value, dst_cell.value):
                raise ValueError(f"{ws.title}!{dst_cell.coordinate} 的值与源表不一致")
            both_empty = value in (None, "") and dst_cell.value in (None, "")
            if not both_empty and data_type != dst_cell.data_type:
                raise ValueError(f"{ws.title}!{dst_cell.coordinate} 的数据类型与源表不一致")
            if source.styles[style_code] != style_signature(dst_cell, dst_style_cache):
                raise ValueError(f"{ws.title}!{dst_cell.coordinate} 的格式与源表不一致")
            if hyperlink_signature(dst_cell) is not None:
                raise ValueError(f"{ws.title}!{dst_cell.coordinate} 的超链接与源表不一致")
            if comment_signature(dst_cell) is not None:
                raise ValueError(f"{ws.title}!{dst_cell.coordinate} 的批注与源表不一致")

        dst_dim = ws.row_dimensions.get(dst_row)
        dst_sig = dimension_signature(dst_dim, ROW_DIM_ATTRS) if dst_dim else None
        if row_data.dimension != dst_sig:
            raise ValueError(f"{ws.title} 第 {dst_row} 行的行格式与源表不一致")

    all_columns = set(source.settings.column_dimensions) | set(ws.column_dimensions)
    for key in all_columns:
        dst_dim = ws.column_dimensions.get(key)
        src_sig = source.settings.column_dimensions.get(key)
        dst_sig = dimension_signature(dst_dim, COL_DIM_ATTRS) if dst_dim else None
        if src_sig != dst_sig:
            raise ValueError(f"{ws.title} 的 {key} 列格式与源表不一致")

    if worksheet_settings_signature(ws) != captured_settings_signature(source.settings):
        raise ValueError(f"{ws.title} 的 sheet 设置与源表不一致")

    shops = {
        str(ws.cell(row, source.shop_col).value).strip()
        for row in range(2, ws.max_row + 1)
    }
    if shops != ({shop} if rows else set()):
        raise ValueError(f"{ws.title} 混入了其他店铺数据：{sorted(shops)}")

    orders = []
    missing_orders = 0
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, source.order_col).value
        if value is None or not str(value).strip():
            missing_orders += 1
        else:
            orders.append(str(value).strip())
    return len(rows), len(set(orders)), missing_orders


def material_cell_keys(ws):
    return {
        key
        for key, cell in ws._cells.items()
        if cell.value not in (None, "")
        or cell.has_style
        or cell.hyperlink is not None
        or cell.comment is not None
    }


def validate_preserved_sheet(original, output):
    if worksheet_settings_signature(original) != worksheet_settings_signature(output):
        raise ValueError(f"非目标 sheet {original.title} 的 sheet 设置发生变化")
    if tuple(original.column_dimensions) != tuple(output.column_dimensions):
        raise ValueError(f"非目标 sheet {original.title} 的列维度发生变化")
    for key in set(original.column_dimensions) | set(output.column_dimensions):
        left = original.column_dimensions.get(key)
        right = output.column_dimensions.get(key)
        left_sig = dimension_signature(left, COL_DIM_ATTRS) if left else None
        right_sig = dimension_signature(right, COL_DIM_ATTRS) if right else None
        if left_sig != right_sig:
            raise ValueError(f"非目标 sheet {original.title} 的 {key} 列格式发生变化")
    for key in set(original.row_dimensions) | set(output.row_dimensions):
        left = original.row_dimensions.get(key)
        right = output.row_dimensions.get(key)
        left_sig = dimension_signature(left, ROW_DIM_ATTRS) if left else None
        right_sig = dimension_signature(right, ROW_DIM_ATTRS) if right else None
        if left_sig != right_sig:
            raise ValueError(f"非目标 sheet {original.title} 的第 {key} 行格式发生变化")

    keys = material_cell_keys(original) | material_cell_keys(output)
    left_styles = {}
    right_styles = {}
    for row, col in keys:
        left = original.cell(row, col)
        right = output.cell(row, col)
        both_empty = left.value in (None, "") and right.value in (None, "")
        if not both_empty and not values_equal(left.value, right.value):
            raise ValueError(f"非目标 sheet {original.title}!{right.coordinate} 的值发生变化")
        if not both_empty and left.data_type != right.data_type:
            raise ValueError(f"非目标 sheet {original.title}!{right.coordinate} 的类型发生变化")
        if style_signature(left, left_styles) != style_signature(right, right_styles):
            raise ValueError(f"非目标 sheet {original.title}!{right.coordinate} 的格式发生变化")
        if hyperlink_signature(left) != hyperlink_signature(right):
            raise ValueError(f"非目标 sheet {original.title}!{right.coordinate} 的超链接发生变化")
        if comment_signature(left) != comment_signature(right):
            raise ValueError(f"非目标 sheet {original.title}!{right.coordinate} 的批注发生变化")


def validate_output(output_path, template_path, expected_sheetnames, source_rows, matches, sources):
    original_wb = openpyxl.load_workbook(template_path, data_only=False)
    output_wb = openpyxl.load_workbook(output_path, data_only=False)
    try:
        if output_wb.sheetnames != expected_sheetnames:
            raise ValueError(f"{output_path.name} 的 sheet 名称或顺序不正确")
        for sheet_name, source in sources.items():
            validate_target_sheet(
                output_wb[sheet_name],
                source_rows[sheet_name],
                source,
                matches[sheet_name],
            )
        for sheet_name in original_wb.sheetnames:
            if sheet_name not in sources:
                validate_preserved_sheet(original_wb[sheet_name], output_wb[sheet_name])
        with zipfile.ZipFile(template_path) as original, zipfile.ZipFile(output_path) as output:
            original_parts = package_sheet_parts(original)
            output_parts = package_sheet_parts(output)
            for sheet_name in original_wb.sheetnames:
                if sheet_name not in sources:
                    if original.read(original_parts[sheet_name]) != output.read(output_parts[sheet_name]):
                        raise ValueError(f"非目标 sheet {sheet_name} 的原始 XML 发生变化")
    finally:
        original_wb.close()
        output_wb.close()


def target_files(only, excluded=()):
    files = sorted(
        path
        for path in TEMPLATE_DIR.glob("*.xlsx")
        if not path.name.startswith("~$") and not any(
            all(marker in path.stem for marker in markers)
            for spec in SOURCE_SPECS.values() for markers in spec["filename_markers"]
        )
        and normalized(path.stem) not in {normalized(Path(name).stem) for name in excluded}
    )
    if only:
        wanted = normalized(Path(only).stem)
        files = [path for path in files if normalized(path.stem) == wanted]
    return files


def unique_nonblank_orders(source, rows):
    return len(
        {
            str(value).strip()
            for row in rows
            if (value := row.cells[source.order_col - 1][0]) is not None
            and str(value).strip()
        }
    )


def duplicate_row_count(source):
    counts = Counter(tuple(cell[0] for cell in row.cells)
                     for rows in source.index.values() for row in rows)
    return sum(count - 1 for count in counts.values() if count > 1)


def run(only=None, output_dir=OUTPUT_DIR, excluded=(), skip_review=False):
    if not TEMPLATE_DIR.is_dir():
        raise FileNotFoundError(f"找不到模板目录：{TEMPLATE_DIR}")
    specs = discover_source_files(skip_review)
    files = target_files(only, excluded)
    if not files:
        raise FileNotFoundError("没有找到待处理的虾皮店铺文件")

    sources = {sheet: load_source(spec) for sheet, spec in specs.items()}
    try:
        all_matches = {}
        aliases = []
        for sheet_name, source in sources.items():
            matches = {}
            for path in files:
                shop = resolve_shop_name(path.stem, source)
                matches[path] = shop
                if shop is not None and normalized(path.stem) != normalized(shop):
                    aliases.append((sheet_name, path.name, shop))
            matched_shops = [shop for shop in matches.values() if shop is not None]
            duplicates = [shop for shop, count in Counter(matched_shops).items() if count > 1]
            if duplicates:
                raise ValueError(f"{sheet_name} 有多个模板匹配到同一源店铺：{duplicates}")
            if only is None:
                unmatched_source = sorted(set(source.index) - set(matched_shops))
                if unmatched_source:
                    raise ValueError(f"{sheet_name} 中这些源店铺没有对应模板：{unmatched_source}")
            all_matches[sheet_name] = matches

        totals = {
            sheet: {"rows": 0, "orders": 0, "missing": 0}
            for sheet in sources
        }
        created_sheets = []
        file_results = []
        for path in files:
            matches = {
                sheet: all_matches[sheet][path]
                for sheet in sources
            }
            source_rows = {
                sheet: sources[sheet].index.get(matches[sheet], [])
                if matches[sheet] is not None
                else []
                for sheet in sources
            }

            wb = openpyxl.load_workbook(path, data_only=False)
            try:
                expected_sheetnames = wb.sheetnames[:]
                for sheet_name, source in sources.items():
                    created = rebuild_sheet(
                        wb,
                        sheet_name,
                        source,
                        source_rows[sheet_name],
                    )
                    if created:
                        expected_sheetnames.append(sheet_name)
                        created_sheets.append((path.name, sheet_name))
                output_path = output_dir / path.name
                preserved_sheet_names = [
                    sheet for sheet in wb.sheetnames if sheet not in sources
                ]
                save_atomic(wb, output_path, path, preserved_sheet_names)
            finally:
                wb.close()

            validate_output(
                output_path,
                path,
                expected_sheetnames,
                source_rows,
                matches,
                sources,
            )

            parts = []
            file_result = {"file": path.name, "sheets": {}}
            for sheet_name, source in sources.items():
                rows = source_rows[sheet_name]
                order_count = unique_nonblank_orders(source, rows)
                missing_count = sum(
                    1
                    for row in rows
                    if row.cells[source.order_col - 1][0] is None
                    or not str(row.cells[source.order_col - 1][0]).strip()
                )
                totals[sheet_name]["rows"] += len(rows)
                totals[sheet_name]["orders"] += order_count
                totals[sheet_name]["missing"] += missing_count
                file_result["sheets"][sheet_name] = {
                    "source_shop": matches[sheet_name], "rows": len(rows),
                    "unique_orders": order_count, "blank_orders": missing_count,
                }
                detail = f"{sheet_name} {len(rows)} 行/{order_count} 单"
                if missing_count:
                    detail += f"，{missing_count} 行无订单号"
                parts.append(detail)
            print(f"[通过] {path.name}｜" + "｜".join(parts))
            file_results.append(file_result)

        print(f"\n完成：{len(files)} 个文件；输出目录：{output_dir}")
        for sheet_name, total in totals.items():
            detail = f"{sheet_name}：{total['rows']} 行，{total['orders']} 个唯一订单"
            if total["missing"]:
                detail += f"，{total['missing']} 行缺少订单号"
            print(detail)
        if created_sheets:
            print("自动创建缺失 sheet：" + "；".join(f"{file} -> {sheet}" for file, sheet in created_sheets))
        if aliases:
            print("店铺别名匹配：" + "；".join(f"{sheet}/{file} -> {shop}" for sheet, file, shop in aliases))
        for sheet_name, source in sources.items():
            if source.duplicate_orders:
                details = "，".join(
                    f"{order}×{count}" for order, count in list(source.duplicate_orders.items())[:10]
                )
                print(f"源数据重复订单号（{sheet_name}）：共 {len(source.duplicate_orders)} 个；示例 {details}；完整明细见校验报告")
            if source.blank_order_rows:
                print(
                    f"源数据缺少订单号（{sheet_name}）："
                    f"Excel 行 {', '.join(map(str, source.blank_order_rows))}"
                )
        report = {
            "source_files": {sheet: str(spec["file"]) for sheet, spec in specs.items()},
            "review_mode": "保留模板补单明细，不填入新数据" if skip_review else "按测评原表填写",
            "files": file_results, "totals": totals, "created_sheets": created_sheets,
            "aliases": aliases, "excluded_templates": list(excluded),
            "validation": {"target_values_types_styles_dimensions": True,
                           "non_target_values_styles_and_xml_unchanged": True},
            "source_anomalies": {sheet: {
                "duplicate_orders": source.duplicate_orders,
                "duplicate_order_number_count": len(source.duplicate_orders),
                "extra_rows_with_repeated_order": sum(count - 1 for count in source.duplicate_orders.values()),
                "identical_duplicate_rows": duplicate_row_count(source),
                "blank_order_rows": source.blank_order_rows,
                "global_unique_orders": unique_nonblank_orders(source, [row for rows in source.index.values() for row in rows]),
                "no_data_templates": [path.name for path, shop in all_matches[sheet].items() if shop is None],
            } for sheet, source in sources.items()},
        }
        report_path = output_dir / "填表校验报告.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"校验报告：{report_path}")
    finally:
        pass


def main():
    configure_console()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("shop", nargs="?", help="可选：只处理指定店铺文件名")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR, help="输出目录")
    parser.add_argument("--exclude", action="append", default=[], help="排除已确认属于其他平台的模板，可重复指定")
    parser.add_argument("--skip-review", action="store_true", help="仅处理大卖家，原样保留模板补单明细；须确认本月无需填补单")
    args = parser.parse_args()
    output_dir = args.output_dir if args.output_dir.is_absolute() else BASE / args.output_dir
    try:
        run(args.shop, output_dir.resolve(), args.exclude, args.skip_review)
    except Exception as exc:
        print(f"[失败] {exc}", file=sys.stderr)
        raise SystemExit(1) from exc


if __name__ == "__main__":
    main()
