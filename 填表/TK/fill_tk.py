# -*- coding: utf-8 -*-
"""按店铺拆分 TK 订单数据，并写入各店铺工作簿的“大卖家”sheet。"""

from __future__ import annotations

import argparse
import json
import math
import os
import posixpath
import sys
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl


BASE = Path(__file__).resolve().parent
DATA_DIR = BASE / "数据源"
TARGET_DIR = DATA_DIR
OUTPUT_DIR = BASE / "输出"
TARGET_SHEET = "大卖家"
ORDER_HEADER = "订单号"
SHOP_HEADER = "BigSeller店铺"
REQUIRED_HEADERS = {ORDER_HEADER, SHOP_HEADER, "商品成本", "总成本价"}

STYLE_ATTRS = ("font", "fill", "border", "alignment", "number_format", "protection")
ROW_DIM_ATTRS = (
    "height",
    "hidden",
    "outlineLevel",
    "collapsed",
    "thickTop",
    "thickBot",
)
COL_DIM_ATTRS = ("width", "hidden", "bestFit", "outlineLevel", "collapsed")
VIEW_ATTRS = (
    "showGridLines", "rightToLeft", "zoomScale", "zoomScaleNormal",
    "showRowColHeaders", "showZeros", "view", "topLeftCell",
)
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def configure_console():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", line_buffering=True)


def normalized(value):
    return unicodedata.normalize("NFKC", str(value)).strip().casefold()


def find_source_file():
    candidates = [
        path for path in DATA_DIR.glob("*.xlsx")
        if not path.name.startswith("~$")
        and ("大卖家" in path.stem or "已补商品成本" in path.stem)
    ]
    if len(candidates) != 1:
        raise ValueError(f"数据源目录应保留一个 TK 大卖家原表，实际找到：{[p.name for p in candidates]}")
    return candidates[0]


def find_source_sheet(wb):
    for ws in wb.worksheets:
        headers = {
            str(ws.cell(1, col).value).strip()
            for col in range(1, ws.max_column + 1)
            if ws.cell(1, col).value is not None
        }
        if REQUIRED_HEADERS <= headers:
            return ws
    raise ValueError(f"源文件中找不到包含这些列的订单sheet：{sorted(REQUIRED_HEADERS)}")


def header_columns(ws):
    result = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value is not None:
            result[str(value).strip()] = col
    return result


def build_source_index(ws, shop_col):
    result = defaultdict(list)
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, shop_col).value
        if value is not None and str(value).strip():
            result[str(value).strip()].append(row)
    return dict(result)


def resolve_shop_name(stem, source_shops):
    lookup = defaultdict(list)
    for shop in source_shops:
        lookup[normalized(shop)].append(shop)

    stem = stem.strip()
    candidates = [stem]
    if stem.endswith("店"):
        candidates.append(stem[:-1])
    else:
        candidates.append(stem + "店")

    for candidate in candidates:
        matches = lookup.get(normalized(candidate), [])
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise ValueError(f"店铺名匹配不唯一：{stem} -> {matches}")
    compact_matches = []
    keys = {"".join(normalized(candidate).split()) for candidate in candidates}
    for shop in source_shops:
        if "".join(normalized(shop).split()) in keys:
            compact_matches.append(shop)
    if len(compact_matches) > 1:
        raise ValueError(f"店铺名去空格后匹配不唯一：{stem} -> {compact_matches}")
    return compact_matches[0] if compact_matches else None


def copy_cell(src, dst, style_cache):
    dst._value = src._value
    dst.data_type = src.data_type
    if src.has_style:
        style = style_cache.get(src.style_id)
        if style is None:
            style = tuple(
                copy(getattr(src, attr)) if attr != "number_format" else src.number_format
                for attr in STYLE_ATTRS
            )
            style_cache[src.style_id] = style
        for attr, value in zip(STYLE_ATTRS, style):
            setattr(dst, attr, value)
        dst.quotePrefix = src.quotePrefix
        dst.pivotButton = src.pivotButton
    if src.hyperlink is not None:
        dst._hyperlink = copy(src.hyperlink)
    if src.comment is not None:
        dst.comment = copy(src.comment)


def copy_row_dimension(src_ws, src_row, dst_ws, dst_row):
    src_dim = src_ws.row_dimensions.get(src_row)
    if src_dim is None:
        return
    dst_dim = dst_ws.row_dimensions[dst_row]
    for attr in ROW_DIM_ATTRS:
        setattr(dst_dim, attr, getattr(src_dim, attr))


def copy_sheet_settings(src, dst):
    dst.sheet_format = copy(src.sheet_format)
    dst.sheet_properties = copy(src.sheet_properties)
    dst.page_margins = copy(src.page_margins)
    dst.page_setup = copy(src.page_setup)
    dst.print_options = copy(src.print_options)
    dst.sheet_state = src.sheet_state
    dst.freeze_panes = src.freeze_panes
    dst.auto_filter.ref = src.auto_filter.ref
    for attr in VIEW_ATTRS:
        setattr(dst.sheet_view, attr, getattr(src.sheet_view, attr))

    for key, src_dim in src.column_dimensions.items():
        dst_dim = dst.column_dimensions[key]
        for attr in COL_DIM_ATTRS:
            setattr(dst_dim, attr, getattr(src_dim, attr))

    if src.merged_cells.ranges:
        raise ValueError("当前源订单sheet包含合并单元格，无法安全拆分数据行")


def rebuild_sheet(wb, src_ws, row_indexes):
    created = TARGET_SHEET not in wb.sheetnames
    sheet_index = len(wb.sheetnames) if created else wb.sheetnames.index(TARGET_SHEET)
    active_title = wb.active.title
    if not created:
        del wb[TARGET_SHEET]
    dst_ws = wb.create_sheet(TARGET_SHEET, sheet_index)
    copy_sheet_settings(src_ws, dst_ws)

    style_cache = {}
    for col in range(1, src_ws.max_column + 1):
        copy_cell(src_ws.cell(1, col), dst_ws.cell(1, col), style_cache)
    copy_row_dimension(src_ws, 1, dst_ws, 1)

    for dst_row, src_row in enumerate(row_indexes, start=2):
        for col in range(1, src_ws.max_column + 1):
            copy_cell(src_ws.cell(src_row, col), dst_ws.cell(dst_row, col), style_cache)
        copy_row_dimension(src_ws, src_row, dst_ws, dst_row)

    if active_title in wb.sheetnames:
        wb.active = wb.sheetnames.index(active_title)
    return created


def package_sheet_parts(archive):
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root.findall(f"{{{PKG_REL_NS}}}Relationship")
    }
    result = {}
    for sheet in workbook_root.findall(f".//{{{MAIN_NS}}}sheet"):
        rel_id = sheet.attrib[f"{{{DOC_REL_NS}}}id"]
        target = targets[rel_id].lstrip("/")
        if not target.startswith("xl/"):
            target = posixpath.normpath(posixpath.join("xl", target))
        result[sheet.attrib["name"]] = target
    return result


def sheet_rel_part(sheet_part):
    folder, name = posixpath.split(sheet_part)
    return posixpath.join(folder, "_rels", f"{name}.rels")


def merge_content_types(original_data, generated_data):
    original = ET.fromstring(original_data)
    generated = ET.fromstring(generated_data)
    default_tag = f"{{{CONTENT_NS}}}Default"
    override_tag = f"{{{CONTENT_NS}}}Override"
    defaults = {item.attrib["Extension"] for item in generated.findall(default_tag)}
    overrides = {item.attrib["PartName"] for item in generated.findall(override_tag)}
    for item in original.findall(default_tag):
        if item.attrib["Extension"] not in defaults:
            generated.append(copy(item))
    for item in original.findall(override_tag):
        if item.attrib["PartName"] not in overrides:
            generated.append(copy(item))
    ET.register_namespace("", CONTENT_NS)
    return ET.tostring(generated, encoding="utf-8", xml_declaration=True)


def merge_workbook_relationships(original_data, generated_data):
    original = ET.fromstring(original_data)
    generated = ET.fromstring(generated_data)
    rel_tag = f"{{{PKG_REL_NS}}}Relationship"
    existing = {
        (item.attrib.get("Type"), item.attrib.get("Target"), item.attrib.get("TargetMode"))
        for item in generated.findall(rel_tag)
    }
    used_ids = {item.attrib.get("Id") for item in generated.findall(rel_tag)}
    number = 1
    for item in original.findall(rel_tag):
        key = (
            item.attrib.get("Type"),
            item.attrib.get("Target"),
            item.attrib.get("TargetMode"),
        )
        if key in existing:
            continue
        while f"rId{number}" in used_ids:
            number += 1
        cloned = copy(item)
        cloned.attrib["Id"] = f"rId{number}"
        generated.append(cloned)
        used_ids.add(f"rId{number}")
        existing.add(key)
    ET.register_namespace("", PKG_REL_NS)
    return ET.tostring(generated, encoding="utf-8", xml_declaration=True)


def preserve_original_sheets(generated_path, template_path, preserved_sheet_names):
    patched_path = generated_path.with_suffix(".patched.xlsx")
    try:
        with zipfile.ZipFile(template_path, "r") as original, zipfile.ZipFile(
            generated_path, "r"
        ) as generated:
            original_parts = package_sheet_parts(original)
            generated_parts = package_sheet_parts(generated)
            preserved = {}
            skipped_generated = set()

            for sheet_name in preserved_sheet_names:
                src_part = original_parts[sheet_name]
                dst_part = generated_parts[sheet_name]
                preserved[dst_part] = original.read(src_part)
                skipped_generated.add(dst_part)
                src_rel = sheet_rel_part(src_part)
                dst_rel = sheet_rel_part(dst_part)
                skipped_generated.add(dst_rel)
                if src_rel in original.namelist():
                    preserved[dst_rel] = original.read(src_rel)

            generated_data = {
                info.filename: generated.read(info.filename)
                for info in generated.infolist()
                if info.filename not in skipped_generated
            }
            generated_data["[Content_Types].xml"] = merge_content_types(
                original.read("[Content_Types].xml"),
                generated.read("[Content_Types].xml"),
            )
            generated_data["xl/_rels/workbook.xml.rels"] = merge_workbook_relationships(
                original.read("xl/_rels/workbook.xml.rels"),
                generated.read("xl/_rels/workbook.xml.rels"),
            )

            with zipfile.ZipFile(patched_path, "w") as patched:
                written = set()
                for info in original.infolist():
                    name = info.filename
                    data = preserved.get(name, generated_data.get(name, original.read(name)))
                    patched.writestr(info, data)
                    written.add(name)
                for name, data in {**generated_data, **preserved}.items():
                    if name in written:
                        continue
                    if name in generated.namelist():
                        info = generated.getinfo(name)
                    else:
                        info = zipfile.ZipInfo(name)
                        info.compress_type = zipfile.ZIP_DEFLATED
                    patched.writestr(info, data)
        os.replace(patched_path, generated_path)
    finally:
        if patched_path.exists():
            patched_path.unlink()


def save_atomic(wb, output_path, template_path, preserved_sheet_names):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    )
    temp_path = Path(handle.name)
    handle.close()
    try:
        wb.save(temp_path)
        preserve_original_sheets(temp_path, template_path, preserved_sheet_names)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def style_signature(cell, cache):
    if not cell.has_style:
        return None
    if cell.style_id not in cache:
        cache[cell.style_id] = tuple(
            copy(getattr(cell, attr)) if attr != "number_format" else cell.number_format
            for attr in STYLE_ATTRS
        ) + (cell.quotePrefix, cell.pivotButton)
    return cache[cell.style_id]


def dimension_signature(dim, attrs):
    return tuple(getattr(dim, attr) for attr in attrs)


def values_equal(left, right):
    if left in (None, "") and right in (None, ""):
        return True
    if isinstance(left, float) and isinstance(right, float):
        if math.isnan(left) and math.isnan(right):
            return True
    return left == right


def validate_output(output_path, template_path, original_sheetnames, src_ws, row_indexes, order_col, shop_col, shop):
    wb = openpyxl.load_workbook(output_path, data_only=False)
    try:
        if wb.sheetnames != original_sheetnames:
            raise ValueError("输出文件的sheet名称或顺序发生变化")
        ws = wb[TARGET_SHEET]
        expected_rows = len(row_indexes) + 1
        if ws.max_row != expected_rows or ws.max_column != src_ws.max_column:
            raise ValueError(
                f"sheet尺寸不正确：实际 {ws.max_row}x{ws.max_column}，"
                f"应为 {expected_rows}x{src_ws.max_column}"
            )

        src_style_cache = {}
        dst_style_cache = {}
        row_map = [1, *row_indexes]
        for dst_row, src_row in enumerate(row_map, start=1):
            for col in range(1, src_ws.max_column + 1):
                src_cell = src_ws.cell(src_row, col)
                dst_cell = ws.cell(dst_row, col)
                if not values_equal(src_cell.value, dst_cell.value):
                    raise ValueError(f"单元格值不一致：{dst_cell.coordinate}")
                both_empty = src_cell.value in (None, "") and dst_cell.value in (None, "")
                if not both_empty and src_cell.data_type != dst_cell.data_type:
                    raise ValueError(f"单元格数据类型不一致：{dst_cell.coordinate}")
                if style_signature(src_cell, src_style_cache) != style_signature(
                    dst_cell, dst_style_cache
                ):
                    raise ValueError(f"单元格格式不一致：{dst_cell.coordinate}")

            src_dim = src_ws.row_dimensions.get(src_row)
            dst_dim = ws.row_dimensions.get(dst_row)
            src_sig = dimension_signature(src_dim, ROW_DIM_ATTRS) if src_dim else None
            dst_sig = dimension_signature(dst_dim, ROW_DIM_ATTRS) if dst_dim else None
            if src_sig != dst_sig:
                raise ValueError(f"行格式不一致：第 {dst_row} 行")

        all_columns = set(src_ws.column_dimensions) | set(ws.column_dimensions)
        for key in all_columns:
            src_dim = src_ws.column_dimensions.get(key)
            dst_dim = ws.column_dimensions.get(key)
            src_sig = dimension_signature(src_dim, COL_DIM_ATTRS) if src_dim else None
            dst_sig = dimension_signature(dst_dim, COL_DIM_ATTRS) if dst_dim else None
            if src_sig != dst_sig:
                raise ValueError(f"列格式不一致：{key} 列")

        if ws.freeze_panes != src_ws.freeze_panes:
            raise ValueError("冻结窗格与源表不一致")
        if ws.auto_filter.ref != src_ws.auto_filter.ref:
            raise ValueError("筛选范围与源表不一致")
        for attr in ("sheet_format", "sheet_properties", "page_margins", "page_setup", "print_options", "sheet_state"):
            if getattr(ws, attr) != getattr(src_ws, attr):
                raise ValueError(f"sheet 设置与源表不一致：{attr}")
        for attr in VIEW_ATTRS:
            if getattr(ws.sheet_view, attr) != getattr(src_ws.sheet_view, attr):
                raise ValueError(f"sheet 视图设置与源表不一致：{attr}")

        orders = {
            str(value).strip()
            for row in range(2, ws.max_row + 1)
            if (value := ws.cell(row, order_col).value) is not None and str(value).strip()
        }
        shops = {str(ws.cell(row, shop_col).value).strip() for row in range(2, ws.max_row + 1)}
        if shops != ({shop} if row_indexes else set()):
            raise ValueError(f"输出sheet混入其他店铺：{sorted(shops)}")
        with zipfile.ZipFile(template_path) as original, zipfile.ZipFile(output_path) as output:
            original_parts = package_sheet_parts(original)
            output_parts = package_sheet_parts(output)
            for title, part in original_parts.items():
                if title != TARGET_SHEET and original.read(part) != output.read(output_parts[title]):
                    raise ValueError(f"非目标 sheet {title} 原始 XML 发生变化")
        original_wb = openpyxl.load_workbook(template_path, read_only=True, data_only=False)
        try:
            for attr in ("_cell_styles", "_fonts", "_fills", "_borders", "_alignments", "_protections", "_number_formats"):
                original_styles = list(getattr(original_wb, attr))
                if original_styles != list(getattr(wb, attr))[:len(original_styles)]:
                    raise ValueError(f"模板已有样式定义发生变化：{attr}")
        finally:
            original_wb.close()
        return len(row_indexes), len(orders)
    finally:
        wb.close()


def target_files(only, source_file):
    files = sorted(
        path
        for path in TARGET_DIR.glob("*.xlsx")
        if not path.name.startswith("~$") and path.resolve() != source_file.resolve()
    )
    if only:
        wanted = normalized(Path(only).stem)
        files = [path for path in files if normalized(path.stem) == wanted]
    return files


def run(only=None, output_dir=OUTPUT_DIR):
    if not TARGET_DIR.is_dir():
        raise FileNotFoundError(f"找不到目标目录：{TARGET_DIR}")

    source_file = find_source_file()
    files = target_files(only, source_file)
    if not files:
        raise FileNotFoundError("没有找到待处理的 TK 店铺文件")

    print(f"原表：{source_file.name}")
    source_wb = openpyxl.load_workbook(source_file, data_only=False)
    try:
        source_ws = find_source_sheet(source_wb)
        columns = header_columns(source_ws)
        order_col = columns[ORDER_HEADER]
        shop_col = columns[SHOP_HEADER]
        source_index = build_source_index(source_ws, shop_col)

        matches = {}
        for path in files:
            shop = resolve_shop_name(path.stem, source_index)
            if shop is not None and shop in matches.values():
                raise ValueError(f"多个目标文件匹配到同一店铺：{shop}")
            matches[path] = shop

        missing = sorted(set(source_index) - {s for s in matches.values() if s}) if only is None else []
        report = {
            "source": str(source_file),
            "source_sheet": source_ws.title,
            "source_rows": sum(map(len, source_index.values())),
            "source_shops": len(source_index),
            "files": [],
            "unmatched_source_shops": {s: len(source_index[s]) for s in missing},
            "empty_templates": [],
            "created_sheets": [],
            "aliases": [],
            "blank_shop_rows": [],
            "blank_order_rows": [],
            "source_formula_cells": [],
            "source_error_cells": [],
            "zero_cost_rows": {},
            "blank_cost_rows": {},
            "verification": "逐格值、非空类型、样式、行列格式、sheet设置一致；非目标sheet原始XML逐字节一致",
        }
        source_orders = Counter()
        for row in range(2, source_ws.max_row + 1):
            shop_value = source_ws.cell(row, shop_col).value
            order_value = source_ws.cell(row, order_col).value
            if shop_value is None or not str(shop_value).strip():
                if any(source_ws.cell(row, col).value is not None for col in range(1, source_ws.max_column + 1)):
                    report["blank_shop_rows"].append(row)
            if order_value is None or not str(order_value).strip():
                report["blank_order_rows"].append(row)
            else:
                source_orders[str(order_value).strip()] += 1
            for name in ("商品成本", "总成本价"):
                cell = source_ws.cell(row, columns[name])
                if cell.value == 0:
                    report["zero_cost_rows"].setdefault(name, []).append(row)
                elif cell.value is None or cell.value == "":
                    report["blank_cost_rows"].setdefault(name, []).append(row)
            for col in range(1, source_ws.max_column + 1):
                cell = source_ws.cell(row, col)
                if cell.data_type == "f":
                    report["source_formula_cells"].append(cell.coordinate)
                elif cell.data_type == "e":
                    report["source_error_cells"].append({"cell": cell.coordinate, "value": cell.value})
        report["source_unique_orders"] = len(source_orders)
        report["duplicate_orders"] = {k: v for k, v in source_orders.items() if v > 1}

        total_rows = 0
        total_orders = 0
        for path, shop in matches.items():
            rows = source_index.get(shop, [])
            if shop is None:
                report["empty_templates"].append(path.name)
            wb = openpyxl.load_workbook(path, data_only=False)
            try:
                original_sheetnames = wb.sheetnames[:]
                created = rebuild_sheet(wb, source_ws, rows)
                if created:
                    original_sheetnames.append(TARGET_SHEET)
                    report["created_sheets"].append(path.name)
                output_path = output_dir / path.name
                preserved_sheet_names = [s for s in wb.sheetnames if s != TARGET_SHEET]
                save_atomic(wb, output_path, path, preserved_sheet_names)
            finally:
                wb.close()

            row_count, order_count = validate_output(
                output_path,
                path,
                original_sheetnames,
                source_ws,
                rows,
                order_col,
                shop_col,
                shop,
            )
            total_rows += row_count
            total_orders += order_count
            report["files"].append({"name": path.name, "shop": shop, "rows": row_count, "unique_orders": order_count, "preserved_sheets": len(preserved_sheet_names)})
            alias = "" if shop is None or normalized(path.stem) == normalized(shop) else f"，匹配源店铺：{shop}"
            if alias:
                report["aliases"].append({"file": path.name, "shop": shop})
            print(f"[通过] {path.name}：{row_count} 行，{order_count} 个订单{alias}")

        print(
            f"完成：{len(matches)} 个文件，{total_rows} 行，{total_orders} 个唯一订单；"
            f"输出目录：{output_dir}"
        )
        report["output_rows"] = total_rows
        report["output_unique_orders_sum"] = total_orders
        report["output_file_count"] = len(matches)
        report_path = output_dir / ("TK校验报告.json" if only is None else f"TK校验报告_{Path(only).stem}.json")
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        if missing:
            print("[需注意] 以下源店铺没有模板，未复制：" + "；".join(f"{s} {len(source_index[s])}行" for s in missing))
        if report["empty_templates"]:
            print("[需注意] 以下模板无源订单，已写入表头：" + "；".join(report["empty_templates"]))
        if report["created_sheets"]:
            print("自动创建大卖家sheet：" + "；".join(report["created_sheets"]))
        if report["duplicate_orders"]:
            print(f"[需注意] 源表重复订单：{report['duplicate_orders']}")
        print(f"校验报告：{report_path}")
    finally:
        source_wb.close()


def main():
    configure_console()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("shop", nargs="?", help="可选：只处理指定店铺文件名")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR, help="输出目录")
    args = parser.parse_args()
    output_dir = args.output_dir if args.output_dir.is_absolute() else BASE / args.output_dir
    try:
        run(args.shop, output_dir.resolve())
    except Exception as exc:
        print(f"[失败] {exc}", file=sys.stderr)
        raise SystemExit(1) from exc


if __name__ == "__main__":
    main()
