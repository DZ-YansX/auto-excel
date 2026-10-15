# -*- coding: utf-8 -*-
"""
TK 大卖家(BigSeller) 利润表自动化
==================================
输入: TKdeal 文件夹内的待处理店铺表
      (含 利润表/销售明细表/账单原表/订单明细表/广告费/未匹配/大卖家 sheet)
      + 脚本目录下可选成本表 (万山成本价.xlsx，仅用于大卖家也没有的未匹配已发货订单兜底)
输出: output 文件夹内结果表；文件名含“-未完成”时自动去掉该标记

规则要点:
  - 销售明细 = 账单原表 Transaction type=Order 且订单号能在订单明细表中找到的行，按创建日期升序
  - 销售明细成本(J列): 仅目标月份且已发货订单按大卖家整单"总成本价"计入、每单只计一次
  - 旧版大卖家表没有"总成本价"时兼容读取"商品成本"；多行金额自动求和
  - 未匹配 = 订单明细表中不在账单 Order 行里的订单，按状态拆成取消未发货/取消已发货/发货未到账
  - 未匹配已发货类成本优先取大卖家，缺失时按成本表(区分国家)计算
  - 利润表自动识别月份, 已有该月列则填写, 没有则在最后一个月份列右侧新建(样式/公式复制上一列)
"""
import os
import re
import glob
import math
import posixpath
import shutil
import stat
import sys
import tempfile
import warnings
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile
from copy import copy
from collections import Counter

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Color
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

warnings.filterwarnings('ignore')
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# ================= 配置区域 =================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))      # 脚本所在目录
TARGET_FOLDER_PATH = os.path.join(SCRIPT_DIR, 'TKdeal')      # 只放待处理文件
OUTPUT_FOLDER_PATH = os.path.join(SCRIPT_DIR, 'output')      # 只放脚本输出结果
COST_FILE_NAME = '万山成本价.xlsx'                            # 可选成本表，仅用于未匹配已发货兜底

# 账单币种 -> CNY 汇率 (按 账单原表 Currency 列自动选取)
CURRENCY_RATES = {
    'PHP': 0.1065,
    'MYR': 1.629,
    'THB': 0.201,
    'IDR': 0.0003754,
}
USD_RATE = 7          # 广告费 USD -> CNY
UNIT_FEE = 3.5        # 发单费单价(元/单)
COUNTRY_UNIT_FEES = {'IDR': 3.0, 'PHP': 3.5, 'MYR': 3.5, 'THB': 3.5}
UNMATCHED_ORDER_ID_WIDTHS = {'IDR': 20.375, 'MYR': 16.75}

# 订单明细表 Country 列 -> 成本表"仓库"列 的映射
COUNTRY_MAP = {
    'Philippines': '菲律宾',
    'Malaysia': '马来',
    'Malesia': '马来',
    'Thailand': '泰国',
    'Indonesia': '印尼',
    'ماليزيا': '马来',
}

# SKU 别名: 成本表中查不到时按此映射再查 (用户确认 YY-05x 即 YY-KP)
SKU_ALIAS = {
    'YY-054': 'YY-KP',
    'YY-056': 'YY-KP',
    'YY-057': 'YY-KP',
    'YY-059': 'YY-KP',
}

# KP 前缀兜底: 该前缀的 SKU 在大卖家与成本表都查不到时, 统一按 "<前缀>-KP" 的成本计算
# (用户确认 YY-* 即 YY-KP, XY-* 即 XY-KP; 成本按国家区分)。可按需扩展前缀。
KP_FALLBACK_PREFIXES = ('YY', 'XY')

# 账单原表列名中英别名: 不同店铺导出语言不一致(英文/中文),
# find_col 会依次尝试这些别名 (中文按精确匹配, 英文精确优先再子串)
BILL_COL_ALIASES = {
    'oid':     ['Order/Adjustment ID', '订单ID/调整单ID', 'ID Pesanan/Penyesuaian', 'Order ID', '订单ID'],
    'type':    ['Transaction type', '交易类型', 'Jenis transaksi'],
    'created': ['created time', '订单创建时间', '创建时间', 'Waktu pemesanan'],
    'currency':['Currency', '货币', 'Mata uang'],
    'sub':     ['Subtotal before discounts', 'Subtotal before discount', '享受折扣前小计',
                'Subtotal sebelum diskon'],
    'disc':    ['Seller discounts', '商家折扣', 'Diskon penjual'],
    'refund':  ['Refund subtotal after seller discounts', 'Refund amount after discount',
                '享受商家折扣后的退款小计', 'Total pengembalian dana setelah diskon'],
    'fees':    ['Total Fees', '总费用', 'Total Biaya'],
    'adj':     ['Adjustment amount', '调整金额', 'Jumlah penyesuaian'],
}

# 账单原表"交易类型"值归一化 (中文导出为 订单/预扣税, 印尼语导出为 Pesanan)
TXN_TYPE_ALIASES = {'订单': 'order', '预扣税': 'withholding tax', 'pesanan': 'order'}
# ===========================================

YELLOW_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')


def safe_float(x):
    """强力转浮点数"""
    if x is None:
        return 0.0
    try:
        if isinstance(x, (int, float)):
            if isinstance(x, float) and math.isnan(x):
                return 0.0
            return float(x)
        s = str(x).strip().lower()
        if s in ('nan', '', 'none'):
            return 0.0
        s = s.replace(',', '').replace(' ', '').replace('¥', '').replace('₱', '')
        val = float(s)
        return 0.0 if math.isnan(val) else val
    except Exception:
        return 0.0


def sum_numeric_cell(x):
    """汇总单元格中的一个或多个数值。
    大卖家导出在同一订单包含多个商品时，成本列可能是 '6.58\n6.58\n'；
    这里按换行/分号拆分后求和，普通数字仍按原数值读取。"""
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return safe_float(x)
    parts = [p.strip() for p in re.split(r'[\r\n;；]+', str(x)) if p.strip()]
    if len(parts) <= 1:
        return safe_float(x)
    return sum(safe_float(p) for p in parts)


def num_or_int(v):
    """整数值写成 int, 其余保持 float"""
    return int(v) if float(v).is_integer() else v


def fmt_num(v):
    """数字转字符串: 去掉多余的小数尾零 (用于公式与批注文本)"""
    v = round(float(v), 4)
    if v == int(v):
        return str(int(v))
    return f'{v:f}'.rstrip('0').rstrip('.')


def parse_month(s):
    """从日期(字符串或datetime)提取月份, 失败返回 None"""
    if s is None:
        return None
    if hasattr(s, 'month') and not isinstance(s, str):
        return s.month
    m = re.match(r'\s*(\d{4})[/-](\d{1,2})', str(s))
    if m:
        return int(m.group(2))
    m = re.match(r'\s*(\d{1,2})[/-]\d{1,2}[/-]\d{4}', str(s))  # MM/DD/YYYY
    if m:
        return int(m.group(1))
    return None


def parse_detail_month(value):
    """订单明细日期多为 DD/MM/YYYY；菲律宾导出为 MM/DD/YYYY hh:mm:ss AM/PM。用于判断未匹配订单是否跨月。"""
    if value is None:
        return None
    if hasattr(value, 'month') and not isinstance(value, str):
        return value.month
    text = str(value).strip()
    match = re.match(r'(\d{4})[/-](\d{1,2})', text)
    if match:
        return int(match.group(2))
    match = re.match(r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})', text)
    if match:
        first, second = int(match.group(1)), int(match.group(2))
        us_format = re.search(r'\b[AP]M\b', text, re.I) is not None
        if second > 12 or (us_format and first <= 12):
            return first
        return second
    return None


def find_col(header_cells, keyword, default=None):
    """按表头文本查找列号(1-based)。keyword 可为单个关键词或关键词列表(中英别名)。
    先对所有别名做精确匹配, 再做子串匹配, 都找不到返回 default。
    (精确优先: 否则 'Seller discounts' 会错配到 'Subtotal after seller discounts';
     中文子串同理, '商家折扣' 精确匹配到该列而非 '享受商家折扣后小计')"""
    kws = [keyword.lower()] if isinstance(keyword, str) else [k.lower() for k in keyword]
    for kw in kws:
        for c in header_cells:
            if c.value is not None and str(c.value).strip().lower() == kw:
                return c.column
    for kw in kws:
        for c in header_cells:
            if c.value is not None and kw in str(c.value).lower():
                return c.column
    return default


def norm_txn_type(v):
    """账单"交易类型"归一化为小写英文口径 (中文 订单/预扣税 -> order/withholding tax)"""
    s = str(v or '').strip()
    return TXN_TYPE_ALIASES.get(s.lower(), s.lower())


def strip_one_newline(v):
    """去掉字符串结尾的一个 \\n (大卖家导出值尾部是 \\r\\n, 正确表保留 \\r)"""
    if isinstance(v, str) and v.endswith('\n'):
        return v[:-1]
    return v


def normalize_quantity(value):
    """单一数量写成数字；多商品数量保留换行，与大卖家原始结构一致。"""
    value = strip_one_newline(value)
    if not isinstance(value, str):
        return num_or_int(safe_float(value))
    if '\n' in value or '\r' in value.strip('\r'):
        return value
    text = value.strip()
    try:
        return num_or_int(float(text))
    except ValueError:
        return value


def is_shipping_status(value):
    """识别已发货但尚未进入账单的订单状态。"""
    text = str(value or '').strip().lower()
    return text in {
        'shipping', 'shipped', '已发货', '已發貨', '运输中', '運輸中',
        '运送中', '運送中', '配送中', '出货中', '出貨中',
        'sedang dikirim', 'sedang dihantar',
    }


def output_file_name(file_name):
    """待完成样例去掉文件名标记；其他历史文件仍保留“新_”前缀。"""
    stem, extension = os.path.splitext(file_name)
    cleaned = re.sub(r'\s*-\s*未完成\s*$', '', stem)
    return f'{cleaned}{extension}' if cleaned != stem else f'新_{file_name}'


def strip_minus_zero(formula):
    """去掉公式中手工失误留下的 '-0' 项 (如 =F4-0-F5 -> =F4-F5), 不影响 -0.5 等正常数字"""
    return re.sub(r'-0(?=[+\-*/)])|-0$', '', formula)


def load_cost_table(path):
    """成本表 -> {(SKU, 仓库): 成本单价}"""
    cost_map = {}
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        sku, wh, price = row
        if sku is None or wh is None:
            continue
        cost_map[(str(sku).strip(), str(wh).strip())] = safe_float(price)
    wb.close()
    print(f"   已加载成本表 {len(cost_map)} 条 (SKU+仓库)")
    return cost_map


OOXML_MAIN_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
OOXML_DOC_REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
OOXML_PKG_REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
OOXML_CONTENT_TYPES_NS = 'http://schemas.openxmlformats.org/package/2006/content-types'


def _normalize_zip_part(base_dir, target):
    """把 OOXML relationship Target 转为 zip 内的规范相对路径。"""
    target = str(target).replace('\\', '/')
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(base_dir, target))


def _worksheet_part(zf, sheet_name):
    """从工作簿关系中取得指定 sheet 对应的 zip part 路径。"""
    workbook = ET.fromstring(zf.read('xl/workbook.xml'))
    rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
    rel_map = {rel.attrib['Id']: rel.attrib['Target'] for rel in rels}
    sheets = workbook.find(f'{{{OOXML_MAIN_NS}}}sheets')
    for sheet in sheets:
        if sheet.attrib.get('name') == sheet_name:
            rel_id = sheet.attrib[f'{{{OOXML_DOC_REL_NS}}}id']
            return _normalize_zip_part('xl', rel_map[rel_id])
    raise KeyError(f'找不到工作表: {sheet_name}')


def _rels_part(part_name):
    return posixpath.join(posixpath.dirname(part_name), '_rels',
                          posixpath.basename(part_name) + '.rels')


def restore_sheet_drawings(source_path, output_path, sheet_name):
    """在 openpyxl 保存后恢复指定工作表的原始嵌入图片。

    WPS 导出的广告截图属于 openpyxl 无法识别的 DrawingML 对象；直接保存会删除
    worksheet 的 drawing 关系、xl/drawings 与 xl/media。此函数只恢复绘图关系和
    原始媒体字节，保留 openpyxl 已写好的单元格、公式、样式及其他工作表内容。
    """
    replacements = {}
    source_infos = {}

    with ZipFile(source_path, 'r') as src, ZipFile(output_path, 'r') as out:
        src_sheet_part = _worksheet_part(src, sheet_name)
        out_sheet_part = _worksheet_part(out, sheet_name)
        src_sheet_rels_part = _rels_part(src_sheet_part)
        out_sheet_rels_part = _rels_part(out_sheet_part)
        if src_sheet_rels_part not in src.namelist():
            return 0

        src_rels = ET.fromstring(src.read(src_sheet_rels_part))
        drawing_rels = [
            rel for rel in src_rels
            if rel.attrib.get('Type', '').endswith('/drawing')
        ]
        if not drawing_rels:
            return 0

        if out_sheet_rels_part in out.namelist():
            out_rels = ET.fromstring(out.read(out_sheet_rels_part))
        else:
            out_rels = ET.Element(f'{{{OOXML_PKG_REL_NS}}}Relationships')
        used_ids = {rel.attrib.get('Id') for rel in out_rels}

        assigned = []
        copied_parts = set()
        for src_rel in drawing_rels:
            rel_id = src_rel.attrib['Id']
            if rel_id in used_ids:
                n = 1
                while f'rId{n}' in used_ids:
                    n += 1
                rel_id = f'rId{n}'
            used_ids.add(rel_id)
            attrs = dict(src_rel.attrib)
            attrs['Id'] = rel_id
            ET.SubElement(out_rels, f'{{{OOXML_PKG_REL_NS}}}Relationship', attrs)
            assigned.append(rel_id)

            drawing_part = _normalize_zip_part(posixpath.dirname(src_sheet_part),
                                               src_rel.attrib['Target'])
            copied_parts.add(drawing_part)
            drawing_rels_part = _rels_part(drawing_part)
            if drawing_rels_part in src.namelist():
                copied_parts.add(drawing_rels_part)
                drawing_part_rels = ET.fromstring(src.read(drawing_rels_part))
                for rel in drawing_part_rels:
                    target_part = _normalize_zip_part(posixpath.dirname(drawing_part),
                                                      rel.attrib['Target'])
                    if target_part in src.namelist():
                        copied_parts.add(target_part)

        # 在 openpyxl 输出的广告费 worksheet 上重新挂回 drawing 节点。
        out_sheet = ET.fromstring(out.read(out_sheet_part))
        for node in list(out_sheet.findall(f'{{{OOXML_MAIN_NS}}}drawing')):
            out_sheet.remove(node)
        for rel_id in assigned:
            ET.SubElement(out_sheet, f'{{{OOXML_MAIN_NS}}}drawing',
                          {f'{{{OOXML_DOC_REL_NS}}}id': rel_id})
        ET.register_namespace('', OOXML_MAIN_NS)
        ET.register_namespace('r', OOXML_DOC_REL_NS)
        replacements[out_sheet_part] = ET.tostring(
            out_sheet, encoding='utf-8', xml_declaration=True)
        ET.register_namespace('', OOXML_PKG_REL_NS)
        replacements[out_sheet_rels_part] = ET.tostring(
            out_rels, encoding='utf-8', xml_declaration=True)

        for part in copied_parts:
            replacements[part] = src.read(part)
            source_infos[part] = src.getinfo(part)

        # 补回图片扩展名与 drawing part 的 Content Types 声明。
        src_types = ET.fromstring(src.read('[Content_Types].xml'))
        out_types = ET.fromstring(out.read('[Content_Types].xml'))
        existing_defaults = {
            node.attrib.get('Extension') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default')
        }
        existing_overrides = {
            node.attrib.get('PartName') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override')
        }
        copied_extensions = {
            posixpath.splitext(part)[1].lstrip('.').lower()
            for part in copied_parts if part.startswith('xl/media/')
        }
        copied_part_names = {'/' + part for part in copied_parts
                             if part.startswith('xl/drawings/') and part.endswith('.xml')}
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default'):
            ext = node.attrib.get('Extension', '').lower()
            if ext in copied_extensions and ext not in existing_defaults:
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Default',
                              dict(node.attrib))
                existing_defaults.add(ext)
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override'):
            part_name = node.attrib.get('PartName')
            if part_name in copied_part_names and part_name not in existing_overrides:
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Override',
                              dict(node.attrib))
                existing_overrides.add(part_name)
        ET.register_namespace('', OOXML_CONTENT_TYPES_NS)
        replacements['[Content_Types].xml'] = ET.tostring(
            out_types, encoding='utf-8', xml_declaration=True)

        out_infos = {info.filename: info for info in out.infolist()}
        fd, temp_path = tempfile.mkstemp(
            prefix='.drawing_restore_', suffix='.xlsx',
            dir=os.path.dirname(os.path.abspath(output_path)))
        os.close(fd)
        with ZipFile(temp_path, 'w', ZIP_DEFLATED, allowZip64=True) as dst:
            for info in out.infolist():
                if info.filename not in replacements:
                    dst.writestr(info, out.read(info.filename))
            for part, data in replacements.items():
                info = out_infos.get(part) or source_infos.get(part)
                if info is not None:
                    dst.writestr(info, data)
                else:
                    dst.writestr(part, data)
    # Windows 不允许替换仍被 ZipFile 打开的文件，因此必须退出上面的 with 后再替换。
    try:
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
    return len(assigned)


def restore_legacy_annotations(source_path, output_path):
    """保留原表批注的 VML 形状、位置及关系，避免 openpyxl 重建后外观变化。"""
    replacements = {}
    source_infos = {}
    copied_parts = set()
    removed_parts = set()

    with ZipFile(source_path, 'r') as src, ZipFile(output_path, 'r') as out:
        src_names = set(src.namelist())
        out_names = set(out.namelist())
        annotation_types = {
            f'{OOXML_DOC_REL_NS}/comments',
            f'{OOXML_DOC_REL_NS}/vmlDrawing',
        }

        for rel_name in src_names:
            if not (rel_name.startswith('xl/worksheets/_rels/') and rel_name.endswith('.rels')):
                continue
            src_rels = ET.fromstring(src.read(rel_name))
            selected = [node for node in src_rels
                        if node.attrib.get('Type') in annotation_types]
            if not selected:
                continue

            if rel_name in out_names:
                out_rels = ET.fromstring(out.read(rel_name))
            else:
                out_rels = ET.Element(f'{{{OOXML_PKG_REL_NS}}}Relationships')
            for node in list(out_rels):
                if node.attrib.get('Type') in annotation_types:
                    out_rels.remove(node)
            for node in selected:
                ET.SubElement(out_rels, f'{{{OOXML_PKG_REL_NS}}}Relationship',
                              dict(node.attrib))

            sheet_name = posixpath.basename(rel_name)[:-5]  # 去掉 .rels
            sheet_part = posixpath.join('xl/worksheets', sheet_name)
            if sheet_part in src_names and sheet_part in out_names:
                src_sheet = ET.fromstring(src.read(sheet_part))
                out_sheet = ET.fromstring(out.read(sheet_part))
                for node in list(out_sheet.findall(f'{{{OOXML_MAIN_NS}}}legacyDrawing')):
                    out_sheet.remove(node)
                source_legacy = src_sheet.find(f'{{{OOXML_MAIN_NS}}}legacyDrawing')
                if source_legacy is not None:
                    out_sheet.append(ET.fromstring(ET.tostring(source_legacy)))
                ET.register_namespace('', OOXML_MAIN_NS)
                ET.register_namespace('r', OOXML_DOC_REL_NS)
                replacements[sheet_part] = ET.tostring(
                    out_sheet, encoding='utf-8', xml_declaration=True)

            ET.register_namespace('', OOXML_PKG_REL_NS)
            replacements[rel_name] = ET.tostring(
                out_rels, encoding='utf-8', xml_declaration=True)

            for node in selected:
                target = _normalize_zip_part(posixpath.dirname(sheet_part),
                                             node.attrib['Target'])
                if target in src_names:
                    copied_parts.add(target)
                    replacements[target] = src.read(target)
                    source_infos[target] = src.getinfo(target)

        if not copied_parts:
            return 0

        removed_parts.update(
            name for name in out_names
            if ((name.startswith('xl/comments') and name.endswith('.xml'))
                or (name.startswith('xl/drawings/') and name.endswith('.vml')))
        )

        src_types = ET.fromstring(src.read('[Content_Types].xml'))
        out_types = ET.fromstring(out.read('[Content_Types].xml'))
        for node in list(out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override')):
            part_name = node.attrib.get('PartName', '').lstrip('/')
            if part_name.startswith('xl/comments'):
                out_types.remove(node)
        source_part_names = {'/' + part for part in copied_parts}
        existing_overrides = {
            node.attrib.get('PartName') for node in
            out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override')
        }
        for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Override'):
            part_name = node.attrib.get('PartName')
            if part_name in source_part_names and part_name not in existing_overrides:
                ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Override',
                              dict(node.attrib))
                existing_overrides.add(part_name)
        source_vml_default = next(
            (node for node in src_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default')
             if node.attrib.get('Extension', '').lower() == 'vml'), None)
        has_vml_default = any(
            node.attrib.get('Extension', '').lower() == 'vml'
            for node in out_types.findall(f'{{{OOXML_CONTENT_TYPES_NS}}}Default'))
        if source_vml_default is not None and not has_vml_default:
            ET.SubElement(out_types, f'{{{OOXML_CONTENT_TYPES_NS}}}Default',
                          dict(source_vml_default.attrib))
        ET.register_namespace('', OOXML_CONTENT_TYPES_NS)
        replacements['[Content_Types].xml'] = ET.tostring(
            out_types, encoding='utf-8', xml_declaration=True)

        out_infos = {info.filename: info for info in out.infolist()}
        fd, temp_path = tempfile.mkstemp(
            prefix='.annotation_restore_', suffix='.xlsx',
            dir=os.path.dirname(os.path.abspath(output_path)))
        os.close(fd)
        with ZipFile(temp_path, 'w', ZIP_DEFLATED, allowZip64=True) as dst:
            for info in out.infolist():
                if info.filename in removed_parts or info.filename in replacements:
                    continue
                dst.writestr(info, out.read(info.filename))
            for part, data in replacements.items():
                info = out_infos.get(part) or source_infos.get(part)
                dst.writestr(info or part, data)

    try:
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
    return len(copied_parts)


def split_skus(cell):
    """把一个单元格里的多个 SKU 拆成列表 (换行/中英逗号/分号/顿号分隔)。
    注意: 不按空格或 '+'/'*' 拆分 —— SKU 名本身可能含空格(如 'La Roche-B3 30ml')、
    '+'(如 'SPF50+')、'*'(倍数或 'Coffee-*2' 这类含*的正式SKU名)。"""
    s = str(cell or '').strip()
    if not s:
        return []
    return [p.strip() for p in re.split(r'[\r\n,，;；、]+', s) if p.strip()]


def kp_fallback_sku(sku):
    """YY-*/XY-* 等前缀 SKU 的兜底目标 '<前缀>-KP'; 不在配置前缀内返回 None。"""
    m = re.match(r'^([A-Za-z]+)-', str(sku).strip())
    if m and m.group(1).upper() in KP_FALLBACK_PREFIXES:
        return m.group(1).upper() + '-KP'
    return None


def lookup_sku_unit(token, country, cost_map):
    """成本表中解析单个 SKU token 的单位成本(区分国家/仓库):
      1) 整个 token 精确匹配 (兼容 'Coffee-*2' 这类名字里带*的正式SKU)
      2) 'base*N' -> 基础SKU成本 * N
      3) SKU_ALIAS 显式别名
      4) YY-/XY- 前缀兜底到 '<前缀>-KP'
    都查不到返回 None。"""
    token = str(token).strip()
    v = cost_map.get((token, country))
    if v is not None:
        return v
    base, mult = token, 1
    if '*' in token:
        b, _, m = token.rpartition('*')
        if m.strip().isdigit():
            base, mult = b.strip(), int(m.strip())
            v = cost_map.get((base, country))
            if v is not None:
                return v * mult
    alias = SKU_ALIAS.get(base)
    if alias is not None:
        v = cost_map.get((alias, country))
        if v is not None:
            return v * mult
    kp = kp_fallback_sku(base)
    if kp is not None:
        v = cost_map.get((kp, country))
        if v is not None:
            return v * mult
    return None


def resolve_sku_cost(sku_cell, country, cost_map):
    """解析一个 SKU 单元格在指定国家的"单位成本合计"(单元格内多个不同SKU自动求和,
    每个 SKU 支持 'base*N' 倍数)。返回 (成本合计, 未解析token列表)。"""
    total = 0.0
    missing = []
    for token in split_skus(sku_cell):
        unit = lookup_sku_unit(token, country, cost_map)
        if unit is None:
            missing.append(token)
        else:
            total += unit
    return total, missing


def order_cost_from_cost_table(oid, det_lines, country, cost_map):
    """按订单明细行汇总某订单的成本表成本 = Σ(每行SKU单位成本合计 * 该行数量)。
    返回 (订单成本, 未解析token列表)。"""
    total = 0.0
    missing = []
    for sku_cell, qty in det_lines.get(oid, []):
        unit, miss = resolve_sku_cost(sku_cell, country, cost_map)
        missing.extend(miss)
        total += unit * safe_float(qty)
    return round(total, 6), missing


def order_sku_qty(oid, det_lines):
    """大卖家表无该订单时, 从订单明细行拼出 SKU 与数量(与大卖家一行一单口径一致):
    多行则 SKU 换行合并、数量求和。返回 (SKU字符串或None, 数量或None)。"""
    lines = det_lines.get(oid, [])
    skus = [str(s).strip() for s, q in lines if str(s or '').strip()]
    qty = sum(safe_float(q) for s, q in lines)
    sku_str = '\n'.join(skus) if skus else None
    return sku_str, (num_or_int(qty) if lines else None)


def find_pending_pnl_month(ws):
    """查找利润表中已经建好、但“实际产品成本”仍为空的最新月份列。
    本月模板通常会预先建立空月份列；优先使用它可避免少量跨月订单让众数误写历史月份。"""
    product_cost_row = None
    for cell in ws['A']:
        label = str(cell.value or '').replace(' ', '').replace('\t', '')
        if '实际产品成本' in label:
            product_cost_row = cell.row
            break
    if product_cost_row is None:
        return None

    pending = []
    for row in ws.iter_rows(min_row=1, max_row=5):
        for cell in row:
            match = re.fullmatch(r'(\d{1,2})\s*月', str(cell.value or '').strip())
            if match and ws.cell(row=product_cost_row, column=cell.column).value in (None, ''):
                pending.append(int(match.group(1)))
    return max(pending) if pending else None


def find_ads_value(ws_ads):
    """广告费 sheet: 找'广告费'标签下方的数值; 找不到则扫描第一个正数"""
    if ws_ads is None:
        return 0.0
    for row in ws_ads.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() == '广告费':
                below = ws_ads.cell(row=c.row + 1, column=c.column).value
                v = safe_float(below)
                if v > 0:
                    return v
    for row in ws_ads.iter_rows():
        for c in row:
            v = safe_float(c.value)
            if v > 0:
                return v
    return 0.0


def process_file(file_path, cost_map=None):
    cost_map = cost_map or {}
    file_name = os.path.basename(file_path)
    print(f"\n{'=' * 60}")
    print(f">>> [1/5] 正在处理: {file_name}")

    wb_src = openpyxl.load_workbook(file_path)

    # ---- 映射 Sheet ----
    sheet_map = {}
    for name in wb_src.sheetnames:
        if '销售明细' in name:
            sheet_map['sales'] = name
        elif '账单' in name:
            sheet_map['bill'] = name
        elif '订单明细' in name:
            sheet_map['detail'] = name
        elif '补单' in name:
            sheet_map['supplement'] = name
        elif '广告' in name:
            sheet_map['ads'] = name
        elif '未匹配' in name:
            sheet_map['unmatched'] = name
        elif '大卖家' in name:
            sheet_map['dmj'] = name
        elif '利润' in name:
            sheet_map['pnl'] = name
    missing = [k for k in ('bill', 'detail', 'dmj', 'pnl', 'sales') if k not in sheet_map]
    if missing:
        print(f"[错误] 缺少必要 Sheet: {missing}")
        wb_src.close()
        return
    dmj_header = wb_src[sheet_map['dmj']][1]
    if (find_col(dmj_header, '订单号') is None
            or find_col(dmj_header, ['总成本价', '商品成本']) is None):
        wb_src.close()
        raise ValueError('大卖家工作表缺少订单号或总成本价/商品成本列；'
                         '请提供订单成本明细，SKU单价表不能代替整单成本。')
    wb_src.close()

    print(f">>> [2/5] 读取与计算...")

    # 复制一份用于写入 (读与写同一工作簿, 保住全部原样式)
    new_file_name = output_file_name(file_name)
    os.makedirs(OUTPUT_FOLDER_PATH, exist_ok=True)
    output_path = os.path.join(OUTPUT_FOLDER_PATH, new_file_name)
    try:
        if os.path.exists(output_path):
            os.chmod(output_path, stat.S_IWRITE)
            os.remove(output_path)
        shutil.copy(file_path, output_path)
    except PermissionError:
        print(f"[错误] 无法写入 {new_file_name}, 请先在 Excel/WPS 中关闭该文件")
        return
    os.chmod(output_path, stat.S_IWRITE)

    # rich_text=True: 保留单元格内多颜色/字号的富文本 (如利润表 A1 标题)
    wb = openpyxl.load_workbook(output_path, rich_text=True)

    # ================= 读取账单原表 =================
    ws_bill = wb[sheet_map['bill']]
    hdr = ws_bill[1]
    c_oid = find_col(hdr, BILL_COL_ALIASES['oid'], 1)
    c_type = find_col(hdr, BILL_COL_ALIASES['type'], 2)
    c_created = find_col(hdr, BILL_COL_ALIASES['created'], 3)
    c_sub = find_col(hdr, BILL_COL_ALIASES['sub'])
    c_disc = find_col(hdr, BILL_COL_ALIASES['disc'])
    c_refund = find_col(hdr, BILL_COL_ALIASES['refund'])
    c_fees = find_col(hdr, BILL_COL_ALIASES['fees'])
    c_adj = find_col(hdr, BILL_COL_ALIASES['adj'])
    c_cur = find_col(hdr, BILL_COL_ALIASES['currency'], 5)

    currency = None
    bill_rows = []
    for r in range(2, ws_bill.max_row + 1):
        oid = ws_bill.cell(row=r, column=c_oid).value
        if oid is None:
            continue
        if currency is None and c_cur:
            cv = ws_bill.cell(row=r, column=c_cur).value
            if cv:
                currency = str(cv).strip().upper()
        bill_rows.append({
            'r': r,
            'oid': str(oid).strip(),
            'type': norm_txn_type(ws_bill.cell(row=r, column=c_type).value),
            'created': ws_bill.cell(row=r, column=c_created).value,
            'sub': safe_float(ws_bill.cell(row=r, column=c_sub).value) if c_sub else 0.0,
            'disc': safe_float(ws_bill.cell(row=r, column=c_disc).value) if c_disc else 0.0,
            'refund': safe_float(ws_bill.cell(row=r, column=c_refund).value) if c_refund else 0.0,
            'fees': safe_float(ws_bill.cell(row=r, column=c_fees).value) if c_fees else 0.0,
            'adj': safe_float(ws_bill.cell(row=r, column=c_adj).value) if c_adj else 0.0,
        })
    order_rows = [b for b in bill_rows if b['type'] == 'order']

    rate = CURRENCY_RATES.get(currency)
    if rate is None:
        print(f"   [警告] 币种 {currency!r} 未在 CURRENCY_RATES 配置, 汇率按 1.0")
        rate = 1.0
    print(f"   账单行: {len(bill_rows)} (Order: {len(order_rows)}) | 币种: {currency} | 汇率: {rate}")

    # ================= 读取大卖家 =================
    ws_dmj = wb[sheet_map['dmj']]
    hdr = ws_dmj[1]
    c_d_oid = find_col(hdr, '订单号', 1)
    c_d_sku = find_col(hdr, '商品SKU', 4)
    c_d_qty = find_col(hdr, ['实际销售数量', '数量'], 13)
    # 本月新版导出使用"总成本价"作为订单整单成本；旧模板没有该列时兼容"商品成本"。
    c_d_cost = find_col(hdr, ['总成本价', '商品成本'], 26)
    dmj = {}
    for r in range(2, ws_dmj.max_row + 1):
        oid = ws_dmj.cell(row=r, column=c_d_oid).value
        if oid is None:
            continue
        oid = str(oid).strip()
        row_data = {
            'sku': ws_dmj.cell(row=r, column=c_d_sku).value,
            'qty': ws_dmj.cell(row=r, column=c_d_qty).value,
            'cost': sum_numeric_cell(ws_dmj.cell(row=r, column=c_d_cost).value),
        }
        if oid in dmj:
            # 极少数导出会把同一订单拆成多行；整单总成本应把各行相加。
            dmj[oid]['cost'] += row_data['cost']
        else:
            dmj[oid] = row_data
    print(f"   大卖家订单: {len(dmj)} | 成本列: {ws_dmj.cell(row=1, column=c_d_cost).value}")

    # ================= 读取订单明细表 =================
    ws_det = wb[sheet_map['detail']]
    hdr = ws_det[1]
    c_t_oid = find_col(hdr, 'Order ID', 1)
    c_t_status = find_col(hdr, 'Order Status', 2)
    c_t_sku = find_col(hdr, 'Seller SKU', 7)
    c_t_qty = find_col(hdr, 'Quantity', 10)
    c_t_ship = find_col(hdr, 'Shipped Time', 28)
    c_t_created = find_col(hdr, ['Created Time', 'Order Creation Date'])
    c_t_country = find_col(hdr, 'Country', 42)
    # 订单明细表最后保留到 'Checked Status' 列 (其后的列不进未匹配表)
    c_t_last = find_col(hdr, 'Checked Status', ws_det.max_column)

    det_rows = []          # 数据行(第3行起; 第2行是英文说明行)
    order_shipped = {}     # oid -> 任一明细行已发货
    order_statuses = {}    # oid -> 订单状态集合
    country_raw = None
    for r in range(3, ws_det.max_row + 1):
        oid = ws_det.cell(row=r, column=c_t_oid).value
        if oid is None:
            continue
        oid = str(oid).strip()
        shipped = str(ws_det.cell(row=r, column=c_t_ship).value or '').strip()
        status = str(ws_det.cell(row=r, column=c_t_status).value or '').strip()
        order_shipped[oid] = order_shipped.get(oid, False) or bool(shipped)
        order_statuses.setdefault(oid, set()).add(status)
        if country_raw is None:
            cv = ws_det.cell(row=r, column=c_t_country).value
            if cv:
                country_raw = str(cv).strip()
        det_rows.append({
            'r': r,
            'oid': oid,
            'sku': str(ws_det.cell(row=r, column=c_t_sku).value or '').strip(),
            'qty': safe_float(ws_det.cell(row=r, column=c_t_qty).value),
            'shipped': shipped,
            'status': status,
            'created': ws_det.cell(row=r, column=c_t_created).value if c_t_created else None,
        })
    det_oids = set(d['oid'] for d in det_rows)
    det_lines = {}         # oid -> [(SKU单元格, 数量), ...] 用于成本表兜底
    for d in det_rows:
        det_lines.setdefault(d['oid'], []).append((d['sku'], d['qty']))
    country = COUNTRY_MAP.get(country_raw, country_raw)
    print(f"   订单明细行: {len(det_rows)} (订单 {len(det_oids)} 个) | 国家: {country_raw} -> {country}")

    # ================= 匹配 =================
    matched = [b for b in order_rows if b['oid'] in det_oids]
    if not matched:
        print("[错误] 账单与订单明细无任何匹配订单, 跳过")
        wb.close()
        return

    # 月份自动识别: 默认取匹配行创建时间月份众数；若利润表已预建空月份列，则优先写该月份，
    # 避免少量跨月订单导致覆盖历史月份（如本期7月订单中夹有6月下单）。
    months = Counter(m for m in (parse_month(b['created']) for b in matched) if m)
    created_month = months.most_common(1)[0][0]
    pending_month = find_pending_pnl_month(wb[sheet_map['pnl']])
    target_month = pending_month if pending_month is not None else created_month
    if pending_month is not None and pending_month != created_month:
        print(f"   识别月份: {target_month}月 (利润表待填月份优先；订单创建月份分布 {dict(months)})")
    else:
        print(f"   识别月份: {target_month}月 (订单创建月份分布 {dict(months)})")

    # 利润表金额按账单原始顺序汇总；销售明细的日期排序只改变展示顺序。
    billing_sum_refund = sum(row['refund'] for row in matched)
    billing_sum_efg = sum(row['sub'] + row['disc'] + row['refund'] for row in matched)
    billing_platform_fees = sum(row['fees'] for row in matched)

    # 正确表按账单创建日期升序，日期相同则保持账单原行序。
    matched.sort(key=lambda row: (str(row['created'] or ''), row['r']))

    # ================= 生成销售明细数据 =================
    # 成本口径: 销售订单号与大卖家订单号匹配后，直接取大卖家整单"总成本价"。
    # 同一订单即使账单出现多行也只计一次；大卖家无该订单时成本记0并标黄，不再查外部成本表。
    sales_rows = []
    sales_missing_skus = []
    sales_cost_credited = set()
    for b in matched:
        m = dmj.get(b['oid'])
        cross_month = parse_month(b['created']) not in (None, target_month)
        not_shipped = not order_shipped.get(b['oid'], False)
        cost_bearing = b['sub'] != 0 and not cross_month and not not_shipped
        cost_missing = False
        if m is not None:
            sku_val, qty_val = m['sku'], m['qty']
            if cost_bearing and b['oid'] not in sales_cost_credited:
                cost_val = m['cost']
                sales_cost_credited.add(b['oid'])
            else:
                cost_val = 0.0
        else:
            # 大卖家没有该订单: SKU/数量从订单明细表补齐，成本记0并标黄。
            sku_val, qty_val = order_sku_qty(b['oid'], det_lines)
            cost_val = 0.0
            if cost_bearing:
                cost_missing = True
                missing_labels = split_skus(sku_val) or [f'订单:{b["oid"]}']
                sales_missing_skus.extend(missing_labels)
                print(f"   [警告] 销售订单 {b['oid']} 不在大卖家表，成本填0并标黄")
        sales_rows.append({
            'date': b['created'],
            'oid': b['oid'],
            'sku': strip_one_newline(sku_val),
            'qty': normalize_quantity(qty_val),
            'sales': num_or_int(b['sub']),
            'disc': num_or_int(b['disc']),
            'refund': num_or_int(b['refund']),
            'cost': num_or_int(cost_val) if cost_bearing else 0,
            'cost_missing': cost_missing,
            'not_shipped': not_shipped,
            'cross_month': cross_month,
        })

    unique_orders = len(set(s['oid'] for s in sales_rows))
    sales_row_count = len(sales_rows)
    cross_month_rows = sum(1 for s in sales_rows if s['cross_month'])
    not_shipped_orders = len(set(s['oid'] for s in sales_rows if s['not_shipped']))
    sales_cost_total = sum(safe_float(s['cost']) for s in sales_rows)
    sum_refund = billing_sum_refund
    sum_efg = billing_sum_efg
    platform_fees = billing_platform_fees
    print(f"   销售明细: {len(sales_rows)} 行 / {unique_orders} 单 (未发货 {not_shipped_orders} 单)")
    print(f"   大卖家成本匹配: {len(set(s['oid'] for s in sales_rows) & set(dmj))}/{unique_orders} 单"
          f" | 商品成本总额: {fmt_num(sales_cost_total)}")

    # ================= 未匹配 =================
    bill_oids = set(b['oid'] for b in order_rows)
    unmatched = [d for d in det_rows if d['oid'] not in bill_oids]
    category_orders = {'取消未发货': set(), '取消已发货': set(), '发货未到账': set()}
    for item in unmatched:
        oid = item['oid']
        statuses = order_statuses.get(oid, set())
        if any(is_shipping_status(status) for status in statuses):
            label = '发货未到账'
        elif order_shipped.get(oid, False):
            label = '取消已发货'
        else:
            label = '取消未发货'
        item['label'] = label
        item['cost_bearing'] = label in ('取消已发货', '发货未到账')
        month = parse_detail_month(item.get('created'))
        item['cross_month'] = month is not None and month != target_month
        category_orders[label].add(oid)

    category_counts = {label: len(oids) for label, oids in category_orders.items()}
    shipped_unmatched_orders = len(category_orders['取消已发货'] | category_orders['发货未到账'])
    print(f"   未匹配: {len(unmatched)} 行 / {len(set(d['oid'] for d in unmatched))} 单 "
          f"(取消未发货 {category_counts['取消未发货']}, "
          f"取消已发货 {category_counts['取消已发货']}, "
          f"发货未到账 {category_counts['发货未到账']})")

    # 未匹配已发货成本: 同样先取大卖家整单"总成本价"(每订单只计一次),
    # 大卖家无该订单时按成本表(万山)兜底(SKU*N/多SKU求和/YY-·XY-→*-KP, 区分国家)
    unmatched_cost_total = 0.0
    missing_skus = []
    unmatched_cost_credited = set()
    for d in unmatched:
        if not d['cost_bearing']:
            continue
        oid = d['oid']
        if oid in unmatched_cost_credited:
            d['write_cost'] = False
            d['cost'] = None
            continue
        unmatched_cost_credited.add(oid)
        d['write_cost'] = True
        if d['cross_month']:
            d['cost'] = 0.0
            continue
        dm = dmj.get(oid)
        if dm is not None:
            d['cost'] = round(safe_float(dm['cost']), 6)
            unmatched_cost_total += d['cost']
            continue
        order_cost, miss = order_cost_from_cost_table(oid, det_lines, country, cost_map)
        if miss:
            d['cost'] = None
            d['cost_missing'] = True
            missing_skus.extend(miss)
            print(f"   [警告] 成本表({country})找不到 SKU: {miss}, 已标黄")
        else:
            d['cost'] = round(order_cost, 6)
            unmatched_cost_total += d['cost']

    # 调整金额: 仅 Withholding tax 且创建时间在目标月
    adj_total = sum(b['adj'] for b in bill_rows
                    if b['type'] == 'withholding tax'
                    and parse_month(b['created']) == target_month)

    # 广告费
    ads_value = find_ads_value(wb[sheet_map['ads']] if sheet_map.get('ads') else None)
    print(f"   广告费原值: {ads_value} USD | 调整金额({target_month}月 Withholding tax): {adj_total}")

    # ================= 写入 =================
    print(f">>> [3/5] 写入销售明细表...")
    write_sales_sheet(wb[sheet_map['sales']], sales_rows, rate)

    print(f">>> [4/5] 写入未匹配表...")
    write_unmatched_sheet(wb, sheet_map, ws_det, unmatched, c_t_last, currency)

    print(f">>> [5/5] 写入利润表...")
    dispatch_unit_fee = COUNTRY_UNIT_FEES.get(currency, UNIT_FEE)
    write_pnl_sheet(
        wb[sheet_map['pnl']], target_month, rate,
        sum_efg=sum_efg, sum_refund=sum_refund,
        sales_cost_total=sales_cost_total, unmatched_cost_total=unmatched_cost_total,
        sales_row_count=sales_row_count, cross_month_rows=cross_month_rows,
        not_shipped_orders=not_shipped_orders,
        shipped_unmatched_orders=shipped_unmatched_orders,
        category_counts=category_counts, dispatch_unit_fee=dispatch_unit_fee,
        platform_fees=platform_fees, adj_total=adj_total, ads_value=ads_value,
    )

    wb.save(output_path)
    wb.close()
    if sheet_map.get('ads'):
        restored = restore_sheet_drawings(file_path, output_path, sheet_map['ads'])
        if restored:
            print(f"   广告费表图片已保留: {restored} 张")
    restored_annotations = restore_legacy_annotations(file_path, output_path)
    if restored_annotations:
        print(f"   原表批注样式已保留: {restored_annotations} 个部件")
    all_missing = sales_missing_skus + missing_skus
    if all_missing:
        print(f"   [注意] {len(all_missing)} 个 SKU 无成本价(已标黄): {sorted(set(all_missing))}"
              f"  [销售明细 {len(sales_missing_skus)} / 未匹配 {len(missing_skus)}]")
    print(f">>> [成功] 完成: {new_file_name}")


def write_sales_sheet(ws, sales_rows, rate):
    """销售明细严格继承模板第2行样式，仅写业务值与公式。"""
    template_styles = [copy(ws.cell(2, c)._style) for c in range(1, 12)]
    clear_to = max(ws.max_row, len(sales_rows) + 2)
    for row in ws.iter_rows(min_row=2, max_row=clear_to, max_col=max(12, ws.max_column)):
        for cell in row:
            cell.value = None
            cell.comment = None

    mark_font = Font(name='宋体', size=11, bold=True, charset=134, color='FF000000')
    mark_alignment = Alignment(vertical='center')

    for r, item in enumerate(sales_rows, start=2):
        values = {
            1: item['date'], 2: item['oid'], 3: item['sku'], 4: item['qty'],
            5: item['sales'], 6: item['disc'], 7: item['refund'],
            8: f"=E{r}+F{r}+G{r}", 9: f"=H{r}*{rate}",
            10: item['cost'], 11: f"=I{r}-J{r}",
        }
        for c, v in values.items():
            cell = ws.cell(row=r, column=c)
            cell._style = copy(template_styles[c - 1])
            cell.value = v
        if item.get('cost_missing'):
            ws.cell(row=r, column=10).fill = YELLOW_FILL
        if item['not_shipped']:
            cell = ws.cell(row=r, column=12)
            cell.value = '未发货'
            cell.font = mark_font
            cell.alignment = mark_alignment

    last = len(sales_rows) + 1
    total_r = last + 1
    for c in range(5, 12):
        col = get_column_letter(c)
        cell = ws.cell(row=total_r, column=c)
        cell.value = f"=SUM({col}2:{col}{last})"
        cell.comment = None


def write_unmatched_sheet(wb, sheet_map, ws_det, unmatched, c_t_last, currency):
    """复制订单明细原样式，并附加未匹配分类与整单成本。"""
    if sheet_map.get('unmatched') in wb.sheetnames:
        ws = wb[sheet_map['unmatched']]
        if ws.max_row >= 1:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet('未匹配')
        sheet_map['unmatched'] = ws.title

    label_font = Font(name='宋体', size=11, charset=134)
    v_center = Alignment(vertical='center')
    col_label = c_t_last + 1
    col_cost = c_t_last + 2

    def copy_cell(src_r, dst_r, column):
        source = ws_det.cell(src_r, column)
        target = ws.cell(dst_r, column)
        target.value = source.value
        target._style = copy(source._style)
        target_font = copy(target.font)
        target_font.charset = 134
        target_font.color = None
        target.font = target_font
        if source.comment is not None:
            target.comment = copy(source.comment)
        if source.hyperlink is not None:
            target._hyperlink = copy(source.hyperlink)

    # 未匹配表保持默认列宽，仅订单号列按站点模板留出完整显示空间。
    ws.column_dimensions['A'].width = UNMATCHED_ORDER_ID_WIDTHS.get(currency, 18.0)

    # 表头行 + 英文说明行
    for src_r, dst_r in ((1, 1), (2, 2)):
        for c in range(1, c_t_last + 1):
            copy_cell(src_r, dst_r, c)
        ws.row_dimensions[dst_r].height = ws_det.row_dimensions[src_r].height or 15

    for i, d in enumerate(unmatched):
        r = i + 3
        for c in range(1, c_t_last + 1):
            copy_cell(d['r'], r, c)
        ws.row_dimensions[r].height = ws_det.row_dimensions[d['r']].height or 15
        cell = ws.cell(row=r, column=col_label)
        cell.value = d['label']
        cell.font = label_font
        cell.alignment = v_center
        if d['cost_bearing']:
            cost_cell = ws.cell(row=r, column=col_cost)
            cost_cell.font = label_font
            cost_cell.alignment = v_center
            if d.get('write_cost', True):
                if d.get('cost_missing'):
                    cost_cell.value = 0
                    cost_cell.fill = YELLOW_FILL
                else:
                    cost_cell.value = num_or_int(d['cost'])


def write_pnl_sheet(ws, target_month, rate, *, sum_efg, sum_refund,
                    sales_cost_total, unmatched_cost_total,
                    sales_row_count, cross_month_rows, not_shipped_orders,
                    shipped_unmatched_orders, category_counts, dispatch_unit_fee,
                    platform_fees, adj_total, ads_value):
    """利润表: 按目标月份填写金额、发单费公式及未匹配备注。"""
    # ---- A1 标题: 富文本, '利润表'宋体12黑色 + 第二行'单位：CNY'宋体8红色 ----
    ws['A1'] = CellRichText(
        TextBlock(InlineFont(rFont='宋体', charset=134, scheme='minor',
                             b=True, sz=12, color=Color(theme=1)), '利润表'),
        TextBlock(InlineFont(rFont='宋体', charset=134, scheme='minor',
                             b=True, sz=8, color=Color(rgb='FFFF0000')),
                  '\n' + ' ' * 10 + '单位：CNY'),   # 缩进10个空格, 保证整格显示
    )

    # ---- A列行标签映射 ----
    row_map = {}
    for row in ws.iter_rows(min_row=1, max_col=1):
        v = row[0].value
        if v is not None:
            row_map[str(v).replace(' ', '').replace('\t', '')] = row[0].row

    def find_row(key):
        for k, r in row_map.items():
            if key in k:
                return r
        return None

    # ---- 定位/新建月份列 ----
    header_row = None
    month_cols = {}   # 月份 -> 列号
    for row in ws.iter_rows(min_row=1, max_row=5):
        for cell in row:
            if cell.value is None:
                continue
            m = re.fullmatch(r'(\d{1,2})\s*月', str(cell.value).strip())
            if m:
                header_row = cell.row
                month_cols[int(m.group(1))] = cell.column
    if not month_cols:
        print("   [错误] 利润表中找不到任何 'N月' 表头")
        return

    if target_month in month_cols:
        col = month_cols[target_month]
        print(f"   利润表已有 {target_month}月 列: 第{col}列, 直接填写")
        # 清理该列公式中手工失误的 '-0' (如 =F4-0-F5)
        for r in range(header_row, max(row_map.values()) + 1):
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = strip_minus_zero(cell.value)
    else:
        prev_col = max(month_cols.values())
        col = prev_col + 1
        prev_letter = get_column_letter(prev_col)
        new_letter = get_column_letter(col)
        print(f"   利润表新建 {target_month}月 列: 第{col}列 (样式/公式复制自 {prev_letter} 列)")
        last_label_row = max(row_map.values())
        for r in range(header_row, last_label_row + 1):
            src = ws.cell(row=r, column=prev_col)
            dst = ws.cell(row=r, column=col)
            dst._style = copy(src._style)
            if isinstance(src.value, str) and src.value.startswith('='):
                # 公式列字母平移 (=E4-E5 -> =F4-F5), 并清理手工失误的 '-0'
                dst.value = strip_minus_zero(
                    re.sub(rf'\b{prev_letter}(\d+)', rf'{new_letter}\1', src.value))
        ws.cell(row=header_row, column=col).value = f'{target_month}月'
        r15 = find_row('团队运营费用')
        if r15:
            ws.cell(row=r15, column=col).value = 0

    letter = get_column_letter(col)

    def put(key, value, formula=None):
        r = find_row(key)
        if r is None:
            print(f"   [警告] 利润表找不到行: {key}")
            return
        cell = ws.cell(row=r, column=col)
        cell.value = formula if formula is not None else value
        cell.comment = None

    refund_cny = abs(sum_refund) * rate
    put('本期销售', round(sum_efg * rate + refund_cny, 10))
    put('销售退款', round(refund_cny, 10))
    put('实际产品成本', round(sales_cost_total + unmatched_cost_total, 4))
    put('发单费', None,
        formula=(f'=({sales_row_count}-{cross_month_rows}-{not_shipped_orders}'
                 f'+{shipped_unmatched_orders})*{fmt_num(dispatch_unit_fee)}'))
    put('广告费', None, formula=f'={fmt_num(ads_value)}*{fmt_num(USD_RATE)}')
    put('平台费用', round(abs(platform_fees) * rate, 10))
    put('调整金额', round(abs(adj_total) * rate, 10))

    # ---- 底部红字备注: '四、本期利润' 的下一行 A 列 ----
    profit_row = find_row('本期利润')
    if profit_row:
        note_cell = ws.cell(row=profit_row + 1, column=1)
        parts = [f'{label}{category_counts[label]}'
                 for label in ('取消未发货', '取消已发货', '发货未到账')
                 if category_counts.get(label)]
        if not parts:
            note_cell.value = None
            return
        note_cell.value = f'{target_month}月未匹配：' + '，'.join(parts)
        note_cell.font = Font(name='宋体', size=10, bold=True,
                              charset=134, color='FFFF0000')
        note_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def main():
    if not os.path.exists(TARGET_FOLDER_PATH):
        print(f"找不到文件夹: {TARGET_FOLDER_PATH}")
        return
    os.makedirs(OUTPUT_FOLDER_PATH, exist_ok=True)
    print(f">>> 待处理目录: {TARGET_FOLDER_PATH}")
    print(f">>> 输出目录: {OUTPUT_FOLDER_PATH}")

    # 成本表是脚本公共配置，不属于待处理店铺文件，放在脚本目录下。
    cost_path = next((p for p in (os.path.join(SCRIPT_DIR, COST_FILE_NAME),
                                  os.path.join(SCRIPT_DIR, '..', '补成本', COST_FILE_NAME))
                      if os.path.exists(p)), None)
    if cost_path:
        print(f">>> 加载可选成本表: {cost_path}")
        cost_map = load_cost_table(cost_path)
    else:
        # 销售明细成本只取大卖家；没有外部成本表不影响正常处理。
        # 仅当未匹配已发货订单也不在大卖家时，相关成本会填0并标黄。
        cost_map = {}
        print(">>> 未提供万山成本价.xlsx：销售成本按大卖家整单成本计算")

    xlsx_files = sorted(glob.glob(os.path.join(TARGET_FOLDER_PATH, "*.xlsx")))
    source_files = []
    for f in xlsx_files:
        name = os.path.basename(f)
        if name == COST_FILE_NAME or name.startswith('新_') or name.startswith('~$'):
            continue
        source_files.append(f)
    print(f"共找到 {len(source_files)} 个原表格: {[os.path.basename(f) for f in source_files]}")
    for f in source_files:
        try:
            process_file(f, cost_map)
        except Exception:
            import traceback
            print(f"[失败] 处理 {os.path.basename(f)} 出错:")
            traceback.print_exc()
    print("\n全部执行完毕！")
    input("按回车键退出...")


if __name__ == '__main__':
    main()
